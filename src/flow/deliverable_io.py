"""Round-trip verified writes to a deliverable — trust the DISK, never memory.

Born from the Bowen external-verification scare: a report claimed edits that a
verifier could not see. Whatever the verifier actually read, the process gap was
real — nothing ever re-read the file FROM DISK after saving to prove the bytes
carry the changes. From now on that proof is mandatory:

- every save goes through ``save_and_verify``: save, then a FRESH load from disk
  asserting the expected sheet set, edited cells, and audit-trail row count;
- the final gate always runs over the file on disk (``revalidate_from_state``
  loads from ``state.ddr_path``), never over an in-memory object.

A failed assertion raises ``RoundTripError`` naming exactly what the disk does
not show — a report can then never describe an edit the disk does not carry.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any


class RoundTripError(AssertionError):
    """The file on disk does not show what the save was supposed to write."""


class FrozenArtifactError(PermissionError):
    """A write was attempted inside a delivered, frozen package."""


# Folders that hold SHIPPED work. Bowen A and Maxwell A live on only as
# regression fixtures; the packages under "<UNIT>/SUBMISSION PLM/" are the
# artefacts the client received. Nothing this repo runs may write a byte inside
# them — not a regenerated sidecar, not a re-packaged tree, not a "harmless"
# refresh. Reading is always fine.
FROZEN_MARKERS: tuple[str, ...] = ("SUBMISSION PLM",)

# Delivered trees that the marker alone does not catch. Josh B was shipped as
# "<unit>/1. Landman Submitted/" with no SUBMISSION PLM wrapper, so the folder
# name proves nothing — the unit has to be named. Populated from config
# (`frozen_paths`), which is gitignored because it holds client paths.
_FROZEN_PATHS: list[Path] = []


def register_frozen_paths(paths) -> list[Path]:
    """Declare additional delivered trees as read-only. Idempotent."""
    for raw in paths or ():
        p = Path(raw).resolve()
        if p not in _FROZEN_PATHS:
            _FROZEN_PATHS.append(p)
    return list(_FROZEN_PATHS)


def frozen_paths() -> list[Path]:
    return list(_FROZEN_PATHS)


def clear_frozen_paths() -> None:
    """Test seam only — production never un-freezes a delivered tree."""
    _FROZEN_PATHS.clear()


def assert_not_frozen(path: str | Path) -> Path:
    """Raise FrozenArtifactError if `path` lies inside a delivered package.

    Two fences: a folder MARKER ("SUBMISSION PLM") and an explicit list of
    delivered trees registered from config. The second exists because Josh B
    shipped without the marker — a delivered unit must be protected by what it
    IS, not by what its folder happens to be called.

    Called by every write path this repo owns (deliverable saves, validation
    sidecars, submission staging). A replay reads shipped work; it never
    touches it."""
    p = Path(path)
    parts = {part.upper() for part in p.parts}
    for marker in FROZEN_MARKERS:
        if marker.upper() in parts:
            raise FrozenArtifactError(
                f"refusing to write inside a delivered package: {p}\n"
                f"'{marker}' folders are frozen — shipped work is read-only "
                f"(CLAUDE.md invariant 16). Write elsewhere, or read instead.")
    resolved = p.resolve() if p.is_absolute() else (Path.cwd() / p).resolve()
    for frozen in _FROZEN_PATHS:
        if resolved == frozen or frozen in resolved.parents:
            raise FrozenArtifactError(
                f"refusing to write inside a delivered unit: {p}\n"
                f"{frozen} is registered as delivered (config frozen_paths) — "
                f"shipped work is read-only (CLAUDE.md invariant 16).")
    # The write-roots fence (invariant 26) shares this chokepoint: every write
    # path already consults the frozen guard, so declaring it here means a new
    # writer cannot forget the manifest any more than it can forget the freeze.
    from src.knowledge.write_roots import record_write
    record_write(resolved)
    return p


def roundtrip_verify(
    path: str | Path,
    *,
    expect_sheets: set[str] | None = None,
    expect_absent_sheets: set[str] | None = None,
    expect_cells: dict[tuple[str, str], Any] | None = None,
    expect_audit_rows: int | None = None,
    audit_sheet: str = "DD_Audit_Trail",
) -> None:
    """Fresh-load ``path`` from disk and assert the expectations hold.

    expect_sheets: the workbook's EXACT visible+hidden tab set, when given.
    expect_absent_sheets: tabs that must NOT exist (e.g. a declared exclusion —
        catches a resurrected copy of an older file).
    expect_cells: {(sheet, "S15"): value} — compared as strings, so a formula
        expectation "=1/12" matches exactly what openpyxl stores.
    expect_audit_rows: exact count of non-empty audit-trail entries.

    Raises RoundTripError listing every expectation the DISK file fails."""
    import openpyxl

    path = Path(path)
    problems: list[str] = []
    if not path.exists():
        raise RoundTripError(f"{path} does not exist on disk")
    wb = openpyxl.load_workbook(path)
    try:
        names = set(wb.sheetnames)
        if expect_sheets is not None and names != set(expect_sheets):
            problems.append(
                f"sheet set mismatch: extra={sorted(names - set(expect_sheets))} "
                f"missing={sorted(set(expect_sheets) - names)}")
        for absent in sorted(expect_absent_sheets or ()):
            if absent in names:
                problems.append(f"sheet {absent!r} must be absent but exists on disk")
        for (sheet, cell), want in (expect_cells or {}).items():
            if sheet not in names:
                problems.append(f"cell {sheet}!{cell}: sheet missing on disk")
                continue
            got = wb[sheet][cell].value
            if str(got) != str(want):
                problems.append(f"cell {sheet}!{cell}: disk={got!r}, expected={want!r}")
        if expect_audit_rows is not None:
            if audit_sheet not in names:
                problems.append(f"audit sheet {audit_sheet!r} missing on disk")
            else:
                ws = wb[audit_sheet]
                n = sum(1 for row in ws.iter_rows(min_row=2, values_only=True) if row[0])
                if n != expect_audit_rows:
                    problems.append(
                        f"audit rows: disk has {n}, expected {expect_audit_rows}")
    finally:
        wb.close()
    if problems:
        raise RoundTripError(
            f"disk file {path.name} failed round-trip verification:\n  - "
            + "\n  - ".join(problems))


def save_and_verify(workbook: Any, path: str | Path, **expectations: Any) -> None:
    """Save the workbook to ``path`` and immediately prove the DISK carries the
    expected state (see roundtrip_verify). The verification load is a brand-new
    openpyxl object — nothing is trusted from the one just saved."""
    path = assert_not_frozen(path)
    workbook.save(path)
    roundtrip_verify(path, **expectations)
