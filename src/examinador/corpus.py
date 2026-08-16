"""The corpus, read in place — inventory, pages, search.

Three dumb operations. None of them decides anything: they do not rank, filter
by relevance, guess a category or pick a page. The model asks for a document and
a page and a mode, and gets exactly that.

`read_page` has three modes and the model chooses between them, because that
choice is a judgement about cost and legibility:

  texto    the embedded text layer, when the file has one. Cheapest by far.
  ocr      tesseract over the rendered page, for scans.
  imagem   the page itself, in pixels, for the model's own eyes.

A mode that yields nothing returns nothing and says so. It never silently falls
through to a costlier mode — that is the code choosing, and the whole point of
this package is that it does not.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from src.examinador.cache import PageCache, digest, key_for
from src.knowledge.corpus_reader import CorpusIndex

IMAGE_SUFFIXES = {".tif", ".tiff", ".png", ".jpg", ".jpeg", ".bmp"}
# Files that ARE text. They have no pages: the whole file is page 1.
TEXT_SUFFIXES = {".csv", ".txt", ".html", ".htm", ".json", ".xml", ".md"}
# Workbooks. Their "pages" are SHEETS, windowed — an XREF runs to tens of
# thousands of rows and the model should ask for the window it needs, not the
# file. The convention the model uses in `pagina`:
#     0       the listing: sheet names, row counts, window counts
#     N       sheet N (1-based), first window
#     "N:J"   sheet N, window J
XLSX_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm"}
WINDOW_ROWS = 200
RENDER_DPI = 200
MAX_EDGE_PX = 7800


@dataclass
class Document:
    name: str                       # the logical path the model uses to ask
    folder: str
    suffix: str
    pages: int | None
    size: int

    def to_dict(self) -> dict[str, Any]:
        return {"documento": self.name, "pasta": self.folder,
                "tipo": self.suffix, "paginas": self.pages, "bytes": self.size}


class Corpus:
    def __init__(self, files_dir: str | Path,
                 cache: PageCache | None = None) -> None:
        self.index = CorpusIndex(files_dir)
        self.cache = cache or PageCache()
        self._bytes: dict[str, bytes] = {}

    # --- bytes ------------------------------------------------------------
    def read_bytes(self, name: str) -> bytes:
        if name not in self._bytes:
            self._bytes[name] = self.index.read_bytes(name)
        return self._bytes[name]

    def names(self) -> list[str]:
        return sorted(self.index._members)          # noqa: SLF001 — the index IS the list

    # --- the three operations ---------------------------------------------
    def list_documents(self, contains: str = "") -> list[dict[str, Any]]:
        """Every document, or every document whose path contains a substring.

        A plain substring, not a relevance ranking: the model does the choosing
        and needs to see what it is choosing from."""
        needle = str(contains or "").lower()
        out: list[dict[str, Any]] = []
        for name in self.names():
            if needle and needle not in name.lower():
                continue
            hit = self.index.resolve(name)
            archive, member = hit if hit else (None, "")
            folder = name.rsplit("/", 1)[0] if "/" in name else ""
            out.append(Document(
                name=name, folder=folder, suffix=Path(name).suffix.lower(),
                pages=None,
                size=(archive.stat().st_size if archive and not member else 0),
            ).to_dict())
        return out

    def page_count(self, name: str) -> int:
        data = self.read_bytes(name)
        suffix = Path(name).suffix.lower()
        if suffix == ".pdf":
            import fitz
            with fitz.open(stream=data, filetype="pdf") as doc:
                return doc.page_count
        if suffix in IMAGE_SUFFIXES:
            from PIL import Image
            img = Image.open(io.BytesIO(data))
            return int(getattr(img, "n_frames", 1))
        if suffix in XLSX_SUFFIXES:
            return self._xlsx_sheet_count(name, data)
        return 1

    def read_page(self, name: str, page: int | str, mode: str = "texto"
                  ) -> dict[str, Any]:
        data = self.read_bytes(name)
        if Path(name).suffix.lower() in XLSX_SUFFIXES:
            return self._read_xlsx(name, data, page, mode)
        try:
            page = int(str(page))
        except ValueError:
            raise ValueError(
                f"pagina {page!r}: o formato aba:janela so vale para "
                f"planilhas; {name} usa numero de pagina")
        doc_key = digest(data)
        if mode == "imagem":
            key = key_for(doc_key, page, "png", RENDER_DPI)
            png = self.cache.get_bytes(key)
            if png is None:
                png = self._render(data, name, page)
                self.cache.put_bytes(key, png)
            return {"kind": "image", "png": png, "documento": name,
                    "pagina": page}

        key = key_for(doc_key, page, mode)
        text = self.cache.get_text(key)
        if text is None:
            text = (self._text_layer(data, name, page) if mode == "texto"
                    else self._ocr(data, name, page))
            self.cache.put_text(key, text, {"documento": name, "pagina": page,
                                            "modo": mode})
        return {"kind": "text", "text": text, "documento": name,
                "pagina": page, "modo": mode, "caracteres": len(text)}

    def search(self, term: str) -> dict[str, Any]:
        """Where a string appears — in file paths, and in text layers.

        Files with no text layer cannot be searched without paying to OCR them,
        so they are LISTED as unsearched rather than quietly treated as
        non-matches. An empty result over a corpus of scans means "not found in
        what could be searched", and the model has to be told which is which."""
        needle = str(term or "").lower()
        in_names, in_text, unsearchable = [], [], []
        for name in self.names():
            if needle in name.lower():
                in_names.append(name)
            suffix = Path(name).suffix.lower()
            if suffix in TEXT_SUFFIXES:
                body = self.read_bytes(name).decode("utf-8", errors="replace")
                if needle in body.lower():
                    in_text.append({"documento": name, "paginas": [1]})
                continue
            if suffix in XLSX_SUFFIXES:
                # As "páginas" devolvidas são especificações aba:janela — o
                # modelo as passa direto a read_page para ver as linhas.
                try:
                    data = self.read_bytes(name)
                    hits: list[str] = []
                    for i in range(1, self._xlsx_sheet_count(name, data) + 1):
                        _, linhas = self._xlsx_sheet(name, data, i)
                        janelas = sorted({
                            idx // WINDOW_ROWS + 1
                            for idx, linha in enumerate(linhas)
                            if needle in linha.lower()})
                        hits.extend(f"{i}:{j}" for j in janelas)
                    if hits:
                        in_text.append({"documento": name, "paginas": hits})
                except Exception as exc:
                    unsearchable.append(f"{name} ({exc.__class__.__name__})")
                continue
            if suffix != ".pdf":
                unsearchable.append(name)
                continue
            try:
                data = self.read_bytes(name)
                import fitz
                with fitz.open(stream=data, filetype="pdf") as doc:
                    pages = []
                    empty = 0
                    for i in range(doc.page_count):
                        text = doc[i].get_text()
                        if not text.strip():
                            empty += 1
                        if needle in text.lower():
                            pages.append(i + 1)
                    if empty == doc.page_count:
                        unsearchable.append(name)
                    elif pages:
                        in_text.append({"documento": name, "paginas": pages})
            except Exception as exc:                # unreadable is a fact, not a match
                unsearchable.append(f"{name} ({exc.__class__.__name__})")
        return {"termo": term, "no_nome": in_names, "no_texto": in_text,
                "nao_pesquisavel_sem_ocr": unsearchable,
                "nota": ("Arquivos sem camada de texto nao foram pesquisados. "
                         "Ausencia aqui nao e prova de ausencia no corpus.")}

    # --- workbooks ---------------------------------------------------------
    # Extraction is transport, not interpretation: cells joined by " | ",
    # rows by newline, values as openpyxl computes them (data_only). Cached
    # per (file hash, sheet) in the GLOBAL cache — an XREF read for one unit
    # is paid once for all 300.

    def _read_xlsx(self, name: str, data: bytes, page: int | str,
                   mode: str) -> dict[str, Any]:
        if mode != "texto":
            return {"kind": "text", "text": "", "documento": name,
                    "pagina": str(page), "modo": mode, "caracteres": 0,
                    "nota": ("planilha nao tem ocr nem imagem; use modo "
                             "`texto` — pagina 0 lista as abas, N le a aba "
                             "N, 'N:J' le a janela J")}
        sheet, window = _parse_sheet_spec(page)
        total = self._xlsx_sheet_count(name, data)
        if sheet == 0:
            text = self._xlsx_listing(name, data)
            return {"kind": "text", "text": text, "documento": name,
                    "pagina": 0, "modo": "texto", "abas_total": total,
                    "caracteres": len(text)}
        if not 1 <= sheet <= total:
            return {"kind": "text", "text": "", "documento": name,
                    "pagina": f"{sheet}:{window}", "modo": "texto",
                    "caracteres": 0,
                    "nota": f"aba {sheet} nao existe; o arquivo tem {total} "
                            f"aba(s) — pagina 0 lista todas"}
        titulo, linhas = self._xlsx_sheet(name, data, sheet)
        janelas = max(1, -(-len(linhas) // WINDOW_ROWS))
        if not 1 <= window <= janelas:
            return {"kind": "text", "text": "", "documento": name,
                    "pagina": f"{sheet}:{window}", "modo": "texto",
                    "aba": titulo, "janelas_total": janelas, "caracteres": 0,
                    "nota": f"janela {window} nao existe; a aba {titulo!r} "
                            f"tem {len(linhas)} linhas em {janelas} janela(s)"}
        corpo = "\n".join(linhas[(window - 1) * WINDOW_ROWS:
                                 window * WINDOW_ROWS])
        return {"kind": "text", "text": corpo, "documento": name,
                "pagina": f"{sheet}:{window}", "modo": "texto",
                "aba": titulo, "abas_total": total, "janela": window,
                "janelas_total": janelas, "linhas_na_aba": len(linhas),
                "caracteres": len(corpo)}

    def _xlsx_sheet_count(self, name: str, data: bytes) -> int:
        key = key_for(digest(data), "xlsx-abas")
        got = self.cache.get_text(key)
        if got is not None:
            return int(got)
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
        try:
            n = len(wb.worksheets)
        finally:
            wb.close()
        self.cache.put_text(key, str(n), {"documento": name})
        return n

    def _xlsx_sheet(self, name: str, data: bytes, sheet: int
                    ) -> tuple[str, list[str]]:
        """(título, linhas) da aba `sheet` (1-based), do cache ou do arquivo."""
        doc_key = digest(data)
        key = key_for(doc_key, "xlsx-aba", sheet)
        key_nome = key_for(doc_key, "xlsx-nome", sheet)
        corpo = self.cache.get_text(key)
        titulo = self.cache.get_text(key_nome)
        if corpo is not None and titulo is not None:
            return titulo, corpo.splitlines()
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True,
                                    data_only=True)
        try:
            ws = wb.worksheets[sheet - 1]
            titulo = ws.title
            linhas: list[str] = []
            for row in ws.iter_rows(values_only=True):
                vals = ["" if v is None else str(v) for v in row]
                while vals and vals[-1] == "":
                    vals.pop()
                linhas.append(" | ".join(vals))
            while linhas and not linhas[-1].strip():
                linhas.pop()
        finally:
            wb.close()
        self.cache.put_text(key, "\n".join(linhas),
                            {"documento": name, "aba": titulo})
        self.cache.put_text(key_nome, titulo, {"documento": name})
        return titulo, linhas

    def _xlsx_listing(self, name: str, data: bytes) -> str:
        total = self._xlsx_sheet_count(name, data)
        out = [f"{name} — {total} aba(s). Leia com pagina N ou 'N:J' "
               f"(janelas de {WINDOW_ROWS} linhas)."]
        for i in range(1, total + 1):
            titulo, linhas = self._xlsx_sheet(name, data, i)
            janelas = max(1, -(-len(linhas) // WINDOW_ROWS))
            out.append(f"  {i}. {titulo} — {len(linhas)} linhas, "
                       f"{janelas} janela(s)")
        return "\n".join(out)

    # --- the dumb transformations -----------------------------------------
    def _text_layer(self, data: bytes, name: str, page: int) -> str:
        suffix = Path(name).suffix.lower()
        if suffix in TEXT_SUFFIXES:
            # A CSV is text. The first round's `texto` mode returned nothing for
            # one because only PDFs were handled, and the model — correctly,
            # from what it could see — wrote a knowledge card telling every
            # future unit not to waste a call on them. A tool that lies about
            # what it cannot read teaches the wrong lesson to 300 units.
            return data.decode("utf-8", errors="replace").strip()
        if suffix != ".pdf":
            return ""
        import fitz
        with fitz.open(stream=data, filetype="pdf") as doc:
            if not 1 <= page <= doc.page_count:
                return ""
            return doc[page - 1].get_text().strip()

    def _ocr(self, data: bytes, name: str, page: int) -> str:
        from PIL import Image

        from src.knowledge.classifier import default_ocr
        png = self._render(data, name, page)
        return str(default_ocr()(Image.open(io.BytesIO(png)))).strip()

    def _render(self, data: bytes, name: str, page: int) -> bytes:
        suffix = Path(name).suffix.lower()
        if suffix == ".pdf":
            import fitz
            with fitz.open(stream=data, filetype="pdf") as doc:
                pix = doc[page - 1].get_pixmap(dpi=RENDER_DPI)
                return _fit(pix.tobytes("png"))
        from PIL import Image
        Image.MAX_IMAGE_PIXELS = None
        img = Image.open(io.BytesIO(data))
        img.seek(page - 1)
        buf = io.BytesIO()
        img.convert("RGB").save(buf, format="PNG")
        return _fit(buf.getvalue())


def _parse_sheet_spec(page: int | str) -> tuple[int, int]:
    """'N:J' -> (aba N, janela J); N -> (N, 1); 0 -> a listagem."""
    text = str(page).strip()
    if ":" in text:
        aba, janela = text.split(":", 1)
        return int(aba), int(janela)
    return int(text), 1


def _fit(png: bytes) -> bytes:
    """The API refuses an image whose side exceeds its limit. A page too large
    to send is downscaled, never dropped."""
    from PIL import Image
    Image.MAX_IMAGE_PIXELS = None
    img = Image.open(io.BytesIO(png))
    if max(img.size) <= MAX_EDGE_PX:
        return png
    scale = MAX_EDGE_PX / max(img.size)
    img = img.resize((int(img.width * scale), int(img.height * scale)),
                     Image.LANCZOS)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()
