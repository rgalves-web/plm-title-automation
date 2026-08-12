"""Round 1 — the Examiner is born and writes the whole DDR.

One agent, one corpus, one workbook. No gate, no vocabulary, no formula, none of
the twenty modules the architecture retires. What comes out is a number: how many
cells the model gets right on its own against the 1,178 evidenced cells of the
current pipeline and against the hand-built DDR.

The answer sheet is quarantined for the whole run. A replay that reads the
answers is not a replay, and the number it produces is not evidence.
"""
from __future__ import annotations

import json
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl

from src.examinador.cache import PageCache
from src.examinador.corpus import Corpus
from src.examinador.ledger import Ledger
from src.examinador.loop import run_loop
from src.examinador.okf import KnowledgeBase
from src.examinador.tools import TOOLS, build_dispatch
from src.knowledge.quarantine import quarantined

PROJECT = Path("C:/Users/Renato/PURPLE LAND MGMT/BKV122TX - Project Emirates")
FILES_DIR = PROJECT / "JOSH B UNIT" / "FILES"
TEMPLATE = PROJECT / "DDR - Template (PLM).xlsx"
OUT_DIR = Path("data/units/_josh_b_replay/KIT_IA")
UNIT = "JOSH B UNIT"
COUNTY = "Montague County, Texas"
MODEL = "claude-opus-5"

SYSTEM = """Voce e o landman responsavel por esta unidade. Nao e um assistente
que ajuda um landman: o trabalho e seu, e o entregavel sai com o seu nome.

O corpus da unidade esta nas ferramentas. Peca os documentos que precisar, na
ordem que quiser, quantas vezes precisar.

Tres coisas valem sempre:

1. Toda celula precisa de fonte. Um valor sem citacao de documento, pagina e
   trecho literal nao e trabalho de title — e chute com formatacao boa.
2. O que voce nao puder provar, deixe em branco com motivo e dono, ou pergunte.
   Branco declarado e melhor que numero plausivel: quem recebe este relatorio
   toma decisao de milhoes com ele, e um numero errado custa mais caro que um
   branco honesto.
3. O que voce aprender que vale alem desta unidade, grave na base de
   conhecimento. Sao 300 unidades do mesmo Seller, no mesmo condado, com os
   mesmos formularios e as mesmas cadeias de cessao. O que voce ler uma vez, as
   outras 299 consultam em vez de reler.

Ninguem vai conferir automaticamente o que voce escrever. Nao ha portao, nao ha
validacao, nao ha formula em codigo. O que voce escrever e o que vai ao cliente.
"""


@dataclass
class RoundResult:
    workbook: Path
    ledger: Path
    log: Path
    stop_reason: str
    counts: dict[str, Any]

    def to_dict(self) -> dict[str, Any]:
        return {"workbook": str(self.workbook), "ledger": str(self.ledger),
                "log": str(self.log), "parou_porque": self.stop_reason,
                "contagens": self.counts}


def template_map(path: str | Path = TEMPLATE, per_cell: int = 60) -> str:
    """Every non-empty cell of the template, as text.

    The model cannot open a workbook — it has nine tools and none of them reads
    xlsx. So the template arrives as data in the assignment: the tabs, and what
    is already written in them. This function transcribes; it does not decide
    which cells are structure, which are another unit's leftovers, or which
    should survive into the deliverable. That judgement is the examiner's, and
    round 1 measures whether it gets made."""
    wb = openpyxl.load_workbook(Path(path), data_only=True)
    lines: list[str] = []
    try:
        for ws in wb.worksheets:
            filled = [(c.coordinate, c.value) for row in ws.iter_rows()
                      for c in row if c.value not in (None, "")]
            lines.append(f"\n### aba `{ws.title}` — {len(filled)} celulas "
                         f"preenchidas")
            for coord, value in filled:
                text = " ".join(str(value).split())
                if len(text) > per_cell:
                    text = text[:per_cell] + "…"
                lines.append(f"{coord} = {text}")
    finally:
        wb.close()
    return "\n".join(lines)


def assignment(unit: str = UNIT, county: str = COUNTY) -> str:
    return f"""ORDEM DE TRABALHO — {unit}, {county}

Produza o Due Diligence Report (DDR) desta unidade.

O entregavel e um workbook com as abas abaixo. Ele nasce como COPIA do template
da casa, e o template vem preenchido com dados de EXEMPLO de outra unidade
(A.E. Carroll). Tudo que voce nao sobrescrever com write_cell nem apagar com
record_blank sai no arquivo entregue exatamente como esta. Dados de outra
unidade num relatorio entregue sao um defeito grave — o que fazer com eles e
decisao sua.

Mapa do template, celula a celula:
{template_map()}

A base de conhecimento ja tem fichas herdadas da base escrita a mao — ontologia,
taxonomia de documentos, vocabulario de motivos, estrutura do template. Elas sao
ponto de partida, nao lei: se um documento contradisser uma ficha, o documento
manda. `read_knowledge` com um nome que nao existe devolve a lista completa de
fichas, entao uma consulta qualquer te mostra o que ha.

Comece pelo que quiser. Quando terminar, pare de pedir ferramentas e escreva um
resumo do que fez, do que ficou em branco e do que perguntou.
"""


def run(*, max_turns: int = 400, max_cost_usd: float | None = 120.0,
        out_dir: str | Path = OUT_DIR, model: str = MODEL) -> RoundResult:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)

    corpus = Corpus(FILES_DIR, cache=PageCache())
    kb = KnowledgeBase()
    ledger = Ledger()
    dispatch = build_dispatch(corpus, kb, ledger)

    # A run of several hours that only writes at the end is a run that loses
    # everything when it does not reach the end. Each turn is flushed.
    progress = out / "progress.jsonl"
    progress.write_text("", encoding="utf-8")

    def flush(record: Any) -> None:
        with progress.open("a", encoding="utf-8") as fh:
            fh.write(json.dumps({"turno": record.turn,
                                 "chamadas": [c["tool"] for c in record.calls],
                                 "uso": record.usage,
                                 "texto": record.text[:300]},
                                ensure_ascii=False) + "\n")
        ledger.save(out / "ledger.json")

    task = assignment()
    with quarantined() as attempts:
        log, _ = run_loop(system=SYSTEM, task=task, tools=TOOLS,
                          dispatch=dispatch, model=model, max_turns=max_turns,
                          max_cost_usd=max_cost_usd, on_turn=flush)
    if attempts:                        # pragma: no cover — the guard's whole point
        raise AssertionError(f"o run abriu o gabarito: {attempts}")

    workbook = write_workbook(ledger, out / f"DDR - {UNIT.title()} (IA).xlsx")
    ledger_path = ledger.save(out / "ledger.json")
    log_path = log.save(out / "agent_log.json")
    (out / "cache_stats.json").write_text(
        json.dumps({"paginas": corpus.cache.stats(), "conhecimento": kb.stats()},
                   indent=2, ensure_ascii=False), encoding="utf-8")

    return RoundResult(workbook=workbook, ledger=ledger_path, log=log_path,
                       stop_reason=log.stop_reason,
                       counts={**ledger.counts(), **log.usage.to_dict(model)})


def write_workbook(ledger: Ledger, target: str | Path) -> Path:
    """The template, plus what the Examiner wrote, minus what it blanked.

    Nothing else happens here. A cell the model never addressed keeps whatever
    the template had — including another unit's sample data. That is not a
    tolerated bug, it is the measurement: the old pipeline had a gate for it,
    and this round asks whether judgement covers what the gate covered."""
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(TEMPLATE, target)

    wb = openpyxl.load_workbook(target)
    try:
        for (tab, coord), cell in ledger.final_cells().items():
            if tab not in wb.sheetnames:
                wb.create_sheet(tab)
            wb[tab][coord] = cell.valor
        for blank in ledger.blanks:
            if blank.aba in wb.sheetnames:
                wb[blank.aba][blank.celula].value = None
        # The trail ships INSIDE the artefact: a workbook whose sources live in
        # a side file is not auditable alone, and it is the workbook that goes
        # to the client.
        trail = wb.create_sheet("DD_Audit_Trail")
        trail.append(["ABA", "CELULA", "VALOR", "DOCUMENTO", "PAGINA",
                      "TRECHO", "CONFIANCA", "MOTIVO SE BRANCO", "DONO"])
        for cell in ledger.cells:
            trail.append([cell.aba, cell.celula, str(cell.valor),
                          cell.citacao.documento, str(cell.citacao.pagina),
                          cell.citacao.trecho, cell.confianca, "", ""])
        for blank in ledger.blanks:
            trail.append([blank.aba, blank.celula, "", "", "", "", "",
                          blank.motivo, blank.dono])
        wb.save(target)
    finally:
        wb.close()
    return target


if __name__ == "__main__":                          # pragma: no cover
    result = run()
    print(json.dumps(result.to_dict(), indent=2, ensure_ascii=False))
