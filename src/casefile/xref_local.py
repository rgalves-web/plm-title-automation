"""Indexador LOCAL de workbooks de referência — linhas por chave, sem LLM.

Condição 2 da aprovação do plano v1 (02/09/2026): o XREF é processado
integralmente por indexação local; a IA interpreta APENAS as linhas
recuperadas. Este módulo é a metade determinística disso: streaming das abas
inteiras (openpyxl read-only), resolução de colunas POR CABEÇALHO com falha
ruidosa (doutrina de src/dd/tabular_source.py — 4 regressões reais nasceram
de leitura posicional), e match por chaves declaradas. Cada chave pesquisada
vira uma linha em `xref_lookups` (achou/não achou é registro, nunca
inferência); cada linha que casou vai INTEIRA para `xref_rows`.

Cabeçalhos duplicados (o XREF repete ACTIVE FLAG/RECORD TYPE/... nos lados
lease e contrato) são desambiguados por ordem: a 2ª ocorrência vira
`NOME__2`. Nenhuma linha daqui decide pertencimento — quem casa é a chave
que o operador mandou buscar.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Iterator


class ErroDeCabecalho(ValueError):
    """O workbook não tem as colunas que a indexação exige — falha ruidosa."""


@dataclass
class ChaveDeBusca:
    kind: str                    # contract | unit_name | lease_no | recording
    value: str                   # como registrado no lookup
    casa: Callable[[dict[str, str]], bool]


def desambiguar_cabecalhos(header: tuple[Any, ...]) -> list[str]:
    vistos: dict[str, int] = {}
    out: list[str] = []
    for h in header:
        nome = str(h or "").strip()
        vistos[nome] = vistos.get(nome, 0) + 1
        out.append(nome if vistos[nome] == 1 else f"{nome}__{vistos[nome]}")
    return out


def exigir_colunas(cabecalhos: list[str], exigidas: list[str],
                   arquivo: str) -> None:
    faltam = [c for c in exigidas if c not in cabecalhos]
    if faltam:
        raise ErroDeCabecalho(
            f"{arquivo}: colunas exigidas ausentes {faltam}. "
            f"Cabecalhos encontrados: {cabecalhos}")


def normalizar_lease(texto: str) -> str:
    """'0220604/001' e '0220604-001 ' viram a mesma chave."""
    return re.sub(r"[\s/]+", "-", str(texto or "").strip().upper()).replace(
        "--", "-")


def linhas_do_workbook(path: Path) -> Iterator[tuple[str, int, dict[str, str]]]:
    """(aba, nº da linha 1-based, {cabecalho: valor}) — streaming, read-only."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            try:
                header = next(rows)
            except StopIteration:
                continue
            cabecalhos = desambiguar_cabecalhos(header)
            for n, row in enumerate(rows, start=2):
                valores = {cabecalhos[i]: ("" if i >= len(row) or
                                           row[i] is None else str(row[i]))
                           for i in range(len(cabecalhos))}
                yield ws.title, n, valores
    finally:
        wb.close()


def cabecalhos_do_workbook(path: Path) -> dict[str, list[str]]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        out = {}
        for ws in wb.worksheets:
            header = next(ws.iter_rows(values_only=True, max_row=1), ())
            out[ws.title] = desambiguar_cabecalhos(header)
        return out
    finally:
        wb.close()


def indexar(repo: Any, unit: dict[str, Any], path: Path, sha256: str,
            chaves: list[ChaveDeBusca],
            colunas_exigidas: list[str]) -> dict[str, Any]:
    """Varre o workbook inteiro uma vez; persiste lookups e linhas casadas."""
    contagem = {f"{c.kind}:{c.value}": 0 for c in chaves}
    linhas_casadas = 0
    abas: set[str] = set()
    cabecalhos_ok = False
    for aba, numero, valores in linhas_do_workbook(path):
        if aba not in abas:
            abas.add(aba)
            exigir_colunas(list(valores), colunas_exigidas,
                           f"{path.name}!{aba}")
            cabecalhos_ok = True
        casaram = [f"{c.kind}:{c.value}" for c in chaves if c.casa(valores)]
        if casaram:
            for chave in casaram:
                contagem[chave] += 1
            repo.registrar_xref_row(unit["id"], sha256, aba, numero,
                                    valores, casaram)
            linhas_casadas += 1
    if not cabecalhos_ok:
        raise ErroDeCabecalho(f"{path.name}: nenhuma aba com linhas")
    for c in chaves:
        repo.registrar_xref_lookup(unit["id"], sha256, c.kind, c.value,
                                   contagem[f"{c.kind}:{c.value}"])
    return {"arquivo": path.name, "abas": sorted(abas),
            "linhas_casadas": linhas_casadas,
            "chaves": len(chaves),
            "chaves_sem_linha": sorted(k for k, n in contagem.items()
                                       if n == 0)}


# --- construtores de chaves (o QUE buscar veio da ordem do operador) --------

def chave_contract(valor: str, coluna: str = "WELL OR CONTRACT NB"
                   ) -> ChaveDeBusca:
    alvo = valor.strip().upper()
    return ChaveDeBusca(
        "contract", valor,
        lambda v: v.get(coluna, "").strip().upper() == alvo)


def chave_unit_name(valor: str, coluna: str = "WELL OR CONTRACT NAME"
                    ) -> ChaveDeBusca:
    alvo = valor.strip().upper()
    return ChaveDeBusca(
        "unit_name", valor,
        lambda v: alvo in v.get(coluna, "").strip().upper())


def chave_lease_no(valor: str, coluna: str = "LEASE NO") -> ChaveDeBusca:
    alvo = normalizar_lease(valor)
    return ChaveDeBusca(
        "lease_no", valor,
        lambda v: normalizar_lease(v.get(coluna, "")) == alvo)


def chave_recording(vol: str, page: str, coluna: str) -> ChaveDeBusca:
    """Vol/página num campo de texto multi-condado: os dois números, nessa
    ordem, separados por pouca coisa ('Vol 531, Pg 840')."""
    padrao = re.compile(
        rf"\b0*{re.escape(str(vol).lstrip('0') or '0')}\D{{0,12}}"
        rf"0*{re.escape(str(page).lstrip('0') or '0')}\b")
    return ChaveDeBusca(
        "recording", f"{vol}/{page}",
        lambda v: bool(padrao.search(v.get(coluna, ""))))
