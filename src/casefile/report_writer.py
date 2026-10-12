"""Projects reviewed cell proposals into an identified template; protects bytes."""
from __future__ import annotations
import io
import json
from pathlib import Path

from src.casefile.storage import checked_artifact,write_bytes
from src.casefile.training import digest


def formula_set(workbook):
    return {(sheet.title,cell.coordinate):cell.value for sheet in workbook for row in sheet for cell in row if cell.data_type=='f'}


def inspected(path,offset=0,limit=100):
    import openpyxl
    workbook=openpyxl.load_workbook(path)
    total=sum(1 for s in workbook for row in s for c in row if c.value is not None)
    cursor=0
    selected={}
    for s in workbook:
        selected[s.title]=[]
        for row in s:
            for c in row:
                if c.value is None: continue
                if max(0,int(offset))<=cursor<max(0,int(offset))+min(max(int(limit),1),100):
                    selected[s.title].append(c)
                cursor+=1
    return {'total_cells':total,'next_offset':offset+limit if offset+limit<total else None,
            'sheets':[{'name':s.title,'rows':s.max_row,'columns':s.max_column,
                      'merges':[str(m) for m in s.merged_cells.ranges],'print_area':str(s.print_area),
                      'cells':[{'cell':c.coordinate,'value':c.value,'style':c.style_id,
                                'wrap':c.alignment.wrap_text,'horizontal':c.alignment.horizontal,
                                'vertical':c.alignment.vertical,'fill':c.fill.fgColor.index,
                                'font':{'name':c.font.name,'size':c.font.sz,'bold':c.font.bold,'italic':c.font.italic},
                                'border':{side:getattr(getattr(c.border,side), 'style', None) for side in ('left','right','top','bottom')}} for c in selected[s.title]]}
                     for s in workbook], 'formula_sha256':digest(sorted((s,c,str(v)) for (s,c),v in formula_set(workbook).items()))}


def validate(repo,task,output):
    case=repo.operation_case(task['unit_id'])
    template_id=case['profile'].get('template_artifact_id')
    if not template_id or not case['profile'].get('lead_example_artifact_id'):
        return ['template_and_lead_example_not_selected']
    if case['profile'].get('sample_data_clearance') not in ('blank_template_confirmed','clear_declared_cells'):
        return ['template_sample_data_clearance_required']
    example=repo.operation_artifact(case['profile']['lead_example_artifact_id'])
    if not example or example['id'] not in {a['id'] for a in repo.operation_artifacts(task['unit_id'])}:
        return ['lead_example_outside_case']
    checked_artifact(example)
    artifact=repo.operation_artifact(template_id)
    if not artifact or artifact['id'] not in {a['id'] for a in repo.operation_artifacts(task['unit_id'])}:
        return ['template_outside_case']
    import re
    label=artifact['metadata'].get('display_name','')+' '+artifact['metadata'].get('original_path','')
    if re.search(r'project[ _-]*log',label,re.I):
        return ['project_log_cannot_be_writing_template']
    path=checked_artifact(artifact)
    if path.suffix.lower()!='.xlsx':
        return ['xlsx_template_required']
    import openpyxl
    workbook=openpyxl.load_workbook(path)
    errors=[]
    locations=set()
    for item in case['profile'].get('sample_cells',[]):
        if item.get('sheet') not in workbook.sheetnames:
            errors.append('sample_cell_unknown_sheet')
        elif workbook[item['sheet']][item['cell_ref']].data_type=='f' or workbook[item['sheet']][item['cell_ref']].__class__.__name__=='MergedCell':
            errors.append('protected_sample_cell')
    for item in output.get('celulas',[])+output.get('brancos',[]):
        key=(item['sheet'],item['cell_ref'])
        if key in locations:
            errors.append('duplicate_cell:'+str(key))
        locations.add(key)
        if item['sheet'] not in workbook.sheetnames:
            errors.append('unknown_template_sheet:'+item['sheet'])
            continue
        cell=workbook[item['sheet']][item['cell_ref']]
        if cell.data_type=='f' or cell.__class__.__name__=='MergedCell':
            errors.append('protected_template_cell:'+str(key))
        if str(item.get('value','')).startswith('='):
            errors.append('formula_injection:'+str(key))
        for cid in item.get('claim_ids',[]):
            claim=repo.claim(cid)
            if not claim or claim['unit_id']!=task['unit_id'] or claim['status']!='ativo':
                errors.append('claim_outside_case_or_superseded:'+str(cid))
                continue
            if claim['materiality']=='alta':
                parents=repo.conn.execute('SELECT t.status FROM task_claims k JOIN tasks t ON t.id=k.task_id WHERE k.claim_id=%s AND t.verificacao=%s',
                                           (cid,'obrigatoria_e_operador')).fetchall()
                if not parents or any(p['status']!='aceita' for p in parents):
                    errors.append('material_claim_without_operator_acceptance:'+str(cid))
            reviews=[v for v in repo.verificacoes_do_claim(cid) if v.get('status','valida')=='valida']
            if not reviews or reviews[-1]['verdict']!='concordancia':
                errors.append('claim_without_independent_review:'+str(cid))
    return errors


def persist(repo,task,output,session_id,**_):
    errors=validate(repo,task,output)
    if errors:
        raise ValueError(';'.join(errors))
    import openpyxl
    from src.casefile.repo import _json
    case=repo.operation_case(task['unit_id'])
    template=repo.operation_artifact(case['profile']['template_artifact_id'])
    workbook=openpyxl.load_workbook(checked_artifact(template))
    before=formula_set(workbook)
    if case['profile']['sample_data_clearance']=='clear_declared_cells':
        for entry in case['profile'].get('sample_cells',[]):
            sample=workbook[entry['sheet']][entry['cell_ref']]
            sample.value=None
            sample.hyperlink=None
            sample.comment=None
    for entry in output['celulas']:
        cell=workbook[entry['sheet']][entry['cell_ref']]
        cell.value=entry['value']
    for entry in output['brancos']:
        cell=workbook[entry['sheet']][entry['cell_ref']]
        cell.value=None
    buffer=io.BytesIO()
    workbook.save(buffer)
    data=buffer.getvalue()
    reread=openpyxl.load_workbook(io.BytesIO(data))
    if formula_set(reread)!=before:
        raise ValueError('template_formulas_changed')
    identity=digest(output)
    path,sha=write_bytes(task['unit_id'],f"trabalho/runsheet_{task['id']}_{identity[:12]}.xlsx",data)
    artifact=repo.operation_add_artifact(task['unit_id'],'working_report',path,sha,len(data),
                                          {'status':'WORKING DRAFT','template_sha256':template['sha256'],
                                           'formula_sha256':digest(sorted((s,c,str(v)) for (s,c),v in before.items())),
                                           'sample_data_clearance':case['profile'].get('sample_data_clearance','pending'),
                                           'provenance':output,'session_id':session_id},task['id'])
    version=1+len(repo.deliverables(task['unit_id'],'surface_runsheet'))
    repo.registrar_deliverable(task['unit_id'],'surface_runsheet',version,str(path),sha,lifecycle='draft')
    deliverable=repo.deliverables(task['unit_id'],'surface_runsheet')[-1]
    for item in output['celulas']:
        repo.conn.execute('INSERT INTO derived_cells(deliverable_version_id,sheet,cell_ref,value,claim_ids) VALUES (%s,%s,%s,%s,%s)',
                           (deliverable['id'],item['sheet'],item['cell_ref'],str(item['value']),item['claim_ids']))
    repo.operation_report_requirements(task['unit_id'],sha)
    return [artifact['id']]
