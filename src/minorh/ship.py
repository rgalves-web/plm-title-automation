"""`minorh ship` — a memória vira DDR, e nada embarca sem auditoria.

Três passos, na ordem:

1. PROJEÇÃO. O Examinador lê a memória inteira e escreve as células com
   `write_cell`, carregando as citações das fichas. Campo branco ou pendente
   vira `record_blank` na célula correspondente — nunca buraco silencioso.
   Quem decide qual campo cai em qual célula é o modelo, com o mapa do
   template diante dele.

2. AUDITORIA DE 100%. Toda célula escrita é conferida por outra cabeça
   (`auditor.py`), que reabre a página citada. Reprovada → volta ao Examinador
   com o motivo; reprovada pela segunda vez → vira branco declarado, e isso
   fica no relatório. Nada embarca reprovado — regra de embarque, aplicada
   pelo código porque é do operador, não juízo de title.

3. KIT. O DDR, o relatório da auditoria célula a célula, o manifesto de
   lacunas (brancos, pendentes, perguntas abertas) e o custo.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import openpyxl

from src.examinador.ledger import Ledger
from src.examinador.loop import run_loop
from src.examinador.run import write_workbook
from src.minorh.auditor import Auditor, Veredito

FERRAMENTAS_PROJECAO = [
    {
        "name": "write_cell",
        "description": (
            "Escreve uma celula do DDR a partir de um campo da memoria. A "
            "citacao vem da ficha — documento, pagina e trecho literais."),
        "input_schema": {"type": "object", "properties": {
            "aba": {"type": "string"}, "celula": {"type": "string"},
            "valor": {"type": "string"}, "documento": {"type": "string"},
            "pagina": {"type": "string"}, "trecho": {"type": "string"},
            "confianca": {"type": "string",
                          "enum": ["alta", "media", "baixa"]}},
            "required": ["aba", "celula", "valor", "documento", "pagina",
                         "trecho", "confianca"]},
    },
    {
        "name": "record_blank",
        "description": (
            "Declara uma celula em branco com motivo e dono — para campos "
            "brancos ou pendentes da memoria, e para limpar sobras do "
            "template. Aceita faixa (ex. A8:X8)."),
        "input_schema": {"type": "object", "properties": {
            "aba": {"type": "string"}, "celula": {"type": "string"},
            "motivo": {"type": "string"}, "dono": {"type": "string"}},
            "required": ["aba", "celula", "motivo", "dono"]},
    },
]

SISTEMA_PROJECAO = """Voce e o landman responsavel pela unidade e vai agora
projetar a MEMORIA (fichas campo a campo, todas com citacao) no workbook do
DDR. Regras:

1. Todo campo ESCRITO da memoria vira write_cell na celula correta do
   template, com a citacao da ficha. Nao invente valor novo: se nao esta na
   memoria, nao existe.
2. Todo campo BRANCO ou PENDENTE vira record_blank na celula correspondente,
   com o motivo e o dono da ficha. Nunca deixe buraco silencioso.
3. Sobras do template (dados de exemplo de outra unidade) sao limpas com
   record_blank, motivo dizendo o que eram.
4. Toda celula escrita sera auditada por outro modelo, que vai reabrir a
   pagina citada. Uma citacao que nao sustenta o valor sera reprovada.

Quando terminar, pare de pedir ferramentas e resuma o que projetou."""

REAUDITORIA = """Voce e o landman responsavel pela unidade. A auditoria
independente REPROVOU as celulas abaixo. Para cada uma: ou corrija com
write_cell (valor e citacao que a pagina de fato sustenta), ou rebaixe a
branco declarado com record_blank (motivo e dono). Corrigir sem fonte real e
pior que rebaixar: a celula sera auditada de novo, e reprovada duas vezes ela
NAO embarca.

Celulas reprovadas:
{reprovadas}

Quando terminar, pare de pedir ferramentas."""


def _mapa_template(template: Path) -> str:
    from src.examinador.run import template_map
    return template_map(template)


def _dispatch_escrita(ledger: Ledger, template: Path) -> dict[str, Any]:
    """write_cell/record_blank com a caneta protegida (B1).

    Célula de fórmula do template → erro nomeado devolvido AO MODELO pelo
    laço; a escrita não vira branco silencioso, o modelo redecide. Faixa de
    record_blank passa — o write_workbook preserva as fórmulas dentro dela."""
    from src.examinador.formulas import erro_celula_de_formula, formula_cells

    protegidas = set(formula_cells(template))

    def _checar(aba: str, celula: str) -> None:
        coord = str(celula).upper()
        if ":" not in coord and (aba, coord) in protegidas:
            raise ValueError(erro_celula_de_formula(aba, coord))

    def _write_cell(aba, celula, valor, documento, pagina, trecho, confianca):
        _checar(aba, celula)
        return ledger.write_cell(aba, celula, valor, documento, pagina,
                                 trecho, confianca)

    def _record_blank(aba, celula, motivo, dono):
        _checar(aba, celula)
        return ledger.record_blank(aba, celula, motivo, dono)

    return {"write_cell": _write_cell, "record_blank": _record_blank}


def _flush_de(pasta: Path, fase: str) -> Any:
    """Escreve cada turno em `_minorh/ship/progresso_<fase>.jsonl` na hora —
    uma sessão que cai no meio (timeout, limite) não pode sumir com a
    contabilidade do que já foi pago (aconteceu 2x em 24/08)."""
    import json as _json
    caminho = pasta / "_minorh" / "ship" / f"progresso_{fase}.jsonl"
    caminho.parent.mkdir(parents=True, exist_ok=True)

    def flush(rec: Any) -> None:
        with caminho.open("a", encoding="utf-8") as fh:
            fh.write(_json.dumps(
                {"turno": rec.turn, "chamadas": [c["tool"] for c in rec.calls],
                 "uso": rec.usage, "texto": rec.text[:200]},
                ensure_ascii=False) + "\n")
    return flush


def projetar(*, pasta: Path, memoria: Any, template: Path, modelo: str,
             dispatch_leitura: dict[str, Any], client: Any = None,
             max_cost_usd: float = 40.0) -> tuple[Ledger, Any]:
    """Sessão 1: a memória no template. Devolve o ledger e o log do laço."""
    ledger = Ledger()

    dispatch = dict(dispatch_leitura)
    dispatch.update(_dispatch_escrita(ledger, template))

    ferramentas_leitura = [t for t in _leitura_schemas()
                           if t["name"] in dispatch_leitura]
    task = (f"MAPA DO TEMPLATE, celula a celula:\n{_mapa_template(template)}"
            f"\n\nMEMORIA DA UNIDADE (o que voce provou, campo a campo):\n"
            f"{memoria.dump_markdown()}\n\nProjete a memoria no template.")
    log, _ = run_loop(system=SISTEMA_PROJECAO, task=task,
                      tools=FERRAMENTAS_PROJECAO + ferramentas_leitura,
                      dispatch=dispatch, model=modelo, max_turns=200,
                      max_cost_usd=max_cost_usd, client=client,
                      on_turn=_flush_de(pasta, "projecao"))
    return ledger, log


def _leitura_schemas() -> list[dict[str, Any]]:
    from src.minorh.ferramentas import FERRAMENTAS
    return [t for t in FERRAMENTAS
            if t["name"] in ("read_page", "read_knowledge", "list_documents",
                             "search_corpus")]


def _pagina_para_auditoria(corpus: Any, documento: str, pagina: Any
                           ) -> tuple[str, str]:
    """A página citada, no modo mais barato que a torna legível.

    Isto é busca de evidência para o auditor, não julgamento: texto, depois
    OCR, depois um aviso honesto de que só a imagem existe."""
    try:
        num: int | str = int(str(pagina).strip() or 1)
    except ValueError:
        # Uma citação de planilha vem como "aba:janela" — vai crua, o corpus
        # sabe abrir. Qualquer outra coisa ilegível vira página 1 declarada.
        num = str(pagina) if ":" in str(pagina) else 1
    for modo in ("texto", "ocr"):
        try:
            got = corpus.read_page(documento, num, modo)
            if len(got.get("text") or "") >= 40:
                return got["text"], modo
        except Exception as exc:
            return (f"[pagina nao pode ser aberta: "
                    f"{exc.__class__.__name__}: {exc}]", "erro")
    return ("[pagina sem texto legivel por texto/ocr — evidencia so em "
            "imagem]", "ilegivel")


def auditar_tudo(*, unidade: str, workbook: Path, ledger: Ledger,
                 corpus: Any, auditor: Auditor,
                 apenas: set[str] | None = None
                 ) -> tuple[list[Veredito], dict[str, int]]:
    """Toda célula que a projeção escreveu E que segue não vazia no arquivo —
    sem amostragem. Uma escrita depois rebaixada a branco não é auditada: ela
    não embarca.

    `apenas` restringe a uma re-auditoria (as reprovadas corrigidas e o que a
    correção escreveu de novo). Células não vazias sem trilha são estrutura
    herdada do template (cabeçalhos e rótulos); contadas e DECLARADAS no
    relatório — não há citação para conferir."""
    finais = ledger.final_cells()

    wb = openpyxl.load_workbook(workbook, data_only=True)
    try:
        nao_vazias = sum(
            1 for ws in wb.worksheets if ws.title != "DD_Audit_Trail"
            for row in ws.iter_rows() for cell in row
            if cell.value not in (None, ""))
        vivas = {
            (aba, coord): cell for (aba, coord), cell in finais.items()
            if aba in wb.sheetnames
            and wb[aba][coord].value not in (None, "")}
    finally:
        wb.close()

    listagem = "\n".join(
        f"  {aba}!{coord} = {str(c.valor)[:60]!r}"
        for (aba, coord), c in sorted(vivas.items()))

    vereditos: list[Veredito] = []
    for (aba, coord), cell in sorted(vivas.items()):
        if apenas is not None and f"{aba}!{coord}" not in apenas:
            continue
        conteudo, modo = _pagina_para_auditoria(
            corpus, cell.citacao.documento, cell.citacao.pagina)
        vereditos.append(auditor.auditar(
            unidade=unidade, celula=f"{aba}!{coord}", valor=str(cell.valor),
            documento=cell.citacao.documento, pagina=cell.citacao.pagina,
            trecho=cell.citacao.trecho, pagina_conteudo=conteudo[:6000],
            modo=modo, listagem=listagem))

    estrutura = nao_vazias - len(vivas)
    return vereditos, {"nao_vazias": nao_vazias, "auditadas": len(vereditos),
                       "vivas": len(vivas),
                       "estrutura_do_template": max(0, estrutura)}


def reprovadas_de(vereditos: list[Veredito]) -> list[Veredito]:
    return [v for v in vereditos if not v.aprovada]


def corrigir(*, pasta: Path, ledger: Ledger, reprovadas: list[Veredito],
             modelo: str, dispatch_leitura: dict[str, Any], template: Path,
             client: Any = None, max_cost_usd: float = 20.0) -> Any:
    """Sessão 2: as reprovadas voltam ao Examinador, com o motivo."""
    dispatch = dict(dispatch_leitura)
    dispatch.update(_dispatch_escrita(ledger, template))
    texto = "\n".join(
        f"- {v.celula} = {v.valor!r} (citou {v.citacao}) — REPROVADA: "
        f"{v.motivo}" for v in reprovadas)
    ferramentas_leitura = [t for t in _leitura_schemas()
                           if t["name"] in dispatch_leitura]
    log, _ = run_loop(
        system=SISTEMA_PROJECAO, task=REAUDITORIA.format(reprovadas=texto),
        tools=FERRAMENTAS_PROJECAO + ferramentas_leitura, dispatch=dispatch,
        model=modelo, max_turns=60, max_cost_usd=max_cost_usd, client=client,
        on_turn=_flush_de(pasta, "correcao"))
    return log


def aplicar_regra_de_embarque(ledger: Ledger,
                              reprovadas_2x: list[Veredito]) -> list[str]:
    """Nada embarca reprovado. Célula reprovada duas vezes vira branco
    declarado — decisão de EMBARQUE do operador, executada pelo código e
    reportada, nunca escondida."""
    removidas = []
    for v in reprovadas_2x:
        aba, coord = v.celula.split("!", 1)
        ledger.record_blank(
            aba, coord,
            f"reprovada 2x na auditoria: {v.motivo}",
            "Examinador na proxima rodada, com documento melhor")
        removidas.append(v.celula)
    return removidas


def escrever_kit(*, pasta: Path, ledger: Ledger, template: Path,
                 vereditos: list[Veredito], contagem: dict[str, int],
                 auditor: Auditor, memoria: Any, estado: Any,
                 rodadas_auditoria: int,
                 historico_rodada_1: list[Veredito] | None = None
                 ) -> dict[str, Path]:
    from src.examinador.formulas import formulas_hash
    from src.knowledge.write_roots import record_write

    ship_dir = pasta / "_minorh" / "ship"
    ship_dir.mkdir(parents=True, exist_ok=True)

    ddr = write_workbook(ledger, ship_dir / "DDR.xlsx", template=template)

    aprovadas = [v for v in vereditos if v.aprovada]
    reprovadas = reprovadas_de(vereditos)
    hash_template = formulas_hash(template)
    hash_ddr = formulas_hash(ddr)
    relatorio = {
        "auditor": {"provedor": auditor.provedor, "modelo": auditor.modelo,
                    "limitacao": auditor.limitacao or None,
                    "tokens": auditor.uso.to_dict(auditor.modelo)},
        # B2: a prova de que a caneta não tocou fórmula — write_workbook já
        # aborta na divergência; aqui o hash entra no relatório do kit.
        "formulas": {"hash_template": hash_template, "hash_ddr": hash_ddr,
                     "identico": hash_ddr == hash_template},
        "cobertura": {**contagem, "amostragem": "nenhuma — 100% das escritas",
                      "rodadas": rodadas_auditoria},
        "resultado": {"aprovadas": len(aprovadas),
                      "reprovadas_ao_final": len(reprovadas)},
        # O que a rodada 1 reprovou, ANTES da correção. O merge por célula
        # substitui esses vereditos pelos da reauditoria; sem este histórico, o
        # relatório não mostra que o erro foi pego — só que foi consertado.
        "reprovadas_na_rodada_1": [v.to_dict()
                                   for v in (historico_rodada_1 or [])],
        "celulas": [v.to_dict() for v in vereditos],
    }
    rel_path = Path(record_write(ship_dir / "AUDITORIA.json"))
    rel_path.write_text(json.dumps(relatorio, indent=2, ensure_ascii=False),
                        encoding="utf-8")

    contagens = memoria.contar()
    linhas = [
        f"# MANIFESTO DE LACUNAS — {pasta.name}", "",
        "O que este DDR declaradamente NAO contem, e por que. Um kit sem",
        "manifesto esconde os buracos; este os numera.", "",
        f"- campos escritos e auditados: {relatorio['resultado']['aprovadas']}",
        f"- brancos declarados na memoria: {contagens['brancos']}",
        f"- pendentes de documento: {contagens['pendentes']}",
        f"- perguntas abertas ao operador: {len(estado.perguntas_abertas())}",
        f"- pedidos abertos: "
        f"{sum(1 for p in estado.pedidos if p['status'] == 'aberto')}", "",
        "## Brancos e pendentes, campo a campo", ""]
    for ent, campos in sorted(memoria.entidades.items()):
        for nome, campo in sorted(campos.items()):
            if campo.estado == "branco":
                linhas.append(f"- {ent} :: {nome} — BRANCO: {campo.motivo} "
                              f"(dono: {campo.dono})")
            elif campo.estado == "pendente":
                linhas.append(f"- {ent} :: {nome} — PENDENTE do pedido "
                              f"{campo.pedido}")
    if auditor.limitacao:
        linhas += ["", "## Limitacao da auditoria", "", auditor.limitacao]
    man_path = Path(record_write(ship_dir / "MANIFESTO_LACUNAS.md"))
    man_path.write_text("\n".join(linhas), encoding="utf-8")

    return {"ddr": ddr, "auditoria": rel_path, "manifesto": man_path}
