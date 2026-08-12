"""The clean-output gate, over bytes re-read from disk.

CLAUDE.md's own trap list says it plainly: *"Template / sample-row data has
leaked into deliverables before. A clean-output check is required on every
write."* The pipeline had this gate — `PLACEHOLDER_MARKERS`, `_drop_delete_tab`,
`_scrub_placeholders`, `check_sheet_set`. A DDR written by loading the template
and typing over one tab ships the template's sample unit in every other tab, its
`delete - …` worksheet, and no audit trail. That is what happened here: 247
contaminated cells, the wrong sheet set, no `DD_Audit_Trail`.

So this module asserts, and it asserts against the FILE, not against the
in-memory workbook that produced it. A gate that trusts the object it just built
is testing its own intentions.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

# The template's sample unit, as it actually appears in the delivered template.
# Every marker here was found in a generated file, not imagined.
PLACEHOLDER_MARKERS: tuple[str, ...] = (
    "***", "Reeves", "Forest Oil", "A.E. Carroll", "636.13", "Section **",
)
DELETE_TAB_PREFIX = "delete - "
AUDIT_TRAIL_SHEET = "DD_Audit_Trail"

_MARKER_RE = re.compile("|".join(re.escape(m) for m in PLACEHOLDER_MARKERS),
                        re.IGNORECASE)


@dataclass
class CleanFinding:
    severity: str            # FAIL
    kind: str
    where: str
    detail: str

    def __str__(self) -> str:
        return f"{self.severity} {self.kind} @ {self.where}: {self.detail}"


@dataclass
class CleanReport:
    findings: list[CleanFinding] = field(default_factory=list)
    sheets: list[str] = field(default_factory=list)
    cells_checked: int = 0
    orphans: int = 0

    @property
    def failed(self) -> bool:
        return any(f.severity == "FAIL" for f in self.findings)

    def summary(self) -> str:
        kinds: dict[str, int] = {}
        for f in self.findings:
            kinds[f.kind] = kinds.get(f.kind, 0) + 1
        return (f"{'FAIL' if self.failed else 'CLEAN'} · "
                f"{self.cells_checked} cells · " + (", ".join(
                    f"{k}={v}" for k, v in sorted(kinds.items())) or "no findings"))


def check_clean_output(path: str | Path, *, expected_sheets: set[str] | None = None,
                       require_audit_trail: bool = True,
                       max_report: int = 40) -> CleanReport:
    """Re-read the saved workbook and prove it carries nothing it should not."""
    path = Path(path)
    report = CleanReport()
    if not path.exists():
        report.findings.append(CleanFinding(
            "FAIL", "missing_file", str(path), "the deliverable is not on disk"))
        return report

    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        report.sheets = list(wb.sheetnames)

        for name in wb.sheetnames:
            if name.lower().startswith(DELETE_TAB_PREFIX):
                report.findings.append(CleanFinding(
                    "FAIL", "delete_tab", name,
                    "a worksheet the template marks for deletion shipped"))

        if expected_sheets is not None:
            actual = set(wb.sheetnames)
            for extra in sorted(actual - expected_sheets):
                report.findings.append(CleanFinding(
                    "FAIL", "sheet_set", extra, "sheet not in the expected set"))
            for missing in sorted(expected_sheets - actual):
                report.findings.append(CleanFinding(
                    "FAIL", "sheet_set", missing, "expected sheet is absent"))

        if require_audit_trail and AUDIT_TRAIL_SHEET not in wb.sheetnames:
            report.findings.append(CleanFinding(
                "FAIL", "no_audit_trail", AUDIT_TRAIL_SHEET,
                "no audit trail in the artefact — the evidence exists only in "
                "the case file, so the deliverable is not traceable on its own"))

        seen = 0
        for name in wb.sheetnames:
            for row in wb[name].iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    report.cells_checked += 1
                    text = str(cell.value)
                    m = _MARKER_RE.search(text)
                    if m and seen < max_report:
                        seen += 1
                        report.findings.append(CleanFinding(
                            "FAIL", "placeholder", f"{name}!{cell.coordinate}",
                            f"{m.group(0)!r} — template sample data: "
                            f"{' '.join(text.split())[:70]}"))
                    elif m:
                        seen += 1
        if seen > max_report:
            report.findings.append(CleanFinding(
                "FAIL", "placeholder", "(more)",
                f"{seen - max_report} further contaminated cells not listed"))
    finally:
        wb.close()
    return report


class NotCleanError(AssertionError):
    """The deliverable carries template data, a delete tab, or no audit trail."""


def assert_clean(path: str | Path, **kwargs: Any) -> CleanReport:
    report = check_clean_output(path, **kwargs)
    if report.failed:
        lines = [str(f) for f in report.findings[:20]]
        raise NotCleanError(
            f"{Path(path).name} is not clean ({len(report.findings)} finding(s)):\n  "
            + "\n  ".join(lines))
    return report


# ---------------------------------------------------------------------------
# The checklist asserts. Every "Yes" and every number in Notes is DERIVED.
# ---------------------------------------------------------------------------
#
# A checklist that says "32 documents uploaded" beside a folder holding one is
# not an optimistic estimate — it is a false statement to the client, made by a
# machine, in a file with our name on it. So the numbers are recomputed from the
# manifest and compared, and a mismatch is a FAIL rather than a note.

_NUMBER_RE = re.compile(r"\b(\d[\d,]*(?:\.\d+)?)\s*%?")


def _numbers_in(text: str) -> list[str]:
    return [m.group(1).replace(",", "") for m in _NUMBER_RE.finditer(str(text or ""))]


def check_checklist(path: str | Path, *, placed: dict[str, list[str]],
                    row_to_folder: dict[int, str | None],
                    derived_numbers: dict[int, set[str]] | None = None
                    ) -> CleanReport:
    """Re-read the checklist and prove each Yes and each number is earned.

    `derived_numbers` maps a row to every number the round actually computed for
    it. A number in Notes that is not in that set was typed, not derived."""
    path = Path(path)
    report = CleanReport()
    if not path.exists():
        report.findings.append(CleanFinding("FAIL", "missing_file", str(path),
                                            "the checklist is not on disk"))
        return report

    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        report.sheets = list(wb.sheetnames)
        for row, folder in row_to_folder.items():
            answer = str(ws.cell(row=row, column=4).value or "").strip().lower()
            notes = str(ws.cell(row=row, column=5).value or "")
            report.cells_checked += 2

            if folder is not None:
                held = len(placed.get(folder, []))
                if answer == "yes" and held == 0:
                    report.findings.append(CleanFinding(
                        "FAIL", "unearned_yes", f"row {row} ({folder})",
                        "answered Yes with nothing in the folder"))
                if answer == "no" and not notes.strip():
                    report.findings.append(CleanFinding(
                        "FAIL", "undeclared_no", f"row {row} ({folder})",
                        "answered No with no pendência in Notes"))

            allowed = (derived_numbers or {}).get(row, set())
            for number in _numbers_in(notes):
                if number not in allowed:
                    report.findings.append(CleanFinding(
                        "FAIL", "undeclared_number", f"row {row}",
                        f"{number!r} in Notes is not a figure this round "
                        f"computed (derived: {sorted(allowed) or 'none'})"))
    finally:
        wb.close()
    return report
