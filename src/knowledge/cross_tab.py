"""Cross-tab consistency, and coverage at BOTH ends.

Two checks the whitelist gate could not make, for the same underlying reason:
it walks the cells that were written and asks whether a source was cited. It
cannot ask whether the right FIELD was read from that source, and it cannot ask
about a cell nobody wrote.

  CONSISTENCY. One identity — a lease, a well — appearing in two tabs must
  carry the same value in both. The PR-10 artefact said `Vol. 567, Pg. 170` for
  lease 0220604-024 in the Lease Review and `Volume 199, Page 356` for the same
  lease in the Runsheet, both fully trailed. This catches that mechanically,
  without anyone opening the workbook.

  COVERAGE AT BOTH ENDS. Invariant 24 says a blank owes a reason, but a gate
  that walks written cells never sees the blank. `Ownership - Worksheet` A5:E5 —
  the unit recap under the headers the template declares, the headline numbers
  of the whole DDR — was empty AND absent from the trail. The check had been
  shaped by what the error looked like instead of by what correctness requires.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

from src.knowledge.clean_output import AUDIT_TRAIL_SHEET, CleanFinding, CleanReport

LEASE_RE = re.compile(r"0\d{6}[-/]\d{3}")


def _norm_recording(value: Any) -> str:
    digits = re.findall(r"\d+", str(value or ""))
    return ".".join(digits[:2]) if len(digits) >= 2 else ""


def _norm_text(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


@dataclass
class Claim:
    """What one tab says about one identity."""
    sheet: str
    cell: str
    identity: str
    field: str
    value: Any

    def key(self) -> tuple[str, str]:
        return (self.identity, self.field)


def _trail_index(wb) -> dict[tuple[str, str], tuple]:
    out: dict[tuple[str, str], tuple] = {}
    if AUDIT_TRAIL_SHEET in wb.sheetnames:
        for row in wb[AUDIT_TRAIL_SHEET].iter_rows(min_row=2, values_only=True):
            if row and row[0] and row[1]:
                out[(str(row[0]), str(row[1]))] = row
    return out


def _lease_from_trail(trail: dict, sheet: str, cell: str) -> str:
    """The lease a row belongs to, taken from its own citation.

    The Runsheet tab does not print a lease number, so a cross-tab check that
    only reads the visible grid can never join it to the Lease Review — which
    is why the `Volume 199, Page 356` contradiction survived a gate that was
    otherwise looking straight at it. The trail knows: the citation names the
    source file, and the corpus path names the lease."""
    row = trail.get((sheet, cell))
    if not row or len(row) < 5:
        return ""
    m = LEASE_RE.search(str(row[4] or ""))
    return m.group(0).replace("/", "-") if m else ""


# How each tab is read for cross-tab claims. Adding a tab means adding it here,
# which is the point: a tab nobody declared is a tab nobody cross-checks.
def collect_claims(wb) -> list[Claim]:
    claims: list[Claim] = []
    trail = _trail_index(wb)

    if "Lease Review" in wb.sheetnames:
        ws = wb["Lease Review"]
        for row in ws.iter_rows(min_row=2):
            ident = LEASE_RE.search(str(row[0].value or "")) if row else None
            if not ident:
                continue
            lease = ident.group(0).replace("/", "-")
            for idx, field_name in ((2, "lessor"), (5, "recording")):
                if idx < len(row) and row[idx].value is not None:
                    claims.append(Claim("Lease Review", row[idx].coordinate,
                                        lease, field_name, row[idx].value))

    if "Ownership - Worksheet" in wb.sheetnames:
        ws = wb["Ownership - Worksheet"]
        for row in ws.iter_rows(min_row=8):
            ident = LEASE_RE.search(str(row[1].value or "")) if len(row) > 1 else None
            if not ident:
                continue
            lease = ident.group(0).replace("/", "-")
            book = row[10].value if len(row) > 10 else None
            page = row[11].value if len(row) > 11 else None
            if book and page:
                claims.append(Claim("Ownership - Worksheet",
                                    f"K{row[0].row}/L{row[0].row}", lease,
                                    "recording", f"{book}.{page}"))
            if len(row) > 4 and row[4].value is not None:
                claims.append(Claim("Ownership - Worksheet", row[4].coordinate,
                                    lease, "lessor", row[4].value))

    if "Runsheet" in wb.sheetnames:
        ws = wb["Runsheet"]
        for row in ws.iter_rows(min_row=4):
            if len(row) < 8 or row[0].value is None:
                continue
            lease = _lease_from_trail(trail, "Runsheet", row[2].coordinate)
            if lease and row[2].value is not None:
                claims.append(Claim("Runsheet", row[2].coordinate, lease,
                                    "recording", row[2].value))
            if lease and len(row) > 6 and row[6].value is not None:
                claims.append(Claim("Runsheet", row[6].coordinate, lease,
                                    "lessor", row[6].value))
    return claims


def check_cross_tab(path: str | Path,
                    lease_recordings: dict[str, str] | None = None) -> CleanReport:
    """Same identity, same value — in every tab that mentions it."""
    report = CleanReport()
    wb = openpyxl.load_workbook(Path(path), data_only=True)
    try:
        claims = collect_claims(wb)
        report.cells_checked = len(claims)
        by_key: dict[tuple[str, str], list[Claim]] = {}
        for claim in claims:
            by_key.setdefault(claim.key(), []).append(claim)

        for (identity, field_name), group in sorted(by_key.items()):
            norm = (_norm_recording if field_name == "recording" else _norm_text)
            values = {norm(c.value) for c in group if c.value is not None}
            values.discard("")
            if len(values) > 1:
                where = " vs ".join(f"{c.sheet}!{c.cell}={str(c.value)[:32]!r}"
                                    for c in group)
                report.findings.append(CleanFinding(
                    "FAIL", "cross_tab", f"{identity} / {field_name}",
                    f"tabs disagree — {where}"))

        # The authority has the last word: a tab that contradicts Exhibit "A"
        # is wrong even when every tab agrees with every other.
        for lease, expected in (lease_recordings or {}).items():
            want = _norm_recording(expected)
            if not want:
                continue
            for claim in by_key.get((lease, "recording"), []):
                if _norm_recording(claim.value) != want:
                    report.findings.append(CleanFinding(
                        "FAIL", "contradicts_authority",
                        f"{claim.sheet}!{claim.cell}",
                        f"{lease}: tab says {str(claim.value)[:32]!r}, "
                        f'Exhibit "A" says {expected!r}'))
    finally:
        wb.close()
    return report


# ---------------------------------------------------------------------------
# Coverage at both ends
# ---------------------------------------------------------------------------

@dataclass
class DataRegion:
    """Cells a tab OWES an answer for — written with a trail, or blank with a
    reason. Declared per sheet, because "the region nobody declared" is exactly
    the region that ships empty and unnoticed."""
    sheet: str
    cells: list[str] = field(default_factory=list)
    why: str = ""


# The unit recap: the headline numbers of the whole DDR, under the template's
# own headers at A4:E4. Empty and untrailed in the PR-10 artefact.
REQUIRED_REGIONS: tuple[DataRegion, ...] = (
    DataRegion("Ownership - Worksheet", ["A5", "B5", "C5", "D5", "E5"],
               "unit recap under the declared headers (Unit Gross · Company Net "
               "Acres · Company Net Revenue Acres · Seller WI · Seller NRI)"),
)


def check_coverage(path: str | Path,
                   regions: tuple[DataRegion, ...] = REQUIRED_REGIONS
                   ) -> CleanReport:
    """Every declared data cell is written-with-trail OR blank-with-reason."""
    report = CleanReport()
    wb = openpyxl.load_workbook(Path(path), data_only=True)
    try:
        trail: dict[tuple[str, str], tuple] = {}
        if AUDIT_TRAIL_SHEET in wb.sheetnames:
            for row in wb[AUDIT_TRAIL_SHEET].iter_rows(min_row=2,
                                                       values_only=True):
                if row and row[0] and row[1]:
                    trail[(str(row[0]), str(row[1]))] = row

        for region in regions:
            if region.sheet not in wb.sheetnames:
                continue
            ws = wb[region.sheet]
            for coord in region.cells:
                report.cells_checked += 1
                entry = trail.get((region.sheet, coord))
                has_value = ws[coord].value is not None
                if has_value and entry:
                    continue
                if not has_value and entry and len(entry) > 5 and entry[5]:
                    continue                     # declared blank with a reason
                report.findings.append(CleanFinding(
                    "FAIL", "uncovered_cell", f"{region.sheet}!{coord}",
                    f"neither written-with-trail nor blank-with-reason — "
                    f"{region.why}"))
    finally:
        wb.close()
    return report


class CrossTabError(AssertionError):
    """Tabs disagree, or a declared data cell is unaccounted for."""


def assert_consistent(path: str | Path,
                      lease_recordings: dict[str, str] | None = None,
                      regions: tuple[DataRegion, ...] = REQUIRED_REGIONS) -> None:
    problems: list[CleanFinding] = []
    problems += check_cross_tab(path, lease_recordings).findings
    problems += check_coverage(path, regions).findings
    if problems:
        raise CrossTabError(
            f"{Path(path).name}: {len(problems)} consistency problem(s):\n  "
            + "\n  ".join(str(f) for f in problems[:20]))
