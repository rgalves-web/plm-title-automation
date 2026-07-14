from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def generate_sample_dataset(base_dir: str | Path = "data/input/sample") -> dict[str, Path]:
    base_dir = Path(base_dir)
    base_dir.mkdir(parents=True, exist_ok=True)

    document_list_path = base_dir / "sample_document_list.xlsx"
    template_path = base_dir / "sample_runsheet_template.xlsx"

    dataframe = pd.DataFrame(
        [
            {
                "Document Type": "Patent",
                "Document Title": "United States Patent",
                "Instrument Number": "1001",
                "Instrument Date": "1901-01-02",
                "Recording Date": "1901-02-01",
                "Grantor": "United States of America",
                "Grantee": "John Smith, a married man",
                "Book": "1",
                "Page": "10",
                "Legal Description": "Section 1 Township 4S Range 64W",
                "Comments": "",
            },
            {
                "Document Type": "Oil and Gas Lease",
                "Document Title": "Oil and Gas Lease - 5 year primary term expired",
                "Instrument Number": "2002",
                "Instrument Date": "1990-06-01",
                "Recording Date": "1990-06-12",
                "Grantor": "Jane Smith, a widow",
                "Grantee": "PLM Energy LLC",
                "Book": "25",
                "Page": "300",
                "Legal Description": "Section 1 Township 4S Range 64W",
                "Comments": "Primary term expired",
            },
            {
                "Document Type": "Easement",
                "Document Title": "Pipeline Easement",
                "Instrument Number": "3003",
                "Instrument Date": "2001-05-03",
                "Recording Date": "2001-05-20",
                "Grantor": "John Smith and Jane Smith, husband and wife",
                "Grantee": "Pipeline Co.",
                "Book": "40",
                "Page": "100",
                "Legal Description": "Section 1 Township 4S Range 64W",
                "Comments": "",
            },
            {
                "Document Type": "Patent",
                "Document Title": "United States Patent copy",
                "Instrument Number": "1001",
                "Instrument Date": "1901-01-02",
                "Recording Date": "1901-02-01",
                "Grantor": "United States of America",
                "Grantee": "John Smith, a married man",
                "Book": "1",
                "Page": "10",
                "Legal Description": "Section 1 Township 4S Range 64W",
                "Comments": "duplicate image",
            },
        ]
    )
    dataframe.to_excel(document_list_path, index=False)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Runsheet"
    headers = [
        "NO.",
        "INSTRUMENT",
        "INSTRUMENT NUMBER",
        "EFFECTIVE DATE",
        "INSTRUMENT DATE",
        "RECORDED DATE",
        "GRANTOR",
        "GRANTEE",
        "ACREAGE",
        "CONVEYANCE",
        "COMMENTS",
    ]
    sheet.append(headers)
    sheet.append([""] * len(headers))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for column in range(1, len(headers) + 1):
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = 20
    workbook.save(template_path)

    return {
        "document_list": document_list_path,
        "template": template_path,
    }


if __name__ == "__main__":
    paths = generate_sample_dataset()
    for name, path in paths.items():
        print(f"{name}: {path}")
