"""Source-truth validation: every cell re-checked against the document it cites.

Structural rules verify FORM; these verify TRUTH. The discipline under test is
the three-way split: a value the source contradicts is a MISMATCH, a value the
source simply doesn't yield is UNVERIFIABLE — never a mismatch, because OCR is
lossy and an accusation on a bad scan is worse than an honest "cannot tell".
"""
from __future__ import annotations

from datetime import date
from fractions import Fraction

import openpyxl
import pytest

from src.dd.source_truth_check import (
    CONFIRMED,
    MISMATCH,
    UNVERIFIABLE,
    SourceIndex,
    check_date_against_xref,
    check_nri_arithmetic,
    check_recording_against_dpu,
    excel_serial_to_date,
    parse_loose_date,
    read_audit_cites,
    run_source_truth,
    states_value,
    truth_table,
)


# --- text matching ---------------------------------------------------------

def test_states_value_ignores_ocr_spacing_and_punctuation():
    assert states_value("Volume 405, Page 343, Montague County", "Volume 405 Page 343")
    assert states_value("EOG RESOURCES, INC.", "EOG Resources Inc")
    assert not states_value("Volume 405, Page 343", "Volume 412, Page 512")


# --- computable cross-checks ----------------------------------------------

def test_nri_equals_one_minus_royalty_minus_orri():
    assert check_nri_arithmetic(0.2, 0, 0.8)[0] == CONFIRMED
    assert check_nri_arithmetic(0.2, 0.017592593, 0.782407407)[0] == CONFIRMED


def test_wrong_nri_is_a_mismatch_against_its_own_inputs():
    outcome, expected = check_nri_arithmetic(0.25, 0, 0.9)
    assert outcome == MISMATCH
    assert expected == Fraction(3, 4)


def test_the_templates_own_nri_formula_is_confirmed_not_recomputed():
    # "=1-G8-W8" IS the rule; there is nothing for it to contradict.
    assert check_nri_arithmetic(0.2, 0, "=1-G8-W8")[0] == CONFIRMED


def test_resolver_maps_a_recording_cite_to_the_volpage_named_instrument(tmp_path):
    # Bowen audit: "DPU K1400725 Exhibit A (Vol 479 Pg 590-596)" must reach the
    # instrument filed as "479.590 - Declaration of Pooled Unit - Bowen A Unit.tif".
    from src.dd.source_truth_check import SourceIndex
    dpu = tmp_path / "479.590 - Declaration of Pooled Unit - Bowen A Unit.tif"
    dpu.write_bytes(b"x")
    idx = SourceIndex([tmp_path], ocr_fn=lambda p: "")
    assert idx.resolve("DPU K1400725 Exhibit A (Vol 479 Pg 590-596)") == dpu
    assert idx.resolve("Amendment (Vol 493, Pg 217)") is None      # não inventa


def test_resolver_maps_an_inner_parenthesized_filename(tmp_path):
    from src.dd.source_truth_check import SourceIndex
    inner = tmp_path / "UNIT DEC - REC.TIF"
    inner.write_bytes(b"x")
    idx = SourceIndex([tmp_path], ocr_fn=lambda p: "")
    assert idx.resolve("DPU K1400725 zip (UNIT DEC - REC.TIF)") == inner


def test_resolver_maps_a_workbook_title_to_the_xlsx_with_copy_suffix(tmp_path):
    from src.dd.source_truth_check import SourceIndex
    import openpyxl
    wb = openpyxl.Workbook(); wb.active["A1"] = "RAMA JOE BOWEN GILLASPIA"
    wb.active["B1"] = "0220618-001"
    sched = tmp_path / "Lease & Mineral Schedule 6.9.26 (2).xlsx"
    wb.save(sched)
    idx = SourceIndex([tmp_path], ocr_fn=lambda p: "")
    got = idx.resolve("Lease & Mineral Schedule 6.9.26")
    assert got == sched
    # verificação tabular: as células viram texto pesquisável, sem OCR
    text = idx.text(got)
    assert "GILLASPIA" in text and "0220618-001" in text


def test_a_citation_never_resolves_across_a_different_abstract(tmp_path):
    # The real Bowen failure: the citation "T-O SUPP Bowen A-227 141.5 ac
    # (T0140210), cert. 10/18/2008" resolved to an A-682 opinion. Two causes,
    # both fixed: Path().stem read the date's slashes as path separators and
    # collapsed the citation to "2008" (so any file naming that year matched),
    # and containment then took whichever candidate traversal happened to hit
    # first — making resolution depend on folder layout.
    from src.dd.source_truth_check import SourceIndex
    for name in ("T-O FIRST SUPP BOWEN UNIT A-682 146.5AC (issued 07-11-2008).pdf",
                 "T-O ORIG BOWEN UNIT A-227 141.5AC (cert 02-04-2008).pdf",
                 "T-O SUPP BOWEN UNIT A-227 141.5AC (cert 10-18-2008).pdf",
                 "T-O ORIG W. EDWARDS A-227 125AC (12-31-2008).pdf"):
        (tmp_path / name).write_bytes(b"x")
    idx = SourceIndex([tmp_path], ocr_fn=lambda p: "")
    got = idx.resolve("T-O SUPP Bowen A-227 141.5 ac (T0140210), cert. 10/18/2008")
    assert got is not None and "A-682" not in got.name.upper()
    assert got.name == "T-O SUPP BOWEN UNIT A-227 141.5AC (cert 10-18-2008).pdf"
    # And the A-682 citation still finds its own document.
    a682 = idx.resolve("T-O FIRST SUPP Bowen Unit A-682 146.5 ac (issued 07-11-2008)")
    assert a682 is not None and "A-682" in a682.name.upper()


def test_a_citation_that_leads_with_the_filename_then_adds_prose(tmp_path):
    # Real Bowen citations carry the document name and then commentary. The
    # longest indexed name the citation STARTS WITH is the document it names.
    from src.dd.source_truth_check import SourceIndex
    (tmp_path / "T-O ORIG BOWEN UNIT A-227, 141.5 AC.PDF").write_bytes(b"x")
    (tmp_path / "Lease & Mineral Schedule 6.9.26 (2).xlsx").write_bytes(b"x")
    idx = SourceIndex([tmp_path], ocr_fn=lambda p: "")
    assert idx.resolve("T-O ORIG BOWEN UNIT A-227, 141.5 AC.PDF (in all 0220649 lease "
                       "folders)").name.startswith("T-O ORIG BOWEN UNIT A-227")
    assert idx.resolve("Lease & Mineral Schedule 6.9.26 TRUE GROSS").suffix == ".xlsx"


def test_same_name_different_bytes_is_not_a_copy(tmp_path):
    # "X (2).pdf is the same document as X.pdf" holds only if the BYTES agree.
    # Same name + different content is two documents wearing one name; naming
    # either would be a guess, so resolution refuses AND says why.
    from src.dd.source_truth_check import SourceIndex
    (tmp_path / "T-O ORIG X A-100.PDF").write_bytes(b"same")
    (tmp_path / "T-O ORIG X A-100 (1).PDF").write_bytes(b"same")
    idx = SourceIndex([tmp_path], ocr_fn=lambda p: "")
    assert idx.resolve("T-O ORIG X A-100.PDF p.1") is not None      # cópias reais

    (tmp_path / "T-O ORIG X A-100 (2).PDF").write_bytes(b"DIFFERENT CONTENT")
    idx2 = SourceIndex([tmp_path], ocr_fn=lambda p: "")
    assert idx2.resolve("T-O ORIG X A-100.PDF p.1") is None
    reason = " ".join(idx2.unresolved_reasons.values())
    assert "bytes differ" in reason and "not copies" in reason


def test_byte_copies_of_one_document_are_one_identity(tmp_path):
    # Four copies of the same opinion are not four rival candidates: the
    # citation's identity is unambiguous, only the copy is. Resolve
    # deterministically instead of refusing.
    from src.dd.source_truth_check import SourceIndex
    for suffix in ("", " (1)", " (2)", " (3)"):
        (tmp_path / f"T-O ORIG BOWEN A-1449 119.2 AC{suffix}.PDF").write_bytes(b"x")
    idx = SourceIndex([tmp_path], ocr_fn=lambda p: "")
    got = idx.resolve("T-O ORIG BOWEN A-1449 119.2 AC.PDF p.1 (read visually)")
    assert got is not None and "A-1449" in got.name
    # Deterministic across runs — never traversal-order dependent.
    assert idx.resolve("T-O ORIG BOWEN A-1449 119.2 AC.PDF p.1 (read visually)") == got


def test_an_ambiguous_citation_resolves_to_nothing_rather_than_guessing(tmp_path):
    # Correct-and-fewer: two equally-plausible candidates name no document.
    from src.dd.source_truth_check import SourceIndex
    for name in ("Assignment Package Copy A.pdf", "Assignment Package Copy B.pdf"):
        (tmp_path / name).write_bytes(b"x")
    idx = SourceIndex([tmp_path], ocr_fn=lambda p: "")
    assert idx.resolve("Assignment Package Copy") is None


def test_a_citation_carrying_a_date_is_not_reduced_to_its_year(tmp_path):
    from src.dd.source_truth_check import _citation_stem, _norm_key
    cite = "T-O SUPP Bowen A-227 141.5 ac (T0140210), cert. 10/18/2008"
    assert _norm_key(_citation_stem(cite)) != "2008"
    assert "BOWEN" in _citation_stem(cite).upper()
    assert _citation_stem("UNIT DEC - REC.TIF") == "UNIT DEC - REC"   # extensão sai


def test_prose_orri_makes_the_identity_unverifiable_not_a_mismatch():
    # Bowen V8: ORRI column carries "Burdens reflected in FIRST SUPP T-O NRI
    # 0.51875 (see opinion)" — prose, not a number. Treating it as zero would
    # manufacture a mismatch out of a type error; the identity is simply not
    # computable. Blank ORRI still means zero; numeric ORRI still verifies.
    outcome, expected = check_nri_arithmetic(
        0.125, "Burdens reflected in FIRST SUPP T-O NRI 0.51875 (see opinion)",
        0.51875)
    assert outcome == UNVERIFIABLE and expected is None
    assert check_nri_arithmetic(0.125, "", 0.875)[0] == CONFIRMED   # blank = 0
    assert check_nri_arithmetic(0.125, 0.1, 0.875)[0] == MISMATCH   # numeric still strict


def test_nri_without_a_royalty_is_unverifiable_not_wrong():
    assert check_nri_arithmetic(None, 0, 0.8)[0] == UNVERIFIABLE


_DPU = {"0190122-001": ("402", "50"), "0220344-006": ("405", "343")}


def test_recording_matching_the_dpu_is_confirmed():
    assert check_recording_against_dpu("0190122-001", "402", "50", _DPU)[0] == CONFIRMED


def test_recording_differing_from_the_dpu_is_a_mismatch():
    # The real defect: Hulme's row carried John A Cox's recording (475/39).
    outcome, src = check_recording_against_dpu("0220344-006", "475", "39", _DPU)
    assert outcome == MISMATCH
    assert src == "Vol 405, Pg 343"


def test_lease_absent_from_the_dpu_is_unverifiable():
    assert check_recording_against_dpu("9999999-999", "1", "2", _DPU)[0] == UNVERIFIABLE


def test_excel_serial_and_loose_dates():
    assert excel_serial_to_date(39322) == date(2007, 8, 28)
    assert parse_loose_date("August 28, 2007") == date(2007, 8, 28)
    assert parse_loose_date("December 20, 2007") == date(2007, 12, 20)
    assert parse_loose_date("2 day of January, 2010") == date(2010, 1, 2)
    assert parse_loose_date("12/20/2007") == date(2007, 12, 20)
    assert parse_loose_date("no date here") is None


def test_contract_date_against_the_xref_serial():
    assert check_date_against_xref("August 28, 2007", 39322)[0] == CONFIRMED
    outcome, src = check_date_against_xref("August 29, 2007", 39322)
    assert outcome == MISMATCH and src == "2007-08-28"
    assert check_date_against_xref("unreadable", 39322)[0] == UNVERIFIABLE


# --- source resolution -----------------------------------------------------

def _docs(tmp_path, files: dict[str, str]):
    for name, body in files.items():
        (tmp_path / name).write_text(body, encoding="utf-8")
    return tmp_path


def test_source_index_resolves_a_reference_to_a_file(tmp_path):
    _docs(tmp_path, {"T-O SUPP LYNN W COX ET AL A-122.txt": "text"})
    index = SourceIndex([tmp_path])
    assert index.resolve("T-O SUPP LYNN W COX ET AL A-122") is not None


def test_source_index_ignores_a_page_cite_in_the_reference(tmp_path):
    _docs(tmp_path, {"T-O SUPP LYNN W COX.txt": "text"})
    index = SourceIndex([tmp_path])
    assert index.resolve("T-O SUPP LYNN W COX pp.5-6 (MINERALS)") is not None


def test_non_document_sources_never_resolve(tmp_path):
    index = SourceIndex([tmp_path])
    for reference in ("manifest", "dd_pipeline", "Lease to Well XREF", "formula"):
        assert index.resolve(reference) is None


def test_page_text_splits_on_form_feeds(tmp_path):
    _docs(tmp_path, {"doc.txt": "PAGE ONE\fPAGE TWO\fPAGE THREE"})
    index = SourceIndex([tmp_path])
    path = index.resolve("doc")
    assert "PAGE TWO" in index.page_text(path, 2)
    assert "PAGE ONE" not in index.page_text(path, 2)


# --- end to end ------------------------------------------------------------

def _wb(agreements: dict[str, object], audit: list[list]):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Agreements")
    for coord, value in agreements.items():
        ws[coord] = value
    aud = wb.create_sheet("DD_Audit_Trail")
    aud.append(["Sheet", "Location", "Source Type", "Source Reference", "Confidence", "Extracted Text"])
    for row in audit:
        aud.append(row)
    return wb


def test_a_value_stated_in_the_cited_source_is_confirmed(tmp_path):
    _docs(tmp_path, {"ASMT.txt": "page one\fThis assignment from JOE BOB SMITH to EOG"})
    wb = _wb({"B4": "JOE BOB SMITH"},
             [["Agreements", "B4", "contract_instrument", "ASMT.TIF", 1, "grantor @ p.2"]])
    report = run_source_truth(wb, [tmp_path])
    hit = next(f for f in report.findings if f.cell == "B4")
    assert hit.outcome == CONFIRMED
    assert hit.source_page == "p.2"


def test_a_value_on_a_different_page_is_confirmed_but_noted(tmp_path):
    _docs(tmp_path, {"ASMT.txt": "JOE BOB SMITH appears here\fpage two"})
    wb = _wb({"B4": "JOE BOB SMITH"},
             [["Agreements", "B4", "contract_instrument", "ASMT.TIF", 1, "grantor @ p.2"]])
    report = run_source_truth(wb, [tmp_path])
    hit = next(f for f in report.findings if f.cell == "B4")
    assert hit.outcome == CONFIRMED and "NOT on the cited p.2" in hit.detail
    assert "names the wrong page" in hit.detail   # a bad cite is itself a defect


def test_a_value_absent_from_the_source_is_unverifiable_not_a_mismatch(tmp_path):
    # OCR is lossy; absence is not contradiction. This is the guard against the
    # over-reach that once produced thousands of false failures.
    _docs(tmp_path, {"ASMT.txt": "wholly unrelated text"})
    wb = _wb({"B4": "JOE BOB SMITH"},
             [["Agreements", "B4", "contract_instrument", "ASMT.TIF", 1, "grantor @ p.1"]])
    report = run_source_truth(wb, [tmp_path])
    hit = next(f for f in report.findings if f.cell == "B4")
    assert hit.outcome == UNVERIFIABLE
    assert report.truthful          # zero mismatches -> still shippable
    assert "not called a contradiction" in hit.detail


def test_a_cell_whose_source_is_not_a_document_is_unverifiable(tmp_path):
    wb = _wb({"B4": "EOG"},
             [["Agreements", "B4", "spreadsheet_row", "Lease to Well XREF", 1, ""]])
    report = run_source_truth(wb, [tmp_path])
    hit = next(f for f in report.findings if f.cell == "B4")
    assert hit.outcome == UNVERIFIABLE
    assert "not a document on disk" in hit.detail


def test_ownership_recording_mismatch_is_caught_end_to_end(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Ownership Report")
    ws.cell(8, 2).value = "0220344-006"
    ws.cell(8, 11).value = "475"      # John A Cox's recording on Hulme's row
    ws.cell(8, 12).value = "39"
    aud = wb.create_sheet("DD_Audit_Trail")
    aud.append(["Sheet", "Location", "Source Type", "Source Reference", "Confidence", "Extracted Text"])
    report = run_source_truth(wb, [tmp_path], dpu_by_lease=_DPU)
    mismatch = next(f for f in report.mismatches if f.field_name == "Record Book/Page")
    assert "Vol 405, Pg 343" in str(mismatch.source_value)
    assert not report.truthful      # a contradicted value blocks the DDR


def test_truth_table_reports_per_sheet_counts(tmp_path):
    _docs(tmp_path, {"ASMT.txt": "JOE BOB SMITH"})
    wb = _wb({"B4": "JOE BOB SMITH", "B5": "NOT IN THE DOC"},
             [["Agreements", "B4", "contract_instrument", "ASMT.TIF", 1, "grantor @ p.1"],
              ["Agreements", "B5", "contract_instrument", "ASMT.TIF", 1, "form @ p.1"]])
    report = run_source_truth(wb, [tmp_path])
    table = truth_table(report)
    assert "Agreements" in table and "confirmed" in table
    assert report.per_sheet["Agreements"].confirmed == 1
    assert report.per_sheet["Agreements"].unverifiable == 1


def test_audit_cites_parse_cell_row_and_page():
    wb = _wb({}, [["Ownership Report", "row 17", "to_pdf", "T-O SUPP COX", 1, "MI @ p.5"],
                  ["Agreements", "B23", "contract_instrument", "ASMT.TIF", 1, "grantor @ p.1"]])
    cites = read_audit_cites(wb)
    assert cites[0].row == 17 and cites[0].cell is None and cites[0].page == 5
    assert cites[1].cell == "B23" and cites[1].row is None and cites[1].page == 1


# --- VERIFIED_BY_RULE: computed cells checked against their own arithmetic ----

from types import SimpleNamespace  # noqa: E402
from src.dd.source_truth_check import (  # noqa: E402
    VERIFIED_BY_RULE, VERIFIED_BY_MANIFEST, check_computed_cells, check_manifest_cells, TruthReport,
)


def _own_wb(cells):
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("Ownership Report")
    for coord, v in cells.items():
        ws[coord] = v
    return wb


def test_net_acre_formula_is_verified_by_rule():
    wb = _own_wb({"B8": "0190122-001", "U8": "=R8*S8*T8"})
    report = TruthReport(); check_computed_cells(wb, report)
    hit = next(f for f in report.findings if f.cell == "U8")
    assert hit.outcome == VERIFIED_BY_RULE and "R*S*T" in hit.field_name


def test_headline_wi_formula_is_verified_by_rule():
    wb = _own_wb({"B8": "0190122-001", "D5": "=B5/A5"})
    report = TruthReport(); check_computed_cells(wb, report)
    assert any(f.cell == "D5" and f.outcome == VERIFIED_BY_RULE for f in report.findings)


def test_nri_consistent_with_its_inputs_is_by_rule():
    wb = _own_wb({"B8": "0190122-001", "G8": 0.2, "W8": 0, "V8": 0.8})
    report = TruthReport(); check_computed_cells(wb, report)
    assert any(f.cell == "V8" and f.outcome == VERIFIED_BY_RULE for f in report.findings)


def test_nri_inconsistent_with_its_inputs_is_a_mismatch_not_by_rule():
    wb = _own_wb({"B8": "0190122-001", "G8": 0.25, "W8": 0, "V8": 0.9})   # should be .75
    report = TruthReport(); check_computed_cells(wb, report)
    assert any(f.cell == "V8" and f.outcome == MISMATCH for f in report.findings)


def test_by_rule_is_not_a_dumping_ground():
    # A non-formula, non-computable cell gets NOTHING from the rule checker.
    wb = _own_wb({"B8": "0190122-001", "O8": "some legal prose"})
    report = TruthReport(); check_computed_cells(wb, report)
    assert not any(f.cell == "O8" for f in report.findings)


# --- VERIFIED_BY_MANIFEST: project values against the manifest ---------------

def _summary_wb(pairs):
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("Summary")
    for i, (label, value) in enumerate(pairs, start=2):
        ws.cell(i, 1).value = label
        ws.cell(i, 2).value = value
    return wb


_MANIFEST = SimpleNamespace(client_name="BKV Corp.", acquisition_name="Project Emirates",
                            operator="EOG Resources, Inc.", seller_name="EOG Resources, Inc.",
                            state="Texas", county="Montague", landman_name="Renato Alves")


def test_client_matching_the_manifest_is_verified_by_manifest():
    wb = _summary_wb([("Client:", "BKV Corp."), ("Operator(s):", "EOG Resources, Inc.")])
    report = TruthReport(); check_manifest_cells(wb, _MANIFEST, report)
    assert all(f.outcome == VERIFIED_BY_MANIFEST for f in report.findings)
    assert {f.cell for f in report.findings} == {"B2", "B3"}


def test_client_differing_from_the_manifest_is_a_mismatch():
    wb = _summary_wb([("Client:", "ACME LLC")])
    report = TruthReport(); check_manifest_cells(wb, _MANIFEST, report)
    hit = next(f for f in report.findings if f.cell == "B2")
    assert hit.outcome == MISMATCH and "BKV Corp." in str(hit.source_value)


def test_operator_casing_still_matches_the_manifest():
    wb = _summary_wb([("Operator(s):", "EOG RESOURCES INC")])
    report = TruthReport(); check_manifest_cells(wb, _MANIFEST, report)
    assert report.findings[0].outcome == VERIFIED_BY_MANIFEST


def test_a_non_manifest_label_is_left_alone():
    wb = _summary_wb([("Potential Defect(s):", "None identified")])
    report = TruthReport(); check_manifest_cells(wb, _MANIFEST, report)
    assert report.findings == []


# --- CONFIRMED_TRACT_LEVEL: WI against the opinion's EOG tract WI -------------

from src.dd.source_truth_check import (  # noqa: E402
    CONFIRMED_TRACT_LEVEL, check_wi_tract_level, check_lease_review_terms,
)


def test_wi_matching_the_opinions_tract_wi_is_confirmed_tract_level():
    opinion = ("OIL, GAS AND MINERAL LEASEHOLD ESTATE (all tracts):\fWorking Net Revenue\n"
               "EOG Resources, Inc. [(49.768518% x 4/5)] .84953704 .66203704")
    wb = _own_wb({"B8": "0220344-013", "T8": 0.84953704})
    report = TruthReport()
    check_wi_tract_level(wb, [("T-O SUPP COX.PDF", opinion)], report)
    hit = next(f for f in report.findings if f.cell == "T8")
    assert hit.outcome == CONFIRMED_TRACT_LEVEL
    assert "SUPP" in hit.source_file


def test_wi_with_no_stated_tract_wi_stays_unverifiable_named():
    wb = _own_wb({"B8": "0220344-013", "T8": 0.84953704})
    report = TruthReport()
    check_wi_tract_level(wb, [("T-O SUPP COX.PDF", "no leasehold table here")], report)
    hit = next(f for f in report.findings if f.cell == "T8")
    assert hit.outcome == UNVERIFIABLE and "no tract-level EOG" in hit.detail


def test_wi_value_present_but_not_in_a_leasehold_context_is_not_confirmed():
    # The decimal appears, but not near EOG/leasehold/working — not the WI table.
    opinion = "some unrelated table lists .84953704 as a survey bearing"
    wb = _own_wb({"B8": "0220344-013", "T8": 0.84953704})
    report = TruthReport()
    check_wi_tract_level(wb, [("doc.PDF", opinion)], report)
    assert report.findings[0].outcome == UNVERIFIABLE


# --- Lease Review term cells: per-lease document, lessor+value ----------------

def test_lease_review_term_confirmed_against_the_lease_document(tmp_path):
    (tmp_path / "LEASE.txt").write_text(
        "Oil and Gas Lease dated October 6, 2006 from Robert W. Hrabovsky to EnCana. "
        "Royalty of 1/5. Primary term two (2) years.", encoding="utf-8")
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("Lease Review")
    ws.cell(2, 1).value = "0190122-005"; ws.cell(2, 3).value = "Robert W. Hrabovsky"
    ws.cell(2, 10).value = 0.2                   # royalty 1/5
    index = SourceIndex([tmp_path])
    report = TruthReport()
    check_lease_review_terms(wb, index, report, lease_doc_finder=lambda ln: tmp_path / "LEASE.txt")
    hit = next(f for f in report.findings if f.cell == "J2")
    assert hit.outcome == CONFIRMED          # 0.2 == 1/5 via normalization


def test_lease_review_memorandum_lease_with_no_document_is_named(tmp_path):
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("Lease Review")
    ws.cell(2, 1).value = "0220344-021"; ws.cell(2, 3).value = "Marilyn J. Parker"
    ws.cell(2, 10).value = 0.2
    report = TruthReport()
    check_lease_review_terms(wb, SourceIndex([tmp_path]), report, lease_doc_finder=lambda ln: None)
    hit = next(f for f in report.findings if f.cell == "J2")
    assert hit.outcome == UNVERIFIABLE and "no recorded document on disk" in hit.detail
