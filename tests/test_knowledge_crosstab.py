"""PR-11 — provenance is not correctness.

Every case is a defect that shipped in the PR-10 artefact WITH a complete audit
trail. The whitelist gate proved a source had been cited; it could not prove the
right field had been read from it, and it could not see a cell nobody wrote.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from src.knowledge.clean_output import AUDIT_TRAIL_SHEET
from src.knowledge.cross_tab import (DataRegion, check_coverage, check_cross_tab,
                                     CrossTabError, assert_consistent)
from src.knowledge.kit_builder import (CellWrite, InadmissibleCitation,
                                       is_screen_capture)
from src.knowledge.recording_authority import (AUTHORITY_NAME,
                                               RECORDING_WRITERS,
                                               RecordingAuthority, normalise)

EXHIBIT_A = {
    "0220604-024": {"recording": "Vol. 567, Pg. 170"},
    "0220604-027": {"recording": "Vol. 571, Pg. 343"},
    "0220604-099": {"recording": ""},
}


def _authority() -> RecordingAuthority:
    return RecordingAuthority.from_exhibit_a(EXHIBIT_A)


# --- item 1: one point of truth for a recording -----------------------------

def test_the_authority_overrules_a_land_description_deed_reference():
    # `Volume 199, Page 356` is the deed reference out of the lease's own land
    # description. It reached Runsheet!C5 while the Lease Review of the same
    # workbook had 567.170 — both fully trailed.
    r = _authority().resolve("0220604-024", "Volume 199, Page 356")
    assert r.value == "Vol. 567, Pg. 170"
    assert r.adopted_from == "authority"
    assert "not adopted" in r.citation and "199" in r.rejected


def test_an_agreeing_extraction_needs_no_apology():
    r = _authority().resolve("0220604-024", "Volume 567, Page 170")
    assert r.value == "Vol. 567, Pg. 170" and "not adopted" not in r.citation
    assert r.citation == AUTHORITY_NAME


def test_where_the_authority_is_silent_the_extraction_is_used_and_says_so():
    r = _authority().resolve("0220604-099", "Vol. 900, Pg. 1")
    assert r.value == "Vol. 900, Pg. 1" and r.adopted_from == "extraction"
    assert "does not state a recording" in r.citation


def test_nothing_at_all_yields_nothing():
    r = _authority().resolve("0220604-099", None)
    assert r.value == "" and r.adopted_from == "none"


def test_normalise_compares_volume_and_page_not_prose():
    assert normalise("Vol. 567, Pg. 170") == normalise("Volume 567, Page 170")
    assert normalise("Volume 199, Page 336 of the Deed Records") == "199.336"


def test_book_and_page_split_out_for_the_worksheet_columns():
    assert _authority().resolve("0220604-027").book_page == ("571", "343")


@pytest.mark.parametrize("module,func", RECORDING_WRITERS)
def test_every_declared_recording_writer_consults_the_authority(module, func):
    # The fifth instance of "fixed at the point of error, not ruled at the point
    # of truth" is what this test exists to make impossible a sixth time.
    import importlib
    import inspect
    source = inspect.getsource(getattr(importlib.import_module(module), func))
    assert "authority" in source, f"{module}.{func} never asks the authority"


# --- item 2: a screen capture is not a citation -----------------------------

@pytest.mark.parametrize("citation", [
    "Captura de tela 2026-08-06 130244.png p.1",
    "screenshot of the county portal",
    "PRINT DA TELA.JPG",
])
def test_a_capture_cannot_source_a_client_facing_value(citation):
    assert is_screen_capture(citation)
    with pytest.raises(InadmissibleCitation, match="invariant 27"):
        CellWrite("Runsheet", 7, "H", "grantee", "CHOCTAW ENERGY", citation)


def test_a_blank_may_still_record_that_only_a_capture_exists():
    # The knowledge that a capture exists is worth keeping; it just cannot be
    # the source of a value.
    c = CellWrite("Runsheet", 7, "H", "grantee", "", "Captura de tela.png",
                  "captura_de_tela")
    assert c.declared_blank


def test_an_ordinary_document_citation_is_admissible():
    assert not is_screen_capture("LEASE - REC.TIF p.1")
    assert CellWrite("Lease Review", 2, "F", "recording", "Vol. 1, Pg. 2",
                     "LEASE - REC.TIF p.1").value


def test_the_reason_is_in_the_closed_vocabulary():
    from src.knowledge.loader import load_knowledge
    assert load_knowledge().is_valid_reason("captura_de_tela")


# --- item 3/4: cross-tab consistency and coverage ---------------------------

def _workbook(tmp_path: Path, *, lease_review: str, runsheet: str,
              runsheet_cite: str, recap: bool = True) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    lr = wb.create_sheet("Lease Review")
    lr["A2"], lr["F2"] = "L25  (0220604-024)", lease_review
    ow = wb.create_sheet("Ownership - Worksheet")
    if recap:
        for col in "ABCDE":
            ow[f"{col}5"] = "x"
    rs = wb.create_sheet("Runsheet")
    rs["A4"], rs["C4"], rs["G4"], rs["H4"] = 1, runsheet, "grantor", "grantee"
    tr = wb.create_sheet(AUDIT_TRAIL_SHEET)
    tr.append(["SHEET", "CELL", "FIELD", "VALUE", "CITED TO", "REASON IF BLANK"])
    tr.append(["Lease Review", "F2", "recording", lease_review, "cite", ""])
    tr.append(["Runsheet", "C4", "runsheet", runsheet, runsheet_cite, ""])
    if recap:
        for col in "ABCDE":
            tr.append(["Ownership - Worksheet", f"{col}5", "unit_recap", "x",
                       "cite", ""])
    path = tmp_path / "ddr.xlsx"
    wb.save(path)
    return path


def test_two_tabs_disagreeing_about_one_lease_fails(tmp_path):
    path = _workbook(tmp_path, lease_review="Vol. 567, Pg. 170",
                     runsheet="Volume 199, Page 356",
                     runsheet_cite="LEASE.TIF p.1 (lease 0220604-024)")
    report = check_cross_tab(path)
    assert report.failed
    assert any(f.kind == "cross_tab" and "0220604-024" in f.where
               for f in report.findings)


def test_two_tabs_agreeing_pass(tmp_path):
    path = _workbook(tmp_path, lease_review="Vol. 567, Pg. 170",
                     runsheet="Volume 567, Page 170",
                     runsheet_cite="LEASE.TIF p.1 (lease 0220604-024)")
    assert not check_cross_tab(path).failed


def test_a_tab_contradicting_exhibit_a_fails_even_when_tabs_agree(tmp_path):
    path = _workbook(tmp_path, lease_review="Vol. 199, Pg. 356",
                     runsheet="Volume 199, Page 356",
                     runsheet_cite="LEASE.TIF p.1 (lease 0220604-024)")
    report = check_cross_tab(path, {"0220604-024": "Vol. 567, Pg. 170"})
    assert any(f.kind == "contradicts_authority" for f in report.findings)


def test_the_runsheet_joins_by_the_lease_named_in_its_own_citation(tmp_path):
    # A citation of just "LEASE.TIF p.1" cannot be joined to anything, which is
    # how the contradiction survived a gate looking straight at both cells.
    path = _workbook(tmp_path, lease_review="Vol. 567, Pg. 170",
                     runsheet="Volume 199, Page 356",
                     runsheet_cite="LEASE.TIF p.1")
    assert not check_cross_tab(path).failed          # unjoinable, so unchecked
    (tmp_path / "b").mkdir()
    path2 = _workbook(tmp_path / "b", lease_review="Vol. 567, Pg. 170",
                      runsheet="Volume 199, Page 356",
                      runsheet_cite="LEASE.TIF p.1 (lease 0220604-024)")
    assert check_cross_tab(path2).failed


def test_an_unwritten_untrailed_recap_cell_fails_coverage(tmp_path):
    path = _workbook(tmp_path, lease_review="Vol. 567, Pg. 170",
                     runsheet="Volume 567, Page 170",
                     runsheet_cite="c (lease 0220604-024)", recap=False)
    report = check_coverage(path)
    assert report.failed and len(report.findings) == 5
    assert all(f.kind == "uncovered_cell" for f in report.findings)


def test_a_declared_blank_satisfies_coverage(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Ownership - Worksheet")
    tr = wb.create_sheet(AUDIT_TRAIL_SHEET)
    tr.append(["SHEET", "CELL", "FIELD", "VALUE", "CITED TO", "REASON IF BLANK"])
    for col in "ABCDE":
        tr.append(["Ownership - Worksheet", f"{col}5", "unit_recap", "", "",
                   "nao_computavel"])
    path = tmp_path / "ddr.xlsx"
    wb.save(path)
    assert not check_coverage(path).failed


def test_assert_consistent_raises_naming_the_problem(tmp_path):
    path = _workbook(tmp_path, lease_review="Vol. 567, Pg. 170",
                     runsheet="Volume 199, Page 356",
                     runsheet_cite="LEASE.TIF p.1 (lease 0220604-024)")
    with pytest.raises(CrossTabError, match="0220604-024"):
        assert_consistent(path)


# --- the standing regression ------------------------------------------------

COLD = Path("data/units/_josh_b_replay/KIT/DDR - Josh B Unit.xlsx")


@pytest.mark.skipif(not COLD.exists(), reason="cold kit not built here")
def test_the_cold_artefact_is_consistent_and_covered():
    import json
    work = Path("data/units/_josh_b_replay/_work/_josh_b_exhibitA.json")
    authority = RecordingAuthority.from_exhibit_a(
        json.loads(work.read_text(encoding="utf-8"))["leases"])
    assert_consistent(COLD, {lease: authority.resolve(lease).value
                             for lease in json.loads(
                                 work.read_text(encoding="utf-8"))["leases"]})


@pytest.mark.skipif(not COLD.exists(), reason="cold kit not built here")
def test_no_cell_in_the_cold_artefact_cites_a_screen_capture():
    wb = openpyxl.load_workbook(COLD, data_only=True)
    try:
        offenders = [
            (row[0], row[1], row[4])
            for row in wb[AUDIT_TRAIL_SHEET].iter_rows(min_row=2, values_only=True)
            if row and row[3] and is_screen_capture(str(row[4] or ""))]
    finally:
        wb.close()
    assert not offenders, offenders[:3]
