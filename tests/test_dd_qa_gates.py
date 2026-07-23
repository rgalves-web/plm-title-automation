"""Permanent DDR QA gates (Collin item 5)."""
from __future__ import annotations

from fractions import Fraction

import openpyxl

from src.dd.dd_qa_gates import (
    mineral_interest_sums_to_one,
    mineral_interest_within_tract_not_over_one,
    no_commentary_in_client_cells,
    no_template_sample_data,
    seller_wi_matches_expected,
    workbook_headline_wi,
)
from src.dd.title_opinion_parser import TOPartyInterest, TOTract, TitleOpinionExtraction
from src.llm_extractor import LLMFieldResult


def _tract(fractions: list[str]) -> TOTract:
    return TOTract(
        tract_legal_description=LLMFieldResult(value="tract", confidence=1.0),
        mineral_owners=[TOPartyInterest(party=f"O{i}", fraction_text=f, confidence=1.0)
                        for i, f in enumerate(fractions)],
    )


def _ext(tract: TOTract) -> TitleOpinionExtraction:
    return TitleOpinionExtraction(file_number="TO1", tracts=[tract])


def test_mi_sums_to_one_passes_when_exact():
    ext = _ext(_tract(["1/2", "1/4", "1/4"]))
    r = mineral_interest_sums_to_one({"TO1": ext})
    assert r.passed
    assert not r.locations


def test_mi_sums_to_one_fails_and_reports_residual():
    ext = _ext(_tract(["1/2", "1/4"]))  # sums to 3/4
    r = mineral_interest_sums_to_one({"TO1": ext})
    assert not r.passed
    assert r.locations and "residual" in r.locations[0]


def test_mi_gate_pinpoints_unparsed_owner_entries():
    # A complex/garbled fraction that won't parse must be named per owner.
    tract = TOTract(
        tract_legal_description=LLMFieldResult(value="t", confidence=1.0),
        mineral_owners=[
            TOPartyInterest(party="Lynn W. Cox", fraction_text="1/2", confidence=1.0),
            TOPartyInterest(party="Croxton Estate", fraction_text="[(1/4 x 1/9) plu", confidence=1.0),
        ],
    )
    r = mineral_interest_sums_to_one({"TO_COX": _ext(tract)})
    assert not r.passed
    pinpoint = [loc for loc in r.locations if "Croxton Estate" in loc and "UNPARSED" in loc]
    assert pinpoint, r.locations


def test_seller_wi_gate_matches():
    r = seller_wi_matches_expected(Fraction("0.816596"), "0.816596")
    assert r.passed


def test_seller_wi_gate_mismatch_reports_location():
    r = seller_wi_matches_expected(Fraction("0.625"), "0.816596")
    assert not r.passed
    assert r.locations and "Ownership Report" in r.locations[0]


def test_seller_wi_gate_skipped_when_unconfigured():
    assert seller_wi_matches_expected(Fraction("0.5"), None).passed


# --- Collin item 1: tract-level MI-sum sanity gate (catches the flattening) ---

def test_mi_within_tract_flags_flattened_over_one():
    # 13 leases in one tract all carrying 5/36 → 65/36 = 1.81 > 1 → must fail.
    rows = [{"lease_no": f"L{i}", "tract_key": "TO_COX::cox", "mi_value": Fraction(5, 36)}
            for i in range(13)]
    r = mineral_interest_within_tract_not_over_one(rows)
    assert not r.passed
    assert r.locations and "1.8" in r.locations[0]


def test_mi_within_tract_passes_distinct_partition():
    rows = [{"lease_no": "L1", "tract_key": "T::x", "mi_value": Fraction(1, 36)},
            {"lease_no": "L2", "tract_key": "T::x", "mi_value": Fraction(5, 72)},
            {"lease_no": "L3", "tract_key": "T::x", "mi_value": Fraction(5, 36)}]
    assert mineral_interest_within_tract_not_over_one(rows).passed


def test_mi_within_tract_ignores_rows_without_mi_or_tract():
    rows = [{"lease_no": "L1", "tract_key": "T::x", "mi_value": None},
            {"lease_no": "L2", "mi_value": Fraction(9, 10)}]  # no tract_key
    assert mineral_interest_within_tract_not_over_one(rows).passed


# --- Collin item 2: workbook headline WI (replicates Ownership Report D5) ---

def test_workbook_headline_wi_matches_when_all_rows_complete():
    # One tract, gross 100 ac, MI summing to 0.8 at WI 1 → net 80 / 100 = 0.8.
    rows = [{"lease_no": "L1", "tract_gross_acres": Fraction(100), "mi_value": Fraction(1, 2), "wi_value": Fraction(1)},
            {"lease_no": "L2", "tract_gross_acres": Fraction(100), "mi_value": Fraction(3, 10), "wi_value": Fraction(1)}]
    r = workbook_headline_wi(rows, Fraction(100), "0.8")
    assert r.passed, r.detail


def test_workbook_headline_wi_flags_understated_with_pending_count():
    # Second row pending MI/acres → headline 50/100 = 0.5 < expected 0.8 → fail,
    # and the gate must state a row is pending.
    rows = [{"lease_no": "L1", "tract_gross_acres": Fraction(100), "mi_value": Fraction(1, 2), "wi_value": Fraction(1)},
            {"lease_no": "L2", "tract_gross_acres": None, "mi_value": None, "wi_value": None}]
    r = workbook_headline_wi(rows, Fraction(100), "0.8")
    assert not r.passed
    assert "pending" in r.detail
    assert any("pending" in loc for loc in r.locations)


def test_workbook_headline_wi_skipped_without_config():
    assert workbook_headline_wi([], None, None).passed


def _wb_with(sheet: str, cell: str, value: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws[cell] = value
    return wb


def test_no_commentary_gate_flags_triage_text():
    wb = _wb_with("Lease Review", "C3", "REVIEW REQUIRED — needs manual check")
    r = no_commentary_in_client_cells(wb)
    assert not r.passed
    assert "Lease Review · C3" in r.locations[0]


def test_no_template_data_gate_flags_placeholder():
    wb = _wb_with("Runsheet", "A2", "636.13 acres in Reeves County")
    r = no_template_sample_data(wb)
    assert not r.passed


def test_gates_ignore_audit_trail_sheet():
    wb = openpyxl.Workbook()
    wb.active.title = "DD_Audit_Trail"
    wb.active["A1"] = "REVIEW REQUIRED note in the provenance layer"
    assert no_commentary_in_client_cells(wb).passed
