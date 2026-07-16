"""R1 (never regress below the approved baseline) and R2 (never overwrite a
manually verified value) — Maxwell lessons standing rules.

The synthetic-workbook tests always run. The real-baseline tests exercise the
approved v5 DDR, which is CLIENT MATERIAL: it lives in data/reference/ (gitignored)
and is never committed, so those tests skip when it is not on disk.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from src.dd.baseline_regression import compare_files, compare_to_baseline
from src.dd.protected_values import (
    apply_protected_values,
    find_protected_values,
    unresolved_protected_locations,
)

_BASELINE = Path("data/reference/MAXWELL_A_UNIT_DDR_v5_BASELINE.xlsx")


def _wb(cells: dict[str, dict[str, object]]):
    """cells = {sheet: {coord: value}}"""
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for sheet_name, values in cells.items():
        sheet = workbook.create_sheet(sheet_name)
        for coordinate, value in values.items():
            sheet[coordinate] = value
    return workbook


# --- R1 -------------------------------------------------------------------

def test_identical_workbook_has_no_regression():
    base = _wb({"Ownership Report": {"S8": "=1/4", "F8": "Lynn W Cox"}})
    assert compare_to_baseline(_wb({"Ownership Report": {"S8": "=1/4", "F8": "Lynn W Cox"}}), base).passed


def test_lost_value_is_a_regression():
    base = _wb({"Ownership Report": {"S8": "=1/4"}})
    result = compare_to_baseline(_wb({"Ownership Report": {"S8": None}}), base)
    assert not result.passed
    assert len(result.lost) == 1
    assert "value lost" in result.lost[0].describe()


def test_changed_verified_value_is_a_regression():
    base = _wb({"Ownership Report": {"S16": "=5/288"}})
    result = compare_to_baseline(_wb({"Ownership Report": {"S16": "=5/36"}}), base)
    assert not result.passed
    assert len(result.changed) == 1


def test_populating_more_than_baseline_is_not_a_regression():
    base = _wb({"Ownership Report": {"S8": "=1/4"}})
    candidate = _wb({"Ownership Report": {"S8": "=1/4", "S9": "=1/8", "S10": "=1/2"}})
    result = compare_to_baseline(candidate, base)
    assert result.passed
    assert result.candidate_populated > result.baseline_populated


def test_template_dashes_do_not_count_as_populated():
    base = _wb({"Lease Review": {"C3": "-", "D3": "--", "E3": "  "}})
    assert compare_to_baseline(_wb({"Lease Review": {}}), base).passed


def test_formula_satisfies_equivalent_numeric_baseline():
    base = _wb({"Ownership Report": {"S8": 0.25}})
    assert compare_to_baseline(_wb({"Ownership Report": {"S8": "=1/4"}}), base).passed


def test_audit_layer_is_not_compared():
    base = _wb({"DD_Audit_Trail": {"A2": "spreadsheet_row"}, "Summary": {"B2": "EOG"}})
    assert compare_to_baseline(_wb({"DD_Audit_Trail": {}, "Summary": {"B2": "EOG"}}), base).passed


def test_missing_sheet_fails():
    base = _wb({"Agreements": {"A20": "K1400049"}})
    result = compare_to_baseline(_wb({"Summary": {"B2": "x"}}), base)
    assert not result.passed and result.missing_sheets == ["Agreements"]


# --- R2 -------------------------------------------------------------------

def _baseline_with_audit():
    workbook = _wb({
        "Ownership Report": {"S16": "=5/288", "S21": "=5/36"},
        "DD_Audit_Trail": {
            "A1": "Sheet", "B1": "Location", "C1": "Source Type", "D1": "Source Reference",
            "A2": "Ownership Report", "B2": "S16", "C2": "manual_verified",
            "D2": "T-O SUPP LYNN W COX pp.5-6; Hulme corrected 5/36->5/288",
            "A3": "Ownership Report", "B3": "cols S/T Cox-tract rows", "C3": "manual_verified",
            "D3": "region — not a single cell",
        },
    })
    return workbook


def test_protected_values_found_from_audit_trail():
    protected = find_protected_values(_baseline_with_audit())
    assert [(p.sheet, p.coordinate, p.value) for p in protected] == [
        ("Ownership Report", "S16", "=5/288")]


def test_region_audit_rows_are_reported_not_guessed():
    unresolved = unresolved_protected_locations(_baseline_with_audit())
    assert len(unresolved) == 1 and "Cox-tract rows" in unresolved[0]


def test_recomputed_value_never_overwrites_manual_and_conflict_is_itemized():
    protected = find_protected_values(_baseline_with_audit())
    # The pipeline re-introduces the very error the human fixed.
    candidate = _wb({"Ownership Report": {"S16": "=5/36"}})
    result = apply_protected_values(candidate, protected)
    assert not result.passed
    assert candidate["Ownership Report"]["S16"].value == "=5/288"  # human wins
    assert "keeping the verified value" in result.conflicts[0]


def test_agreeing_recomputation_carries_forward_without_conflict():
    protected = find_protected_values(_baseline_with_audit())
    candidate = _wb({"Ownership Report": {"S16": "=5/288"}})
    result = apply_protected_values(candidate, protected)
    assert result.passed and len(result.carried_forward) == 1


def test_blank_regen_cell_is_filled_from_protected_without_conflict():
    protected = find_protected_values(_baseline_with_audit())
    candidate = _wb({"Ownership Report": {"S16": None}})
    result = apply_protected_values(candidate, protected)
    assert result.passed
    assert candidate["Ownership Report"]["S16"].value == "=5/288"


# --- real approved baseline (skips when the client file is absent) ---------

@pytest.mark.skipif(not _BASELINE.exists(), reason="approved v5 baseline not on disk (client data)")
def test_v5_baseline_declares_its_manually_verified_values():
    workbook = openpyxl.load_workbook(_BASELINE)
    protected = find_protected_values(workbook)
    unresolved = unresolved_protected_locations(workbook)
    # v5's audit trail records manual work; every such row must be either a
    # protectable cell or explicitly surfaced for a human.
    assert protected or unresolved


@pytest.mark.skipif(not _BASELINE.exists(), reason="approved v5 baseline not on disk (client data)")
def test_v5_baseline_does_not_regress_against_itself():
    result = compare_files(_BASELINE, _BASELINE)
    assert result.passed, result.summary()
    assert result.baseline_populated > 0
