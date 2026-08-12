"""Round 1 — tests for the plumbing, and ONLY for the plumbing.

There is nothing here that asserts a title conclusion, because no code in this
package reaches one. What these tests protect is the promise the architecture
makes: that Python opens files, moves pages, carries messages and saves bytes —
faithfully, and without choosing anything on the model's behalf.

So the assertions are of one kind: the model asked for X and got exactly X. A
mode that yields nothing yields nothing; a knowledge miss stays a miss; a cell
written twice keeps both writes; a run that stops early SAYS it stopped early.
"""
from __future__ import annotations

import json
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from src.examinador.cache import PageCache, digest, key_for
from src.examinador.corpus import Corpus
from src.examinador.ledger import Ledger
from src.examinador.loop import RATES, Usage, run_loop
from src.examinador.okf import Card, KnowledgeBase
from src.examinador.tools import TOOLS, build_dispatch


# --- the corpus -------------------------------------------------------------

def _pdf(text: str) -> bytes:
    import fitz
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


@pytest.fixture()
def corpus(tmp_path: Path) -> Corpus:
    files = tmp_path / "FILES"
    files.mkdir()
    with zipfile.ZipFile(files / "pacote.zip", "w") as zf:
        zf.writestr("0220604-024 (walker)/lease.pdf", _pdf("THIS LEASE Walker"))
        zf.writestr("0220604-025 (jones)/lease.pdf", _pdf("THIS LEASE Jones"))
    return Corpus(files, cache=PageCache(tmp_path / "cache"))


def test_list_documents_filters_by_substring_and_nothing_else(corpus: Corpus):
    everything = corpus.list_documents()
    assert len(everything) == 2
    only = corpus.list_documents("walker")
    assert [d["documento"] for d in only] == \
        ["0220604-024 (walker)/lease.pdf"]


def test_the_text_mode_returns_the_text_layer(corpus: Corpus):
    got = corpus.read_page("0220604-024 (walker)/lease.pdf", 1, "texto")
    assert "THIS LEASE Walker" in got["text"]
    assert got["caracteres"] == len(got["text"])


def test_a_mode_that_yields_nothing_never_falls_through_to_a_costlier_one(
        tmp_path: Path):
    """The code choosing to spend money is the code deciding. A scan read in
    `texto` comes back empty, and the model is the one who escalates."""
    files = tmp_path / "FILES"
    files.mkdir()
    from PIL import Image
    img = Image.new("RGB", (80, 40), "white")
    img.save(files / "scan.tif")
    c = Corpus(files, cache=PageCache(tmp_path / "cache"))
    got = c.read_page("scan.tif", 1, "texto")
    assert got["text"] == ""
    assert got["kind"] == "text"


def test_the_image_mode_returns_pixels(corpus: Corpus):
    got = corpus.read_page("0220604-024 (walker)/lease.pdf", 1, "imagem")
    assert got["kind"] == "image"
    assert got["png"][:8] == b"\x89PNG\r\n\x1a\n"


def test_search_lists_what_it_could_not_search_instead_of_calling_it_absent(
        tmp_path: Path):
    files = tmp_path / "FILES"
    files.mkdir()
    from PIL import Image
    Image.new("RGB", (40, 40), "white").save(files / "scan.tif")
    (files / "born.pdf").write_bytes(_pdf("Volume 569 Page 324"))
    c = Corpus(files, cache=PageCache(tmp_path / "cache"))
    got = c.search("Volume 569")
    assert got["no_texto"] == [{"documento": "born.pdf", "paginas": [1]}]
    assert "scan.tif" in got["nao_pesquisavel_sem_ocr"]
    assert "nao e prova de ausencia" in got["nota"]


def test_the_page_cache_is_global_and_keyed_by_the_page_itself(tmp_path: Path):
    shared = PageCache(tmp_path / "global")
    one = Corpus(_corpus_dir(tmp_path, "unidade_a"), cache=shared)
    two = Corpus(_corpus_dir(tmp_path, "unidade_b"), cache=shared)
    one.read_page("lease.pdf", 1, "texto")
    before = shared.hits
    two.read_page("lease.pdf", 1, "texto")     # same bytes, other unit
    assert shared.hits == before + 1


def _corpus_dir(tmp_path: Path, name: str) -> Path:
    d = tmp_path / name
    d.mkdir()
    # The same BYTES in both folders — a PDF written twice is two different byte
    # strings, and this test is about identical pages, not identical text.
    if not hasattr(_corpus_dir, "bytes"):
        _corpus_dir.bytes = _pdf("THIS LEASE identical bytes")
    (d / "lease.pdf").write_bytes(_corpus_dir.bytes)
    return d


# --- the knowledge base -----------------------------------------------------

def test_a_card_survives_the_round_trip(tmp_path: Path):
    kb = KnowledgeBase(tmp_path / "okf")
    kb.write("cadeia encana eog", "Encana cedeu a EOG.", type="cadeia",
             sources=["K1400170.pdf p.2"], links=["formulario eog barnett"])
    got = kb.read("cadeia encana eog")
    assert got["encontrado"] and got["tipo"] == "cadeia"
    assert got["fontes"] == ["K1400170.pdf p.2"]
    assert got["links"] == ["formulario eog barnett"]


def test_a_miss_returns_candidates_never_the_nearest_card_as_the_answer(
        tmp_path: Path):
    kb = KnowledgeBase(tmp_path / "okf")
    kb.write("formulario eog barnett 2008", "…", type="formulario")
    got = kb.read("formulario eog barnett 2011")
    assert got["encontrado"] is False
    assert "corpo" not in got
    assert got["todas_as_fichas"] == ["formulario-eog-barnett-2008"]


def test_index_and_log_are_maintained_as_the_model_writes(tmp_path: Path):
    kb = KnowledgeBase(tmp_path / "okf")
    kb.write("layout de parecer", "A frase de certificacao fica na ultima "
                                  "pagina.", type="layout")
    kb.write("layout de parecer", "Corrigido: penultima pagina.", type="layout")
    index = (tmp_path / "okf" / "index.md").read_text(encoding="utf-8")
    log = (tmp_path / "okf" / "log.md").read_text(encoding="utf-8")
    assert "layout de parecer" in index
    assert index.count("layout-de-parecer.md") == 1      # one card, not two
    assert "criou" in log and "atualizou" in log


def test_the_frontmatter_is_yaml_and_the_body_is_markdown(tmp_path: Path):
    kb = KnowledgeBase(tmp_path / "okf")
    kb.write("uma coisa", "# titulo\n\ncorpo", type="regra")
    raw = (tmp_path / "okf" / "uma-coisa.md").read_text(encoding="utf-8")
    assert raw.startswith("---\n")
    assert Card.from_markdown(raw).body == "# titulo\n\ncorpo"


# --- the ledger -------------------------------------------------------------

def test_a_cell_written_twice_keeps_both_writes_and_ships_the_last():
    led = Ledger()
    led.write_cell("Lease Review", "c5", "Smith", "a.tif", 1, "SMITH", "media")
    second = led.write_cell("Lease Review", "C5", "Smyth", "b.tif", 2, "SMYTH",
                            "alta")
    assert second["reescreveu"] == 1
    assert len(led.cells) == 2
    assert led.final_cells()[("Lease Review", "C5")].valor == "Smyth"


def test_the_ledger_records_a_blank_with_its_reason_and_owner():
    led = Ledger()
    led.record_blank("Runsheet", "D7", "copia sem carimbo",
                     "operador: confirmar no cartorio")
    assert led.blanks[0].dono == "operador: confirmar no cartorio"
    assert led.counts()["brancos"] == 1


def test_the_ledger_does_not_judge_what_it_records():
    """No vocabulary, no citation check, no plausibility. Round 1 measures the
    model unaided, so the notebook may not quietly correct it."""
    led = Ledger()
    led.write_cell("Aba Que Nao Existe", "ZZ999", "", "", "", "", "seja la")
    led.record_blank("Aba", "A1", "porque sim", "")
    assert led.counts() == {"escritas": 1, "celulas_distintas": 1,
                            "brancos": 1, "perguntas": 0}


# --- the loop ---------------------------------------------------------------

@dataclass
class _Use:
    name: str
    input: dict[str, Any]
    id: str = "u1"
    type: str = "tool_use"


@dataclass
class _Text:
    text: str
    type: str = "text"


@dataclass
class _Usage:
    input_tokens: int = 10
    output_tokens: int = 5
    cache_creation_input_tokens: int = 0
    cache_read_input_tokens: int = 0


@dataclass
class _Message:
    content: list[Any]
    usage: _Usage


class _FakeClient:
    """Replays a script of assistant turns; records what it was sent."""

    def __init__(self, script: list[list[Any]]) -> None:
        self.script = script
        self.seen: list[list[dict[str, Any]]] = []
        self.messages = self

    def create(self, **kwargs: Any) -> _Message:
        # A snapshot, not the live list: the loop keeps mutating the same object
        # and a stored reference would show every turn as the last one.
        import copy
        self.seen.append(copy.deepcopy(kwargs["messages"]))
        return _Message(content=self.script.pop(0), usage=_Usage())


def test_the_loop_dispatches_a_tool_call_and_hands_the_result_back():
    led = Ledger()
    dispatch = {"write_cell": lambda **kw: led.write_cell(**kw)}
    client = _FakeClient([
        [_Use("write_cell", {"aba": "Lease Review", "celula": "C5",
                             "valor": "Smith", "documento": "a.tif",
                             "pagina": "1", "trecho": "SMITH",
                             "confianca": "alta"})],
        [_Text("pronto")],
    ])
    log, messages = run_loop(system="s", task="t", tools=[], dispatch=dispatch,
                             model="claude-opus-5", client=client)
    assert log.stop_reason == "o modelo parou de pedir ferramentas"
    assert log.to_dict()["ordem_das_ferramentas"] == ["write_cell"]
    assert led.final_cells()[("Lease Review", "C5")].valor == "Smith"
    result = messages[2]["content"][0]
    assert result["type"] == "tool_result" and result["is_error"] is False


def test_a_tool_that_raises_is_reported_to_the_model_not_swallowed():
    def boom(**kw: Any):
        raise FileNotFoundError("nao esta no corpus")

    client = _FakeClient([[_Use("read_page", {"documento": "x"})],
                          [_Text("entendi")]])
    log, messages = run_loop(system="s", task="t", tools=[],
                             dispatch={"read_page": boom},
                             model="claude-opus-5", client=client)
    result = messages[2]["content"][0]
    assert result["is_error"] is True
    assert "FileNotFoundError" in result["content"]
    assert log.turns[0].calls[0]["erro"] is True


def test_an_image_result_travels_as_an_image_block():
    png = b"\x89PNG\r\n\x1a\n" + b"0" * 10
    client = _FakeClient([[_Use("read_page", {})], [_Text("vi")]])
    _, messages = run_loop(
        system="s", task="t", tools=[],
        dispatch={"read_page": lambda **kw: {"kind": "image", "png": png,
                                             "documento": "d", "pagina": 1}},
        model="claude-opus-5", client=client)
    blocks = messages[2]["content"][0]["content"]
    assert blocks[0]["type"] == "image"
    assert blocks[0]["source"]["media_type"] == "image/png"


def test_a_run_that_hits_its_turn_ceiling_says_it_is_incomplete():
    client = _FakeClient([[_Use("noop", {})] for _ in range(5)])
    log, _ = run_loop(system="s", task="t", tools=[],
                      dispatch={"noop": lambda **kw: {"ok": True}},
                      model="claude-opus-5", max_turns=3, client=client)
    assert "INCOMPLETA" in log.stop_reason


def test_the_cost_ceiling_stops_the_run_and_says_so():
    client = _FakeClient([[_Use("noop", {})] for _ in range(5)])
    log, _ = run_loop(system="s", task="t", tools=[],
                      dispatch={"noop": lambda **kw: {"ok": True}},
                      model="claude-opus-5", max_turns=5,
                      max_cost_usd=0.0000001, client=client)
    assert "teto de custo" in log.stop_reason
    assert "INCOMPLETA" in log.stop_reason


def test_the_assignment_is_cached_and_the_breakpoint_rolls_forward():
    """By turn fifty the unchanged prefix is most of the bill."""
    client = _FakeClient([[_Use("noop", {})], [_Use("noop", {})],
                          [_Text("fim")]])
    run_loop(system="s", task="ordem longa", tools=[],
             dispatch={"noop": lambda **kw: {"ok": True}},
             model="claude-opus-5", client=client)
    first, last = client.seen[0], client.seen[-1]
    assert first[0]["content"][0]["cache_control"] == {"type": "ephemeral"}
    assert last[0]["content"][0]["cache_control"] == {"type": "ephemeral"}
    marked = [i for i, m in enumerate(last[1:], start=1)
              if isinstance(m.get("content"), list)
              and any(isinstance(b, dict) and "cache_control" in b
                      for b in m["content"])]
    assert marked == [len(last) - 1]


def test_every_model_the_round_uses_has_a_declared_rate():
    usage = Usage(input=1_000_000)
    assert usage.cost("claude-opus-5") == RATES["claude-opus-5"]["input"]
    assert usage.cost("um-modelo-sem-tarifa") is None


# --- the tools ---------------------------------------------------------------

def test_the_nine_tools_are_exactly_the_ones_the_order_names():
    assert [t["name"] for t in TOOLS] == [
        "list_documents", "read_page", "search_corpus", "read_knowledge",
        "write_knowledge", "get_precedent", "write_cell", "record_blank",
        "flag_for_human"]


def test_write_cell_cannot_be_called_without_a_citation():
    required = next(t for t in TOOLS if t["name"] == "write_cell")
    assert set(required["input_schema"]["required"]) == {
        "aba", "celula", "valor", "documento", "pagina", "trecho", "confianca"}


def test_every_tool_has_a_dispatch_entry(corpus: Corpus, tmp_path: Path):
    dispatch = build_dispatch(corpus, KnowledgeBase(tmp_path / "okf"), Ledger())
    assert set(dispatch) == {t["name"] for t in TOOLS}


def test_a_precedent_that_was_never_recorded_comes_back_empty(
        corpus: Corpus, tmp_path: Path):
    kb = KnowledgeBase(tmp_path / "okf")
    kb.write("cert date da a-319", "9/9/2010, dito pelo operador.",
             type="veredito")
    dispatch = build_dispatch(corpus, kb, Ledger())
    assert dispatch["get_precedent"]("cert date")["precedentes"]
    miss = dispatch["get_precedent"]("profundidade")
    assert miss["precedentes"] == []
    assert "Sem precedente" in miss["nota"]


# --- the measurement ---------------------------------------------------------

def _book(path: Path, cells: dict[str, dict[str, Any]]) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for tab, values in cells.items():
        ws = wb.create_sheet(tab)
        for coord, value in values.items():
            ws[coord] = value
    wb.save(path)
    return path


def test_a_cell_the_template_already_held_is_not_credited_to_the_model(
        tmp_path: Path):
    """The easiest way to fake this exercise is to count the template's own
    values as reproductions. They get their own bucket."""
    from src.examinador.compare import compare

    hand = _book(tmp_path / "hand.xlsx", {"Lease Review": {"C5": "Smith",
                                                           "C6": "Jones"}})
    template = _book(tmp_path / "tpl.xlsx", {"Lease Review": {"C6": "Jones"}})
    ia = _book(tmp_path / "ia.xlsx", {"Lease Review": {"C5": "Smith",
                                                       "C6": "Jones"}})
    ledger = tmp_path / "ledger.json"
    ledger.write_text(json.dumps({
        "celulas": [{"aba": "Lease Review", "celula": "C5"}],
        "brancos": []}), encoding="utf-8")

    got = compare(ia, ledger, hand=hand, template=template).to_dict()
    assert got["resultado"]["concordou"] == 1
    assert got["resultado"]["herdado_do_template_nao_conta"] == 1


def test_the_populations_are_declared_before_the_rates(tmp_path: Path):
    from src.examinador.compare import compare

    hand = _book(tmp_path / "hand.xlsx", {"Runsheet": {"A4": "1"}})
    template = _book(tmp_path / "tpl.xlsx", {"Runsheet": {}})
    ia = _book(tmp_path / "ia.xlsx", {"Runsheet": {"A4": "2"}})
    ledger = tmp_path / "ledger.json"
    ledger.write_text(json.dumps({
        "celulas": [{"aba": "Runsheet", "celula": "A4"}], "brancos": []}),
        encoding="utf-8")
    got = compare(ia, ledger, hand=hand, template=template).to_dict()
    names = [p["populacao"] for p in got["populacoes_declaradas_antes"]]
    assert names == ["gabarito", "escritas_pela_ia", "template"]
    assert got["resultado"]["divergiu"] == 1


def test_template_leftovers_are_counted_as_a_defect_of_their_own(tmp_path: Path):
    from src.examinador.compare import leftovers

    template = _book(tmp_path / "tpl.xlsx",
                     {"Summary": {"B2": "*** Energy", "B3": "A.E. Carroll"}})
    ia = _book(tmp_path / "ia.xlsx",
               {"Summary": {"B2": "EOG Resources", "B3": "A.E. Carroll"}})
    ledger = tmp_path / "ledger.json"
    ledger.write_text(json.dumps({
        "celulas": [{"aba": "Summary", "celula": "B2"}], "brancos": []}),
        encoding="utf-8")
    got = leftovers(ia, ledger, template=template)
    assert got["celulas_do_template_ainda_no_entregavel"] == 1
    assert "A.E. Carroll" in got["amostra"][0]


# --- the migration -----------------------------------------------------------

def test_two_files_with_the_same_name_are_two_concepts(tmp_path: Path):
    """`fixtures/a/fixture.yaml` and `fixtures/b/fixture.yaml` are two lessons.
    Naming a card after the file stem collapses one onto the other, silently."""
    from src.examinador.migrate import migrate

    src = tmp_path / "knowledge"
    for name in ("primeira", "segunda"):
        d = src / "fixtures" / name
        d.mkdir(parents=True)
        (d / "fixture.yaml").write_text(f"lesson_id: {name}\n", encoding="utf-8")
    kb = KnowledgeBase(tmp_path / "okf")
    got = migrate(src, kb)
    assert got["fichas_escritas"] == 2
    assert kb.stats()["fichas"] == 2


def test_a_file_of_many_concepts_becomes_many_cards(tmp_path: Path):
    from src.examinador.migrate import migrate

    src = tmp_path / "knowledge"
    (src / "rules").mkdir(parents=True)
    (src / "rules" / "motivos.yaml").write_text(
        "reasons:\n  um:\n    label: A\n  dois:\n    label: B\n",
        encoding="utf-8")
    kb = KnowledgeBase(tmp_path / "okf")
    migrate(src, kb)
    assert kb.read("motivos: um")["encontrado"]
    assert kb.read("motivos: dois")["encontrado"]


def test_every_migrated_card_names_the_file_it_came_from(tmp_path: Path):
    from src.examinador.migrate import migrate

    src = tmp_path / "knowledge"
    (src / "ontology").mkdir(parents=True)
    (src / "ontology" / "entities.yaml").write_text("a: 1\n", encoding="utf-8")
    kb = KnowledgeBase(tmp_path / "okf")
    migrate(src, kb)
    card = kb.read("entities")
    assert card["fontes"] == ["knowledge/ontology/entities.yaml"]
    assert card["tipo"] == "entidade"


# --- the quarantine ----------------------------------------------------------

def test_the_delivered_kit_is_quarantined_like_the_rest_of_the_answer_sheet():
    """The hand-built workbook sits in the project folder beside the corpus the
    run legitimately reads. Only the comparison harness opens it."""
    from src.knowledge.quarantine import is_quarantined
    assert is_quarantined("BKV122TX - Project Emirates/JOSH B UNIT/"
                          "1. Landman Submitted/PLM Report/DDR - Josh B.xlsx")
    assert not is_quarantined("BKV122TX - Project Emirates/JOSH B UNIT/"
                              "FILES/pacote.zip")
