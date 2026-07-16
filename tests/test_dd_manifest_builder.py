"""Tests for src/dd/manifest_builder.py — all fixtures are synthetic, tiny
workbooks built with openpyxl, never the real 65k-row reference files."""
from __future__ import annotations

import zipfile
from pathlib import Path

import openpyxl
import pytest

from src.dd.manifest_builder import (
    ManifestBuildError,
    build_and_write_manifest,
    build_manifest,
)

# ---------------------------------------------------------------------------
# Fixture builders
# ---------------------------------------------------------------------------

PROJECT_LOG_HEADER = [
    "VALUE SORT", "OPERATOR", "COUNTY", "WELL NAME", "UNIT NAME",
    "WELL STATUS", "AOV%", "Scope of Review",
]

XREF_HEADER = [
    "INCLUDE FLAG", "LEASE NO", "LESSOR", "LESSEE", "ACTIVE FLAG",
    "RECORD TYPE", "RIGHTS HELD", "PROPERTY STATUS", "LEASE DATE", "EFF DATE",
    "EXP DATE", "EXT EXP DATE", "STATE", "COUNTY", "WELL OR CONTRACT NB",
    "WELL OR CONTRACT NAME", "ACTIVE FLAG2", "RECORD TYPE3", "RIGHTS HELD4",
    "WELL PROD STATUS", "AGMT DATE", "EFF DATE5", "WELL OPERATOR", "WELL API",
    "STATE6", "COUNTY7", "UNIT LEVEL NAME", "UNIT LEVEL PRELIM AV %",
    "UNIT LEVEL MATCH METHOD", "UNIT LEVEL MATCH KEY",
]

WELL_LIST_HEADER = (
    ["OPERATOR", "COUNTY", "WELL NAME", "PAYOUT", "WELL #",
     "CURRENT WI", "CURRENT NRI"]
    + [f"COL{i}" for i in range(7, 28)]
    + ["UNIT NAME", "WELL STATUS"]
)
# Columns resolve BY HEADER NAME (src/dd/tabular_source), so these names must be
# the real ones the client's export prints — the loader no longer reads by
# position, and a fixture with placeholder names would not exercise real code.
# index: 0=OPERATOR 1=COUNTY 2=WELL NAME ... 5=CURRENT WI 6=CURRENT NRI ... 28=UNIT NAME 29=WELL STATUS


def _make_project_log(tmp_path: Path, rows: list[list]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Project Log"
    ws.append(PROJECT_LOG_HEADER)
    for row in rows:
        ws.append(row)
    path = tmp_path / "project_log.xlsx"
    wb.save(path)
    return path


def _make_xref(tmp_path: Path, rows: list[list], filename: str = "xref.xlsx") -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(XREF_HEADER)
    for row in rows:
        padded = row + [None] * (30 - len(row))
        ws.append(padded)
    path = tmp_path / filename
    wb.save(path)
    return path


def _xref_row(
    lease_no=None, lessor=None, lessee=None, record_type3=None,
    well_or_contract_nb=None, unit_level_name=None, match_method=None,
    rights_held=None, property_status=None,
) -> list:
    row = [None] * 30
    row[1] = lease_no
    row[2] = lessor
    row[3] = lessee
    row[6] = rights_held
    row[7] = property_status
    row[14] = well_or_contract_nb
    row[17] = record_type3
    row[26] = unit_level_name
    row[28] = match_method
    return row


def _make_well_list(tmp_path: Path, rows: list[list]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Wells"
    ws.append(WELL_LIST_HEADER)
    for row in rows:
        padded = row + [None] * (30 - len(row))
        ws.append(padded)
    path = tmp_path / "well_list.xlsx"
    wb.save(path)
    return path


def _well_row(operator=None, county=None, well_name=None, current_wi=None,
              current_nri=None, unit_name=None, well_status=None) -> list:
    row = [None] * 30
    row[0] = operator
    row[1] = county
    row[2] = well_name
    row[5] = current_wi
    row[6] = current_nri
    row[28] = unit_name
    row[29] = well_status
    return row


# ---------------------------------------------------------------------------
# Project Log matching
# ---------------------------------------------------------------------------

def test_build_manifest_finds_unit_in_project_log(tmp_path):
    project_log = _make_project_log(
        tmp_path, [[1, "EOG RESOURCES INC", "MONTAGUE", "TEST UNIT #1H", "TEST UNIT", "PRODUCING", 0.005, "Supplemental Title"]]
    )
    xref = _make_xref(tmp_path, [])
    well_list = _make_well_list(tmp_path, [])

    manifest, _ = build_manifest("TEST UNIT", project_log, xref, well_list)

    assert manifest.operator == "EOG RESOURCES INC"
    assert manifest.county == "MONTAGUE"
    assert manifest.scope_of_review == "Supplemental Title"
    assert manifest.seller_name == "EOG RESOURCES INC"


def test_build_manifest_raises_when_unit_not_found(tmp_path):
    project_log = _make_project_log(tmp_path, [[1, "EOG", "MONTAGUE", "W1", "OTHER UNIT", "PRODUCING", 0.005, ""]])
    xref = _make_xref(tmp_path, [])
    well_list = _make_well_list(tmp_path, [])

    with pytest.raises(ManifestBuildError, match="not found"):
        build_manifest("MISSING UNIT", project_log, xref, well_list)


def test_build_manifest_raises_when_unit_ambiguous(tmp_path):
    project_log = _make_project_log(
        tmp_path,
        [
            [1, "EOG", "MONTAGUE", "W1", "DUP UNIT", "PRODUCING", 0.005, ""],
            [2, "EOG", "MONTAGUE", "W2", "DUP UNIT", "PRODUCING", 0.005, ""],
        ],
    )
    xref = _make_xref(tmp_path, [])
    well_list = _make_well_list(tmp_path, [])

    with pytest.raises(ManifestBuildError, match="ambiguous"):
        build_manifest("DUP UNIT", project_log, xref, well_list)


# ---------------------------------------------------------------------------
# confirm_coverage
# ---------------------------------------------------------------------------

def test_confirm_coverage_true_when_only_semicolon_token_match(tmp_path):
    project_log = _make_project_log(tmp_path, [[1, "EOG", "MONTAGUE", "W1", "TARGET UNIT", "PRODUCING", 0.005, ""]])
    xref = _make_xref(
        tmp_path,
        [
            _xref_row(
                lease_no="L1", record_type3="TITLE OPINIONS",
                well_or_contract_nb="T0000001",
                unit_level_name="OTHER UNIT; TARGET UNIT; THIRD UNIT",
                match_method="Lease-number fallback from direct unit matches",
            )
        ],
    )
    well_list = _make_well_list(tmp_path, [])

    manifest, _ = build_manifest("TARGET UNIT", project_log, xref, well_list)

    assert len(manifest.title_opinions) == 1
    assert manifest.title_opinions[0].confirm_coverage is True


def test_confirm_coverage_false_when_exact_single_value_aa_exists(tmp_path):
    project_log = _make_project_log(tmp_path, [[1, "EOG", "MONTAGUE", "W1", "TARGET UNIT", "PRODUCING", 0.005, ""]])
    xref = _make_xref(
        tmp_path,
        [
            _xref_row(
                lease_no="L1", record_type3="TITLE OPINIONS",
                well_or_contract_nb="T0000002",
                unit_level_name="TARGET UNIT",
                match_method="Direct unit match",
            )
        ],
    )
    well_list = _make_well_list(tmp_path, [])

    manifest, _ = build_manifest("TARGET UNIT", project_log, xref, well_list)

    assert len(manifest.title_opinions) == 1
    assert manifest.title_opinions[0].confirm_coverage is False


# ---------------------------------------------------------------------------
# Dedup + lease collection
# ---------------------------------------------------------------------------

def test_title_opinion_dedup_across_multiple_lessor_rows(tmp_path):
    project_log = _make_project_log(tmp_path, [[1, "EOG", "MONTAGUE", "W1", "TARGET UNIT", "PRODUCING", 0.005, ""]])
    xref = _make_xref(
        tmp_path,
        [
            _xref_row(lease_no="L1", lessor="PARTY A", record_type3="TITLE OPINIONS",
                      well_or_contract_nb="T0000003", unit_level_name="TARGET UNIT; OTHER"),
            _xref_row(lease_no="L2", lessor="PARTY B", record_type3="TITLE OPINIONS",
                      well_or_contract_nb="T0000003", unit_level_name="TARGET UNIT; THIRD"),
        ],
    )
    well_list = _make_well_list(tmp_path, [])

    manifest, build_log = build_manifest("TARGET UNIT", project_log, xref, well_list)

    assert len(manifest.title_opinions) == 1
    to = manifest.title_opinions[0]
    assert to.file_number == "T0000003"
    assert set(to.matched_unit_tokens) == {"TARGET UNIT", "OTHER", "THIRD"}
    assert len(build_log["title_opinion_rows"]["T0000003"]) == 2


def test_leases_collected_independent_of_record_type(tmp_path):
    project_log = _make_project_log(tmp_path, [[1, "EOG", "MONTAGUE", "W1", "TARGET UNIT", "PRODUCING", 0.005, ""]])
    xref = _make_xref(
        tmp_path,
        [
            _xref_row(lease_no="L9", lessor="SOME LESSOR", lessee="SOME LESSEE",
                      record_type3="WELL", well_or_contract_nb="14738-000",
                      unit_level_name="TARGET UNIT"),
        ],
    )
    well_list = _make_well_list(tmp_path, [])

    manifest, _ = build_manifest("TARGET UNIT", project_log, xref, well_list)

    assert len(manifest.title_opinions) == 0
    assert len(manifest.leases) == 1
    assert manifest.leases[0].lease_no == "L9"
    assert manifest.leases[0].lessor == "SOME LESSOR"


def test_reference_contracts_classified_separately_from_title_opinions(tmp_path):
    project_log = _make_project_log(tmp_path, [[1, "EOG", "MONTAGUE", "W1", "TARGET UNIT", "PRODUCING", 0.005, ""]])
    xref = _make_xref(
        tmp_path,
        [
            _xref_row(lease_no="L1", lessor="A", lessee="B", record_type3="FARMOUT AGREEMENT",
                      well_or_contract_nb="K1400170", unit_level_name="TARGET UNIT; OTHER"),
        ],
    )
    well_list = _make_well_list(tmp_path, [])

    manifest, _ = build_manifest("TARGET UNIT", project_log, xref, well_list)

    assert len(manifest.reference_contracts) == 1
    assert manifest.reference_contracts[0].record_type == "FARMOUT AGREEMENT"
    assert len(manifest.title_opinions) == 0


# ---------------------------------------------------------------------------
# PDF availability probing
# ---------------------------------------------------------------------------

def test_pdf_available_false_for_empty_zip(tmp_path):
    project_log = _make_project_log(tmp_path, [[1, "EOG", "MONTAGUE", "W1", "TARGET UNIT", "PRODUCING", 0.005, ""]])
    xref = _make_xref(
        tmp_path,
        [_xref_row(lease_no="L1", record_type3="TITLE OPINIONS",
                  well_or_contract_nb="T0000005", unit_level_name="TARGET UNIT; OTHER")],
    )
    well_list = _make_well_list(tmp_path, [])

    to_dir = tmp_path / "to_files"
    to_dir.mkdir()
    empty_zip = to_dir / "T0000005 (NO DOCUMENTS).zip"
    with zipfile.ZipFile(empty_zip, "w") as zf:
        zf.writestr("T0000005 (NO DOCUMENTS)/", "")

    manifest, _ = build_manifest("TARGET UNIT", project_log, xref, well_list, to_files_dir=to_dir)

    assert manifest.title_opinions[0].pdf_available is False
    assert manifest.title_opinions[0].pdf_path is None


def test_pdf_available_true_for_zip_with_pdf_member(tmp_path):
    project_log = _make_project_log(tmp_path, [[1, "EOG", "MONTAGUE", "W1", "TARGET UNIT", "PRODUCING", 0.005, ""]])
    xref = _make_xref(
        tmp_path,
        [_xref_row(lease_no="L1", record_type3="TITLE OPINIONS",
                  well_or_contract_nb="T0000006", unit_level_name="TARGET UNIT; OTHER")],
    )
    well_list = _make_well_list(tmp_path, [])

    to_dir = tmp_path / "to_files"
    to_dir.mkdir()
    real_zip = to_dir / "T0000006 (SOME OPINION).zip"
    with zipfile.ZipFile(real_zip, "w") as zf:
        zf.writestr("T0000006 (SOME OPINION)/opinion.pdf", b"%PDF-1.4 fake content")

    manifest, _ = build_manifest("TARGET UNIT", project_log, xref, well_list, to_files_dir=to_dir)

    assert manifest.title_opinions[0].pdf_available is True
    assert manifest.title_opinions[0].pdf_path is not None


# ---------------------------------------------------------------------------
# Well List / represented_interest
# ---------------------------------------------------------------------------

def test_represented_interest_from_well_list_single_well(tmp_path):
    project_log = _make_project_log(tmp_path, [[1, "EOG", "MONTAGUE", "W1", "TARGET UNIT", "PRODUCING", 0.005, ""]])
    xref = _make_xref(tmp_path, [])
    well_list = _make_well_list(
        tmp_path,
        [_well_row(operator="EOG", county="MONTAGUE", well_name="TARGET UNIT #1H",
                   current_wi=0.9975, current_nri=0.78813617, unit_name="TARGET UNIT",
                   well_status="PRODUCING")],
    )

    manifest, _ = build_manifest("TARGET UNIT", project_log, xref, well_list)

    assert manifest.represented_interest is not None
    assert manifest.represented_interest.current_wi == "0.9975"
    assert manifest.represented_interest.current_nri == "0.78813617"
    assert manifest.manifest_notes == []


def test_represented_interest_logs_note_when_wells_disagree(tmp_path):
    project_log = _make_project_log(tmp_path, [[1, "EOG", "MONTAGUE", "W1", "TARGET UNIT", "PRODUCING", 0.005, ""]])
    xref = _make_xref(tmp_path, [])
    well_list = _make_well_list(
        tmp_path,
        [
            _well_row(operator="EOG", well_name="TARGET UNIT #1H", current_wi=0.9975,
                      current_nri=0.78813617, unit_name="TARGET UNIT"),
            _well_row(operator="EOG", well_name="TARGET UNIT #2H", current_wi=0.5,
                      current_nri=0.4, unit_name="TARGET UNIT"),
        ],
    )

    manifest, _ = build_manifest("TARGET UNIT", project_log, xref, well_list)

    assert len(manifest.manifest_notes) == 1
    assert "disagree" in manifest.manifest_notes[0]


# ---------------------------------------------------------------------------
# manifest.yaml round trip
# ---------------------------------------------------------------------------

# ---------------------------------------------------------------------------
# Wells - Prelim AV% embedded in Project Log (no separate well_list_path)
# ---------------------------------------------------------------------------

WELLS_PRELIM_HEADER = (
    ["OPERATOR", "COUNTY", "WELL NAME", "PAYOUT", "WELL #",
     "CURRENT WI", "CURRENT NRI"]
    + [f"COL{i}" for i in range(7, 28)]
    + ["UNIT NAME", "WELL STATUS"]
)


def _make_project_log_with_wells_prelim(tmp_path: Path, log_rows: list[list], well_rows: list[list]) -> Path:
    """Project Log workbook that also contains a 'Wells - Prelim AV%' sheet."""
    wb = openpyxl.Workbook()
    ws_log = wb.active
    ws_log.title = "Project Log"
    ws_log.append(PROJECT_LOG_HEADER)
    for row in log_rows:
        ws_log.append(row)

    ws_wells = wb.create_sheet("Wells - Prelim AV%")
    ws_wells.append(WELLS_PRELIM_HEADER)
    for row in well_rows:
        padded = row + [None] * (30 - len(row))
        ws_wells.append(padded)

    path = tmp_path / "project_log_with_wells.xlsx"
    wb.save(path)
    return path


def test_well_list_read_from_embedded_prelim_sheet(tmp_path):
    """When well_list_path is None, wells are read from 'Wells - Prelim AV%' in the Project Log."""
    project_log = _make_project_log_with_wells_prelim(
        tmp_path,
        log_rows=[[1, "EOG", "MONTAGUE", "W1", "TARGET UNIT", "PRODUCING", 0.005, ""]],
        well_rows=[_well_row(operator="EOG", county="MONTAGUE", well_name="TARGET UNIT #1H",
                             current_wi=0.9975, current_nri=0.78813617, unit_name="TARGET UNIT",
                             well_status="PRODUCING")],
    )
    xref = _make_xref(tmp_path, [])

    manifest, _ = build_manifest("TARGET UNIT", project_log, xref, well_list_path=None)

    assert manifest.represented_interest is not None
    assert manifest.represented_interest.current_wi == "0.9975"
    assert manifest.represented_interest.current_nri == "0.78813617"
    assert manifest.wells[0].well_name == "TARGET UNIT #1H"


def test_well_list_embedded_prelim_filters_by_unit(tmp_path):
    """Wells from a different unit in the embedded sheet are excluded."""
    project_log = _make_project_log_with_wells_prelim(
        tmp_path,
        log_rows=[[1, "EOG", "MONTAGUE", "W1", "TARGET UNIT", "PRODUCING", 0.005, ""]],
        well_rows=[
            _well_row(well_name="TARGET UNIT #1H", current_wi=0.9975, current_nri=0.78, unit_name="TARGET UNIT"),
            _well_row(well_name="OTHER UNIT #1H", current_wi=0.5, current_nri=0.4, unit_name="OTHER UNIT"),
        ],
    )
    xref = _make_xref(tmp_path, [])

    manifest, _ = build_manifest("TARGET UNIT", project_log, xref, well_list_path=None)

    assert len(manifest.wells) == 1
    assert manifest.wells[0].well_name == "TARGET UNIT #1H"


def test_source_workbooks_records_project_log_when_no_well_list_path(tmp_path):
    project_log = _make_project_log_with_wells_prelim(
        tmp_path,
        log_rows=[[1, "EOG", "MONTAGUE", "W1", "TARGET UNIT", "PRODUCING", 0.005, ""]],
        well_rows=[],
    )
    xref = _make_xref(tmp_path, [])

    manifest, _ = build_manifest("TARGET UNIT", project_log, xref, well_list_path=None)

    assert manifest.source_workbooks["well_list"] == str(project_log)


# ---------------------------------------------------------------------------
# XREF Sheet2 merging
# ---------------------------------------------------------------------------


def _make_xref_with_sheet2(
    tmp_path: Path,
    sheet1_rows: list[list],
    sheet2_rows: list[list],
    filename: str = "xref.xlsx",
) -> Path:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(XREF_HEADER)
    for row in sheet1_rows:
        ws1.append(row + [None] * (30 - len(row)))

    ws2 = wb.create_sheet("Sheet2")
    ws2.append(XREF_HEADER)
    for row in sheet2_rows:
        ws2.append(row + [None] * (30 - len(row)))

    path = tmp_path / filename
    wb.save(path)
    return path


def test_xref_sheet2_title_opinions_are_merged(tmp_path):
    """Title opinions found only in Sheet2 appear in the manifest."""
    project_log = _make_project_log(tmp_path, [[1, "EOG", "MONTAGUE", "W1", "TARGET UNIT", "PRODUCING", 0.005, ""]])
    xref = _make_xref_with_sheet2(
        tmp_path,
        sheet1_rows=[
            _xref_row(lease_no="L1", record_type3="TITLE OPINIONS",
                      well_or_contract_nb="T0000010", unit_level_name="TARGET UNIT"),
        ],
        sheet2_rows=[
            _xref_row(lease_no="L2", record_type3="TITLE OPINIONS",
                      well_or_contract_nb="T0000011", unit_level_name="TARGET UNIT"),
        ],
    )
    well_list = _make_well_list(tmp_path, [])

    manifest, build_log = build_manifest("TARGET UNIT", project_log, xref, well_list)

    file_numbers = {to.file_number for to in manifest.title_opinions}
    assert "T0000010" in file_numbers
    assert "T0000011" in file_numbers


def test_xref_sheet2_deduplicates_with_sheet1(tmp_path):
    """The same title opinion appearing in both sheets is deduplicated (merged) to one entry."""
    project_log = _make_project_log(tmp_path, [[1, "EOG", "MONTAGUE", "W1", "TARGET UNIT", "PRODUCING", 0.005, ""]])
    xref = _make_xref_with_sheet2(
        tmp_path,
        sheet1_rows=[
            _xref_row(lease_no="L1", lessor="PARTY A", record_type3="TITLE OPINIONS",
                      well_or_contract_nb="T0000012", unit_level_name="TARGET UNIT"),
        ],
        sheet2_rows=[
            _xref_row(lease_no="L2", lessor="PARTY B", record_type3="TITLE OPINIONS",
                      well_or_contract_nb="T0000012", unit_level_name="TARGET UNIT"),
        ],
    )
    well_list = _make_well_list(tmp_path, [])

    manifest, build_log = build_manifest("TARGET UNIT", project_log, xref, well_list)

    assert len(manifest.title_opinions) == 1
    assert manifest.title_opinions[0].file_number == "T0000012"
    assert len(build_log["title_opinion_rows"]["T0000012"]) == 2


def test_xref_sheet2_leases_are_merged(tmp_path):
    """Leases unique to Sheet2 are included in manifest.leases."""
    project_log = _make_project_log(tmp_path, [[1, "EOG", "MONTAGUE", "W1", "TARGET UNIT", "PRODUCING", 0.005, ""]])
    xref = _make_xref_with_sheet2(
        tmp_path,
        sheet1_rows=[
            _xref_row(lease_no="L-SHEET1", lessor="A", record_type3="WELL",
                      unit_level_name="TARGET UNIT"),
        ],
        sheet2_rows=[
            _xref_row(lease_no="L-SHEET2", lessor="B", record_type3="WELL",
                      unit_level_name="TARGET UNIT"),
        ],
    )
    well_list = _make_well_list(tmp_path, [])

    manifest, _ = build_manifest("TARGET UNIT", project_log, xref, well_list)

    lease_nos = {l.lease_no for l in manifest.leases}
    assert "L-SHEET1" in lease_nos
    assert "L-SHEET2" in lease_nos


def test_xref_without_sheet2_still_works(tmp_path):
    """Workbooks with only Sheet1 are unaffected — Sheet2 absence is silently ignored."""
    project_log = _make_project_log(tmp_path, [[1, "EOG", "MONTAGUE", "W1", "TARGET UNIT", "PRODUCING", 0.005, ""]])
    xref = _make_xref(tmp_path, [
        _xref_row(lease_no="L1", record_type3="TITLE OPINIONS",
                  well_or_contract_nb="T0000020", unit_level_name="TARGET UNIT"),
    ])
    well_list = _make_well_list(tmp_path, [])

    manifest, _ = build_manifest("TARGET UNIT", project_log, xref, well_list)

    assert len(manifest.title_opinions) == 1
    assert manifest.title_opinions[0].file_number == "T0000020"


# ---------------------------------------------------------------------------
# manifest.yaml round trip
# ---------------------------------------------------------------------------

def test_manifest_yaml_round_trips_fraction_safe_strings(tmp_path):
    project_log = _make_project_log(tmp_path, [[1, "EOG", "MONTAGUE", "W1", "TARGET UNIT", "PRODUCING", 0.005, ""]])
    xref = _make_xref(tmp_path, [])
    well_list = _make_well_list(
        tmp_path,
        [_well_row(operator="EOG", well_name="TARGET UNIT #1H", current_wi=0.9975,
                   current_nri=0.78813617, unit_name="TARGET UNIT")],
    )

    manifest, manifest_path, _ = build_and_write_manifest(
        "TARGET UNIT", project_log, xref, well_list, output_root=tmp_path / "out"
    )

    import yaml
    reloaded = yaml.safe_load(manifest_path.read_text(encoding="utf-8"))
    assert reloaded["represented_interest"]["current_wi"] == "0.9975"
    assert reloaded["represented_interest"]["current_nri"] == "0.78813617"
