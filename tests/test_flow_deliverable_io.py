"""Round-trip verified saves: the disk, not memory, is what gets asserted."""
from __future__ import annotations

import openpyxl
import pytest

from src.flow.deliverable_io import RoundTripError, roundtrip_verify, save_and_verify


def _book(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Summary"
    wb["Summary"]["B9"] = "07/27/2026"
    own = wb.create_sheet("Ownership Report")
    own["S15"] = 0.08333333
    at = wb.create_sheet("DD_Audit_Trail")
    at.append(["Sheet", "Location", "Type", "Ref", "Conf"])
    at.append(["Summary", "B2", "manifest", "x", "high"])
    p = tmp_path / "DDR.xlsx"
    wb.save(p)
    return wb, p


def test_save_and_verify_passes_when_disk_matches(tmp_path):
    wb, p = _book(tmp_path)
    wb["Summary"]["B9"] = "07/28/2026"
    wb["Ownership Report"]["S15"] = "=1/12"
    wb["DD_Audit_Trail"].append(["Ownership Report", "row 15", "document", "T-O", "high"])
    save_and_verify(
        wb, p,
        expect_sheets={"Summary", "Ownership Report", "DD_Audit_Trail"},
        expect_absent_sheets={"Ownership - 1H (No FO Option)"},
        expect_cells={("Summary", "B9"): "07/28/2026",
                      ("Ownership Report", "S15"): "=1/12"},
        expect_audit_rows=2,
    )


def test_an_unsaved_edit_is_caught_by_the_disk_read(tmp_path):
    # The Bowen failure mode: object edited, disk not carrying it.
    wb, p = _book(tmp_path)
    wb["Summary"]["B9"] = "07/28/2026"      # edited in memory, NEVER saved
    with pytest.raises(RoundTripError, match="B9"):
        roundtrip_verify(p, expect_cells={("Summary", "B9"): "07/28/2026"})


def test_a_resurrected_excluded_sheet_is_caught(tmp_path):
    wb, p = _book(tmp_path)
    wb.create_sheet("Ownership - 1H (No FO Option)").sheet_state = "hidden"
    wb.save(p)
    with pytest.raises(RoundTripError, match="must be absent"):
        roundtrip_verify(p, expect_absent_sheets={"Ownership - 1H (No FO Option)"})


def test_audit_row_count_mismatch_is_caught(tmp_path):
    wb, p = _book(tmp_path)
    with pytest.raises(RoundTripError, match="audit rows"):
        roundtrip_verify(p, expect_audit_rows=33)     # disk has 1


def test_sheet_set_mismatch_is_caught(tmp_path):
    wb, p = _book(tmp_path)
    with pytest.raises(RoundTripError, match="sheet set mismatch"):
        roundtrip_verify(p, expect_sheets={"Summary", "Ownership Report"})


def test_missing_file_is_a_roundtrip_error(tmp_path):
    with pytest.raises(RoundTripError, match="does not exist"):
        roundtrip_verify(tmp_path / "nope.xlsx")
