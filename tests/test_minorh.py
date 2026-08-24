"""O sistema operável — testes do encanamento, e SÓ do encanamento.

Nenhuma asserção aqui conclui title: o pacote não conclui. O que se protege é
a operação — o censo persiste, `continuar` acha só o que é novo (por hash), a
resposta do operador é lida do arquivo, um campo provado não regride, a regra
de embarque nunca deixa célula reprovada no arquivo, e as linhas se alinham
por identidade antes da comparação.
"""
from __future__ import annotations

import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from src.examinador.cache import PageCache
from src.examinador.corpus import Corpus
from src.examinador.ledger import Ledger
from src.examinador.okf import KnowledgeBase
from src.minorh.arquivos import (escrever_pedidos, escrever_perguntas,
                                 ler_respostas)
from src.minorh.estado import Estado
from src.minorh.ferramentas import FERRAMENTAS, montar_dispatch
from src.minorh.memoria import Memoria


# --- memória ----------------------------------------------------------------

def test_a_memoria_conta_os_tres_estados(tmp_path: Path):
    m = Memoria(tmp_path / "memoria.json")
    m.gravar("lease 0999001-001", "royalty", "1/8", "lease.pdf", 1,
             "one-eighth (1/8)", "alta")
    m.branco("lease 0999001-001", "profundidade", "o registro nao diz",
             "operador")
    m.pendente("lease 0999001-004", "lessor", "PED-001")
    assert m.contar() == {"escritos": 1, "brancos": 1, "pendentes": 1}


def test_um_campo_provado_nao_regride_a_pendente(tmp_path: Path):
    """Um pedido amplo não apaga o que já foi provado com citação."""
    m = Memoria(tmp_path / "memoria.json")
    m.gravar("lease 0999001-001", "royalty", "1/8", "lease.pdf", 1, "1/8",
             "alta")
    m.pendente("lease 0999001-001", "royalty", "PED-001")
    campo = m.entidades["lease 0999001-001"]["royalty"]
    assert campo.estado == "escrito" and campo.valor == "1/8"


def test_a_memoria_sobrevive_ao_processo(tmp_path: Path):
    m = Memoria(tmp_path / "memoria.json")
    m.gravar("tract 1", "acres", "40.00", "dpu.pdf", 3, "40.00 acres", "alta")
    de_novo = Memoria(tmp_path / "memoria.json")
    assert de_novo.entidades["tract 1"]["acres"].valor == "40.00"


# --- estado -----------------------------------------------------------------

def _corpus(tmp_path: Path, nomes: dict[str, bytes]) -> Corpus:
    files = tmp_path / "FILES"
    files.mkdir(exist_ok=True)
    for nome, dados in nomes.items():
        (files / nome).write_bytes(dados)
    return Corpus(files, cache=PageCache(tmp_path / "cache"))


def test_continuar_acha_o_novo_pelo_conteudo_nao_pelo_nome(tmp_path: Path):
    corpus = _corpus(tmp_path, {"a.csv": b"conteudo a", "b.csv": b"conteudo b"})
    estado = Estado(tmp_path)
    estado.marcar_lido("a.csv", Estado.hash_de(b"conteudo a"))
    assert estado.novos_documentos(corpus) == ["b.csv"]
    # renomear nao torna novo; substituir o conteudo torna
    corpus2 = _corpus(tmp_path, {"a_renomeado.csv": b"conteudo a"})
    assert "a_renomeado.csv" not in estado.novos_documentos(corpus2)
    corpus3 = _corpus(tmp_path, {"a.csv": b"conteudo TROCADO"})
    assert "a.csv" in estado.novos_documentos(corpus3)


def test_o_censo_persiste_e_o_custo_acumula(tmp_path: Path):
    estado = Estado(tmp_path)
    estado.declarar_censo(42, "a conta")
    estado.registrar_execucao("unidade", {"entrada": 10, "saida": 5,
                                          "cache_escrita": 0,
                                          "cache_leitura": 0,
                                          "custo_usd": 1.5})
    estado.registrar_execucao("continuar", {"entrada": 1, "saida": 1,
                                            "cache_escrita": 0,
                                            "cache_leitura": 0,
                                            "custo_usd": 0.25})
    relido = Estado(tmp_path)
    assert relido.censo["exige"] == 42
    assert relido.custo_acumulado()["custo_usd"] == 1.75


# --- PEDIDOS.md / PERGUNTAS.md ----------------------------------------------

def test_a_resposta_do_operador_e_lida_do_arquivo(tmp_path: Path):
    estado = Estado(tmp_path)
    estado.nova_pergunta("Qual registro vale para a lease 0999001-003?",
                         ["Vol. 100, Pg. 3", "Vol. 100, Pg. 33"],
                         "a certidao do cartorio")
    path = escrever_perguntas(estado)
    # placeholder nao conta como resposta
    assert ler_respostas(estado) == []
    texto = path.read_text(encoding="utf-8").replace(
        "_(escreva aqui e rode `minorh continuar`)_",
        "Vale Vol. 100, Pg. 33 — conferido no cartorio em 10/08.")
    path.write_text(texto, encoding="utf-8")
    respostas = ler_respostas(estado)
    assert len(respostas) == 1
    assert "Pg. 33" in respostas[0]["resposta"]


def test_pergunta_respondida_nao_e_lida_de_novo(tmp_path: Path):
    estado = Estado(tmp_path)
    qid = estado.nova_pergunta("pergunta?", [], "um documento")
    estado.responder(qid, "resposta dada")
    escrever_perguntas(estado)
    assert ler_respostas(estado) == []


def test_pedidos_md_carrega_os_campos_que_o_documento_destrava(tmp_path: Path):
    estado = Estado(tmp_path)
    estado.novo_pedido("Lease 0999001-004", "o instrumento gravado",
                       "Seller / cartorio de Palo Duro", "requisicao ao Seller",
                       [{"entidade": "lease 0999001-004", "campo": "lessor"}])
    texto = escrever_pedidos(estado).read_text(encoding="utf-8")
    assert "PED-001" in texto
    assert "lease 0999001-004 :: lessor" in texto


# --- ferramentas ------------------------------------------------------------

@pytest.fixture()
def ctx(tmp_path: Path) -> dict[str, Any]:
    corpus = _corpus(tmp_path, {"doc.csv": b"API,NOME\n42-999-00001,POCO\n"})
    estado = Estado(tmp_path)
    memoria = Memoria(tmp_path / "_minorh" / "memoria.json")
    kb_u = KnowledgeBase(tmp_path / "_minorh" / "fichas")
    kb_g = KnowledgeBase(tmp_path / "okf_global")
    return {"dispatch": montar_dispatch(corpus, kb_u, kb_g, memoria, estado),
            "estado": estado, "memoria": memoria, "kb_u": kb_u, "kb_g": kb_g}


def test_toda_ferramenta_tem_dispatch(ctx):
    assert set(ctx["dispatch"]) == {t["name"] for t in FERRAMENTAS}


def test_declarar_censo_grava_estado_e_ficha(ctx):
    got = ctx["dispatch"]["declarar_censo"](exige=35, resumo="5 leases x 7")
    assert got == {"censo_declarado": 35}
    assert ctx["estado"].censo["exige"] == 35
    assert ctx["kb_u"].read("censo da unidade")["encontrado"]


def test_pedir_documento_deixa_os_campos_pendentes(ctx):
    ctx["dispatch"]["pedir_documento"](
        titulo="Lease 0999001-005", o_que="o instrumento", onde="Seller",
        metodo="requisicao",
        campos=[{"entidade": "lease 0999001-005", "campo": "lessor"},
                {"entidade": "lease 0999001-005", "campo": "royalty"}])
    assert ctx["memoria"].contar()["pendentes"] == 2
    assert ctx["estado"].pedidos[0]["status"] == "aberto"


def test_ler_pagina_marca_o_documento_como_lido(ctx):
    ctx["dispatch"]["read_page"](documento="doc.csv", pagina=1, modo="texto")
    assert len(ctx["estado"].lidos) == 1


def test_read_knowledge_procura_na_unidade_antes_do_global(ctx):
    ctx["kb_g"].write("formulario x", "versao global", type="formulario")
    assert ctx["dispatch"]["read_knowledge"](
        conceito="formulario x")["origem"] == "global"
    ctx["kb_u"].write("formulario x", "versao da unidade", type="formulario")
    got = ctx["dispatch"]["read_knowledge"](conceito="formulario x")
    assert got["origem"] == "unidade" and got["corpo"] == "versao da unidade"


def test_write_knowledge_roteia_pelo_escopo(ctx):
    ctx["dispatch"]["write_knowledge"](
        conceito="cadeia sintetica", ficha="...", tipo="cadeia",
        escopo="global")
    assert ctx["kb_g"].read("cadeia sintetica")["encontrado"]
    assert not ctx["kb_u"].read("cadeia sintetica")["encontrado"]


# --- auditor ----------------------------------------------------------------

def test_sem_segunda_chave_o_auditor_reserva_e_declarado(monkeypatch,
                                                         tmp_path):
    from src.minorh.auditor import MODELO_RESERVA, Auditor
    # chdir: `escolher` recarrega ".env" do cwd, e o .env real do repo pode
    # ter a chave do segundo provedor — o teste é sobre a AUSÊNCIA dela.
    monkeypatch.chdir(tmp_path)
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.setenv("ANTHROPIC_API_KEY", "chave-de-teste")
    auditor = Auditor.escolher()
    assert auditor.provedor == "anthropic"
    assert auditor.modelo == MODELO_RESERVA
    assert "LIMITACAO REGISTRADA" in auditor.limitacao


def test_veredito_ilegivel_reprova_nunca_aprova(monkeypatch):
    from src.minorh.auditor import Auditor
    auditor = Auditor(provedor="anthropic", modelo="claude-sonnet-5")
    monkeypatch.setattr(auditor, "_chamar", lambda prompt: "nao sei dizer")
    v = auditor.auditar(unidade="X", celula="A!B1", valor="v",
                        documento="d.pdf", pagina=1, trecho="t",
                        pagina_conteudo="...", modo="texto", listagem="")
    assert v.aprovada is False
    assert "veredito legivel" in v.motivo


def test_o_veredito_do_modelo_e_aplicado_sem_reinterpretacao(monkeypatch):
    from src.minorh.auditor import Auditor
    auditor = Auditor(provedor="anthropic", modelo="claude-sonnet-5")
    monkeypatch.setattr(auditor, "_chamar", lambda prompt: json.dumps(
        {"sustenta": True, "contradiz": "Wells!A3", "unidade_certa": True,
         "motivo": "contradiz a aba de pocos"}))
    v = auditor.auditar(unidade="X", celula="A!B1", valor="v",
                        documento="d.pdf", pagina=1, trecho="t",
                        pagina_conteudo="...", modo="texto", listagem="")
    assert v.aprovada is False and v.contradiz == "Wells!A3"


# --- ship: a regra de embarque ----------------------------------------------

def test_nada_embarca_reprovado(tmp_path: Path):
    """Célula reprovada duas vezes vira branco declarado no arquivo final."""
    from src.examinador import run as run_mod
    from src.minorh.auditor import Veredito
    from src.minorh.ship import aplicar_regra_de_embarque

    template = tmp_path / "tpl.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Lease Review"
    wb.save(template)
    wb.close()

    ledger = Ledger()
    ledger.write_cell("Lease Review", "B2", "valor errado", "d.pdf", 1,
                      "trecho", "alta")
    reprovada = Veredito(celula="Lease Review!B2", valor="valor errado",
                         citacao="d.pdf p.1", aprovada=False, sustenta=False,
                         contradiz=None, unidade_certa=True,
                         motivo="a pagina nao diz isso")
    removidas = aplicar_regra_de_embarque(ledger, [reprovada])
    assert removidas == ["Lease Review!B2"]

    out = run_mod.write_workbook(ledger, tmp_path / "ddr.xlsx",
                                 template=template)
    check = openpyxl.load_workbook(out)
    assert check["Lease Review"]["B2"].value is None
    trilha = [str(r[7]) for r in check["DD_Audit_Trail"].iter_rows(
        min_row=2, values_only=True) if r[7]]
    check.close()
    assert any("reprovada 2x" in t for t in trilha)


def test_a_reauditoria_cobre_o_que_a_correcao_escreveu_de_novo(tmp_path: Path):
    """`apenas` restringe; célula rebaixada a branco não é auditada — ela não
    embarca."""
    from src.examinador import run as run_mod
    from src.minorh.ship import auditar_tudo

    template = tmp_path / "tpl.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Aba"
    wb.save(template)
    wb.close()

    ledger = Ledger()
    ledger.write_cell("Aba", "A1", "fica", "d.csv", 1, "x", "alta")
    ledger.write_cell("Aba", "A2", "sai", "d.csv", 1, "y", "alta")
    ledger.record_blank("Aba", "A2", "rebaixada", "ninguem")
    out = run_mod.write_workbook(ledger, tmp_path / "ddr.xlsx",
                                 template=template)

    class _AuditorFake:
        def auditar(self, **kw):
            from src.minorh.auditor import Veredito
            return Veredito(celula=kw["celula"], valor=kw["valor"],
                            citacao="d.csv p.1", aprovada=True, sustenta=True,
                            contradiz=None, unidade_certa=True, motivo="ok")

    corpus = _corpus(tmp_path, {"d.csv": b"x y conteudo longo o suficiente "
                                         b"para os quarenta caracteres"})
    vereditos, contagem = auditar_tudo(
        unidade="X", workbook=out, ledger=ledger, corpus=corpus,
        auditor=_AuditorFake())
    assert [v.celula for v in vereditos] == ["Aba!A1"]
    assert contagem["vivas"] == 1


# --- compare: alinhamento por identidade ------------------------------------

def _livro(path: Path, abas: dict[str, dict[str, Any]]) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for aba, celulas in abas.items():
        ws = wb.create_sheet(aba)
        for coord, valor in celulas.items():
            ws[coord] = valor
    wb.save(path)
    return path


def test_linha_em_posicao_diferente_e_layout_nao_erro(tmp_path: Path):
    """A mesma lease na linha 5 de um lado e na 9 do outro é UMA lease."""
    from src.examinador.compare import compare

    mao = _livro(tmp_path / "mao.xlsx", {"Lease Review": {
        "A5": "0999001-001", "B5": "Ana Souza",
        "A6": "0999001-002", "B6": "Bruno Lima"}})
    tpl = _livro(tmp_path / "tpl.xlsx", {"Lease Review": {}})
    ia = _livro(tmp_path / "ia.xlsx", {"Lease Review": {
        "A9": "0999001-002", "B9": "Bruno Lima",
        "A10": "0999001-001", "B10": "Ana Souza"}})
    ledger = tmp_path / "ledger.json"
    ledger.write_text(json.dumps({"celulas": [
        {"aba": "Lease Review", "celula": c} for c in
        ("A9", "B9", "A10", "B10")], "brancos": []}), encoding="utf-8")
    got = compare(ia, ledger, hand=mao, template=tpl, aliases={}).to_dict()
    assert got["resultado"]["concordou"] == 4
    assert got["resultado"]["divergiu"] == 0
    assert got["linhas_alinhadas_por_identidade"] == 2


def test_identidade_repetida_pareia_em_ordem_e_o_resto_fica_calado(
        tmp_path: Path):
    """A mão repete a lease por trato; a IA tem uma linha. Uma pareia, a outra
    fica como silêncio — nunca um pareamento forçado."""
    from src.examinador.compare import compare

    mao = _livro(tmp_path / "mao.xlsx", {"Worksheet": {
        "A5": "0999001-001", "B5": "tract 1",
        "A6": "0999001-001", "B6": "tract 2"}})
    tpl = _livro(tmp_path / "tpl.xlsx", {"Worksheet": {}})
    ia = _livro(tmp_path / "ia.xlsx", {"Worksheet": {
        "A2": "0999001-001", "B2": "tract 1"}})
    ledger = tmp_path / "ledger.json"
    ledger.write_text(json.dumps({"celulas": [
        {"aba": "Worksheet", "celula": "A2"},
        {"aba": "Worksheet", "celula": "B2"}], "brancos": []}),
        encoding="utf-8")
    got = compare(ia, ledger, hand=mao, template=tpl, aliases={}).to_dict()
    assert got["resultado"]["concordou"] == 2
    assert got["resultado"]["calou_onde_a_mao_escreveu"] == 2


# --- cli: uma sessão com modelo roteirizado ---------------------------------

@dataclass
class _Uso:
    input_tokens: int = 10
    output_tokens: int = 5
    cache_creation_input_tokens: int = 0
    cache_read_input_tokens: int = 0


@dataclass
class _Msg:
    content: list[Any]
    usage: _Uso = field(default_factory=_Uso)


@dataclass
class _Use:
    name: str
    input: dict[str, Any]
    id: str = "u1"
    type: str = "tool_use"


@dataclass
class _Texto:
    text: str
    type: str = "text"


class _ClientRoteirizado:
    def __init__(self, roteiro: list[list[Any]]) -> None:
        self.roteiro = roteiro
        self.messages = self

    def create(self, **kw: Any) -> _Msg:
        return _Msg(content=self.roteiro.pop(0))


def test_o_comando_unidade_fecha_com_os_tres_numeros(tmp_path: Path,
                                                     monkeypatch, capsys):
    """Uma sessão roteirizada: censo, um campo, um pedido, uma pergunta. O
    fecho imprime exige/escrevi/branco/pendente e os arquivos existem."""
    import shutil

    from src.minorh import cli

    monkeypatch.setenv("PLM_PAGE_CACHE", str(tmp_path / "cache"))
    monkeypatch.setattr(cli, "ULTIMA", tmp_path / "ultima.txt")
    unidade = tmp_path / "unidade"
    origem = Path("tests/fixtures/unidade_sintetica")
    shutil.copytree(origem / "FILES", unidade / "FILES")
    shutil.copyfile(origem / "template.xlsx", unidade / "template.xlsx")
    # triagem completa pré-existente: a fase do Triador (F2) não roda e o
    # roteiro abaixo alimenta só o Examinador.
    from src.examinador.cache import PageCache as _PC
    from src.examinador.corpus import Corpus as _Corpus
    nomes = _Corpus(unidade / "FILES", cache=_PC(tmp_path / "cache")).names()
    (unidade / "_minorh").mkdir(parents=True, exist_ok=True)
    (unidade / "_minorh" / "triagem.json").write_text(json.dumps(
        {n: {"rotulo": "doc", "resumo": "fixture", "confianca": "alta",
             "quando": ""} for n in nomes}), encoding="utf-8")

    client = _ClientRoteirizado([
        [_Use("declarar_censo", {"exige": 35, "resumo": "5 leases x 7"})],
        [_Use("gravar_campo", {"entidade": "lease 0999001-001",
                               "campo": "royalty", "valor": "1/8",
                               "documento": "0999001-001 (ana souza) lease.pdf",
                               "pagina": "1", "trecho": "one-eighth (1/8)",
                               "confianca": "alta"})],
        [_Use("pedir_documento", {"titulo": "Lease 0999001-005",
                                  "o_que": "instrumento", "onde": "Seller",
                                  "metodo": "requisicao",
                                  "campos": [{"entidade": "lease 0999001-005",
                                              "campo": "lessor"}]})],
        [_Use("perguntar", {"pergunta": "Qual registro vale?",
                            "candidatos": ["Pg. 3", "Pg. 33"],
                            "o_que_desempataria": "cartorio"})],
        [_Texto("aberto")],
    ])
    cli.cmd_unidade(unidade, client=client)
    saida = capsys.readouterr().out
    assert "exige 35" in saida
    assert "escrevi 1 (com citação)" in saida
    assert "pendente de documento 1" in saida
    assert (unidade / "PEDIDOS.md").exists()
    assert (unidade / "PERGUNTAS.md").exists()
    # e o estado sobreviveu para o proximo comando
    cli.cmd_estado(unidade)
    assert "exige 35" in capsys.readouterr().out
