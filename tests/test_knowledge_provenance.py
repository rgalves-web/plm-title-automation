"""PR-9 — write containment and the provenance whitelist.

The lesson fixture is the PR-8 artefact itself: 759 non-empty cells outside the
trail, seven tabs of another operator's sample unit, approved by a blacklist of
six strings. Fixture bytes live outside git (invariant 7/14); when absent, the
tests SKIP counted, never pass silent.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

import openpyxl
import pytest
import yaml

from src.knowledge.clean_output import AUDIT_TRAIL_SHEET
from src.knowledge.corpus_reader import CorpusIndex, CorpusReadOnly
from src.knowledge.kit_builder import (ABSENCE_TOKENS, AbsenceNotDerived,
                                       CellWrite, DERIVED_EMPTY)
from src.knowledge.template_structure import (STRUCTURE_YAML, SAMPLE_VOCABULARY,
                                              StructureError,
                                              check_checklist_vs_artifact,
                                              check_provenance, load_structure,
                                              tabs_populated)
from src.knowledge.write_roots import (WriteOutsideRoots, active_scope,
                                       record_write, write_scope)

FIXTURE = Path("tests/fixtures/ddr_josh_b_pr8_provenance_defect.xlsx")
FIXTURE_CLAIM = Path("tests/fixtures/ddr_josh_b_pr8_provenance_defect.yaml")
REBUILT = Path("data/units/_josh_b_replay/KIT/DDR - Josh B Unit.xlsx")


# --- item 0.1: write-roots --------------------------------------------------

def test_a_write_outside_the_declared_roots_raises(tmp_path):
    inside, outside = tmp_path / "round", tmp_path / "elsewhere"
    with write_scope(inside):
        record_write(inside / "ok.txt")
        with pytest.raises(WriteOutsideRoots, match="declared"):
            record_write(outside / "stray.txt")


def test_the_manifest_lists_every_recorded_write_even_on_failure(tmp_path):
    root = tmp_path / "round"
    with pytest.raises(RuntimeError):
        with write_scope(root) as scope:
            record_write(root / "a.txt")
            raise RuntimeError("the round crashed")
    manifest = json.loads((root / "write_manifest.json").read_text())
    assert any(w.endswith("a.txt") for w in manifest["written"])


def test_outside_a_scope_record_write_only_resolves(tmp_path):
    assert active_scope() is None
    assert record_write(tmp_path / "x.txt") == (tmp_path / "x.txt").resolve()


def test_the_frozen_guard_is_the_chokepoint():
    import inspect

    from src.flow import deliverable_io
    assert "record_write" in inspect.getsource(deliverable_io.assert_not_frozen)


# --- item 0.2: the corpus is read in place ----------------------------------

def _make_corpus(tmp_path: Path) -> Path:
    import zipfile
    files = tmp_path / "FILES"
    files.mkdir()
    with zipfile.ZipFile(files / "0220604-002 (X).zip", "w") as zf:
        zf.writestr("0220604-002 (X)/LEASE.TIF", b"lease bytes")
    (files / "LOOSE TO.PDF").write_bytes(b"opinion bytes")
    return files


def test_a_logical_corpus_path_resolves_into_the_original_zip(tmp_path):
    idx = CorpusIndex(_make_corpus(tmp_path))
    old_style = r"data\units\_josh_b_replay\corpus\0220604-002 (X)\LEASE.TIF"
    assert idx.exists(old_style)
    assert idx.read_bytes(old_style) == b"lease bytes"
    assert idx.read_bytes("LOOSE TO.PDF") == b"opinion bytes"


def test_materialise_refuses_any_destination_outside_work(tmp_path):
    idx = CorpusIndex(_make_corpus(tmp_path))
    with pytest.raises(CorpusReadOnly, match="_work"):
        idx.materialise("0220604-002 (X)/LEASE.TIF", tmp_path / "standing_copy")
    out = idx.materialise("0220604-002 (X)/LEASE.TIF", tmp_path / "_work")
    assert out.read_bytes() == b"lease bytes"


# --- item 2: the structure baseline -----------------------------------------

def test_the_baseline_exists_and_blesses_no_sample_value():
    data = yaml.safe_load(STRUCTURE_YAML.read_text(encoding="utf-8"))
    assert data["sheets"], "empty baseline"
    for sheet, cells in data["sheets"].items():
        for coord, value in cells.items():
            for marker in SAMPLE_VOCABULARY:
                assert marker.lower() not in str(value).lower(), \
                    f"{sheet}!{coord} blesses sample value {marker!r}"


def test_every_expected_sheet_has_a_rule_and_counts_for_review():
    data = yaml.safe_load(STRUCTURE_YAML.read_text(encoding="utf-8"))
    from src.knowledge.kit_builder import EXPECTED_SHEETS
    assert set(data["sheets"]) == EXPECTED_SHEETS - {AUDIT_TRAIL_SHEET}
    assert all("structure" in c and "sample_stripped" in c
               for c in data["counts"].values())


# --- items 3+4: skeleton and the whitelist gate ------------------------------

def _mini_structure() -> dict:
    return {"Summary": {"A1": "Due Diligence Report", "A2": "Client:"}}


def _book(tmp_path: Path, cells: dict[str, str], trail: list[tuple] = ()) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Summary")
    for coord, value in cells.items():
        ws[coord] = value
    tr = wb.create_sheet(AUDIT_TRAIL_SHEET)
    tr.append(["SHEET", "CELL", "FIELD", "VALUE", "CITED TO", "REASON IF BLANK"])
    for row in trail:
        tr.append(list(row))
    path = tmp_path / "ddr.xlsx"
    wb.save(path)
    return path


def test_a_trailed_value_and_declared_structure_pass(tmp_path):
    path = _book(tmp_path,
                 {"A1": "Due Diligence Report", "A2": "Client:", "B2": "EOG"},
                 trail=[("Summary", "B2", "client", "EOG", "cite", "")])
    report = check_provenance(path, structure=_mini_structure())
    assert not report.failed and report.orphans == 0


def test_an_untrailed_cell_is_an_orphan_and_fails(tmp_path):
    path = _book(tmp_path, {"A1": "Due Diligence Report",
                            "B5": "Top of Bone Spring"})
    report = check_provenance(path, structure=_mini_structure())
    assert report.failed and report.orphans == 1
    assert any("Bone Spring" in f.detail for f in report.findings)


def test_structure_with_a_different_value_fails(tmp_path):
    path = _book(tmp_path, {"A2": "Cliente:"})
    report = check_provenance(path, structure=_mini_structure())
    assert any(f.kind == "structure_altered" for f in report.findings)


def test_the_trail_sheet_itself_needs_no_provenance(tmp_path):
    path = _book(tmp_path, {"A1": "Due Diligence Report"},
                 trail=[("Summary", "Z9", "x", "v", "c", "")])
    assert not check_provenance(path, structure=_mini_structure()).failed


@pytest.mark.skipif(not FIXTURE.exists(),
                    reason="fixture bytes absent — COUNTED SKIP (invariant 14)")
def test_the_pr8_artefact_fails_the_whitelist_citing_the_sample_unit():
    claim = yaml.safe_load(FIXTURE_CLAIM.read_text(encoding="utf-8"))
    digest = hashlib.sha256(FIXTURE.read_bytes()).hexdigest()
    assert digest == claim["files"][0]["sha256"], "fixture bytes drifted"
    report = check_provenance(FIXTURE, max_report=1200)
    assert report.failed
    assert report.orphans > 400        # 525 measured; ~759 minus legitimate structure
    blob = " ".join(str(f) for f in report.findings)
    assert "Bone Spring" in blob and "Ma Chona Adaya" in blob


@pytest.mark.skipif(not REBUILT.exists(), reason="cold kit not built here")
def test_the_rebuilt_artefact_has_zero_orphans():
    report = check_provenance(REBUILT)
    assert not report.failed and report.orphans == 0


def test_the_writer_calls_the_whitelist_gate():
    import inspect

    from src.knowledge import kit_builder
    assert "assert_provenance" in inspect.getsource(kit_builder.write_ddr)


# --- item 5: the checklist is derived from the bytes -------------------------

def test_a_checklist_yes_that_the_artefact_contradicts_fails(tmp_path):
    ddr = _book(tmp_path, {"A1": "Due Diligence Report"})      # nothing trailed
    chk = openpyxl.Workbook()
    ws = chk.active
    ws.cell(row=6, column=4).value = "Yes"
    chk_path = tmp_path / "chk.xlsx"
    chk.save(chk_path)
    report = check_checklist_vs_artifact(chk_path, ddr, {6: "Summary"},
                                         structure=_mini_structure())
    assert any(f.kind == "checklist_contradiction" for f in report.findings)


def test_summary_populated_requires_the_named_fields_not_one_cell(tmp_path):
    ddr = _book(tmp_path, {"A9": "Date Submitted:", "B9": "2026-08-11"},
                trail=[("Summary", "B9", "date_submitted", "2026-08-11", "c", "")])
    populated = tabs_populated(ddr, structure=_mini_structure())
    assert populated["Summary"] is False       # one field of eight is not populated


# --- item 6: absence is derived ----------------------------------------------

@pytest.mark.parametrize("token", sorted(ABSENCE_TOKENS))
def test_an_absence_token_without_a_computed_empty_set_refuses_to_write(token):
    with pytest.raises(AbsenceNotDerived, match="negative conclusion"):
        CellWrite("Summary", 10, "B", "potential_defects", token, "citation")


def test_an_absence_token_with_the_derived_marker_writes(tmp_path):
    c = CellWrite("Summary", 10, "B", "potential_defects", "--",
                  f"{DERIVED_EMPTY}: gap-hunter returned no findings")
    assert c.value == "--"


def test_ordinary_values_are_untouched():
    assert CellWrite("Summary", 2, "B", "client", "EOG Resources").value
