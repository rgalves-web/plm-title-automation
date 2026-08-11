"""`as_of` is replay-only — the ship path can never pin the date.

Pinning the date is legitimate for ONE purpose: re-validating an
already-shipped deliverable without the submission-date rule manufacturing a
FAIL months later. It must never reach a shipping run, because a shipping run
judging against a pinned past date would silently defeat the S8 freshness gate
("Date Submitted" is filled at packaging; a stale date is a structural FAIL).

These tests are the fence: default None everywhere, absent from the ship path
and from config, and loudly announced whenever it IS used.
"""
from __future__ import annotations

import inspect
from datetime import date

import pytest

from src.dd.ddr_validator import run_final_validation, validate_ddr
from src.flow.config import Config
from src.flow.runner import revalidate_from_state


def _signature_default(fn, name):
    return inspect.signature(fn).parameters[name].default


# --- the fence: replay entry points default to None -------------------------

def test_every_as_of_entry_point_defaults_to_none():
    assert _signature_default(run_final_validation, "as_of") is None
    assert _signature_default(revalidate_from_state, "as_of") is None
    assert _signature_default(validate_ddr, "today") is None


# --- the fence: the ship path cannot pin a date -----------------------------

def test_the_packager_has_no_as_of_parameter():
    from src.dd import submission_packager as pkg
    for fn in (pkg.package_unit_submission, pkg.build_submission_plan,
               pkg.build_plan_from_unit_folder, pkg.stage_submission):
        assert "as_of" not in inspect.signature(fn).parameters, fn.__name__


def test_the_ddr_writer_has_no_as_of_parameter():
    from src.dd.ddr_writer import fill_ddr_template
    assert "as_of" not in inspect.signature(fill_ddr_template).parameters


def test_the_workbench_service_never_passes_as_of():
    # The UI drives gates and ship; neither may pin a date.
    import src.webapp.service as service
    source = inspect.getsource(service)
    assert "as_of" not in source
    for name in ("run_structural", "start_source_truth", "ship"):
        assert "as_of" not in inspect.signature(getattr(
            service.WorkbenchService, name)).parameters


def test_as_of_is_not_a_config_key():
    # It is a per-call replay argument, never ambient configuration.
    assert "as_of" not in Config().to_dict()
    assert not hasattr(Config(), "as_of")


# --- the announcement: a replayed verdict is never mistaken for a ship one ---

def test_a_replayed_report_announces_REPLAY_in_summary_and_sidecar(tmp_path):
    import openpyxl
    from src.dd.ddr_validator import console_summary, write_validation_sidecar

    ddr = tmp_path / "DDR.xlsx"
    openpyxl.Workbook().save(ddr)

    replay = run_final_validation(ddr, run_source_truth_check=False,
                                  as_of=date(2026, 7, 28))
    assert replay.replay_as_of == date(2026, 7, 28)
    assert replay.summary().startswith("REPLAY — as_of=2026-07-28")
    assert "REPLAY" in console_summary(replay)
    sidecar = write_validation_sidecar(replay, ddr)
    assert "REPLAY — as_of=2026-07-28" in sidecar.read_text(encoding="utf-8")


def test_a_ship_time_report_says_nothing_about_replay(tmp_path):
    import openpyxl
    from src.dd.ddr_validator import console_summary

    ddr = tmp_path / "DDR.xlsx"
    openpyxl.Workbook().save(ddr)

    shipping = run_final_validation(ddr, run_source_truth_check=False)   # as_of=None
    assert shipping.replay_as_of is None
    assert "REPLAY" not in shipping.summary()
    assert "REPLAY" not in console_summary(shipping)


# --- the reason the fence exists --------------------------------------------

def test_a_stale_submission_date_still_fails_at_ship_time():
    # S8 freshness: at ship time the rule judges against the real today, so a
    # date baked in earlier FAILs. Only an explicit replay may pin it.
    from src.dd.ddr_validator import FAIL, check_submission_date
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Date Submitted:"
    ws["B1"] = "07/28/2026"
    assert [f for f in check_submission_date(wb, today=date(2026, 8, 11))
            if f.severity == FAIL]
    assert not [f for f in check_submission_date(wb, today=date(2026, 7, 28))
                if f.severity == FAIL]
