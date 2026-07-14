"""Tests for source file matching, PDF extraction integration, and QA hard gates."""
from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

import pandas as pd
import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

from src.pdf_reader import (
    EXTRACTION_ERROR,
    LOW_CONFIDENCE_TEXT,
    OCR_REQUIRED,
    TEXT_EXTRACTED,
    apply_extraction_to_record,
    extract_document,
)
from src.pipeline import run_full_pipeline_from_dataframe
from src.qa import run_qa
from src.schemas import ClassificationStatus, LeaseStatus, ProjectMetadata, SourceDocument
from src.source_matcher import match_source_files


fitz = pytest.importorskip("fitz")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _write_pdf(path: Path, text: str) -> Path:
    doc = fitz.open()
    page = doc.new_page()
    if text:
        page.insert_textbox((72, 72, 540, 720), text, fontsize=11)
    doc.save(path)
    doc.close()
    return path


def _write_blank_pdf(path: Path) -> Path:
    return _write_pdf(path, "")


def _zip_with_pdf(zip_path: Path, pdf_path: Path, member_name: str | None = None) -> Path:
    member = member_name or pdf_path.name
    with ZipFile(zip_path, "w") as archive:
        archive.write(pdf_path, member)
    return zip_path


OFFICIAL_HEADERS = [
    "NO.", "INSTRUMENT", "INSTRUMENT NUMBER", "EFFECTIVE DATE",
    "INSTRUMENT DATE", "RECORDED DATE", "GRANTOR", "GRANTEE",
    "ACREAGE", "CONVEYANCE", "COMMENTS",
]


def _create_template(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Runsheet"
    for col, header in enumerate(OFFICIAL_HEADERS, start=1):
        cell = ws.cell(row=1, column=col)
        cell.value = header
        cell.font = Font(bold=True)
    wb.create_sheet("Notes")
    wb.save(path)


def _record(
    status=ClassificationStatus.INCLUDE_FULL,
    document_type="Patent",
    source_document_id="DOC-000001",
    **kwargs,
) -> SourceDocument:
    return SourceDocument(
        source_document_id=source_document_id,
        source_row_number=1,
        document_type=document_type,
        status=status,
        **kwargs,
    )


# ---------------------------------------------------------------------------
# Test 1: source file matched by original_file_name
# ---------------------------------------------------------------------------

def test_source_file_matched_by_original_file_name(tmp_path):
    pdf_path = _write_pdf(tmp_path / "patent_001.pdf", "Patent document text content here.")
    zip_path = _zip_with_pdf(tmp_path / "package.zip", pdf_path)

    record = _record(original_file_name="patent_001.pdf")
    match_source_files([record], document_package_path=zip_path)

    assert record.matched_source_file == "patent_001.pdf"
    assert record.source_match_method == "original_file_name"
    assert record.source_match_confidence == 1.0


# ---------------------------------------------------------------------------
# Test 2: source file matched by instrument_number
# ---------------------------------------------------------------------------

def test_source_file_matched_by_instrument_number(tmp_path):
    pdf_path = _write_pdf(tmp_path / "20240000001.pdf", "Mineral Deed content text here.")
    zip_path = _zip_with_pdf(tmp_path / "package.zip", pdf_path)

    record = _record(
        document_type="Mineral Deed",
        instrument_number="20240000001",
        original_file_name=None,
    )
    match_source_files([record], document_package_path=zip_path)

    assert record.matched_source_file == "20240000001.pdf"
    assert record.source_match_method == "instrument_number"
    assert record.source_match_confidence == 0.95


# ---------------------------------------------------------------------------
# Test 3: unmatched INCLUDE_FULL creates QA error
# ---------------------------------------------------------------------------

def test_unmatched_include_full_creates_qa_error():
    record = _record(status=ClassificationStatus.INCLUDE_FULL)
    # No matched_source_file set; source_verification_required=True

    exceptions = run_qa([record], expected_document_count=1, source_verification_required=True)

    error_codes = {e.code for e in exceptions}
    assert "SOURCE_FILE_NOT_MATCHED" in error_codes
    exc = next(e for e in exceptions if e.code == "SOURCE_FILE_NOT_MATCHED")
    assert exc.severity == "ERROR"
    assert exc.source_document_id == "DOC-000001"


# ---------------------------------------------------------------------------
# Test 4: unreadable PDF creates QA error
# ---------------------------------------------------------------------------

def test_unreadable_pdf_creates_qa_error():
    record = _record(
        status=ClassificationStatus.INCLUDE_FULL,
        pdf_extraction_status=EXTRACTION_ERROR,
        matched_source_file="broken.pdf",
    )

    exceptions = run_qa([record], expected_document_count=1)

    error_codes = {e.code for e in exceptions}
    assert "PDF_EXTRACTION_ERROR" in error_codes
    exc = next(e for e in exceptions if e.code == "PDF_EXTRACTION_ERROR")
    assert exc.severity == "ERROR"
    assert exc.source_document_id == "DOC-000001"


# ---------------------------------------------------------------------------
# Test 5: unreadable/blank document pushes record to REVIEW and emits QA WARNING
# ---------------------------------------------------------------------------

def test_blank_pdf_pushes_record_to_review_and_emits_warning(tmp_path):
    """A blank PDF must push the record to REVIEW and emit a QA WARNING.

    With Tesseract installed: blank page → OCR runs → LOW_CONFIDENCE_TEXT → REVIEW.
    Without Tesseract: OCR_REQUIRED → REVIEW.  Both paths must produce a WARNING.
    """
    pdf_path = _write_blank_pdf(tmp_path / "scanned.pdf")
    result = extract_document(pdf_path)
    assert result.extraction_status in {OCR_REQUIRED, LOW_CONFIDENCE_TEXT}

    record = SourceDocument(
        source_document_id="DOC-000001",
        source_row_number=1,
        document_type="Unknown",
    )
    apply_extraction_to_record(record, result)

    assert record.status == ClassificationStatus.REVIEW
    assert record.pdf_extraction_status in {OCR_REQUIRED, LOW_CONFIDENCE_TEXT}
    assert record.review_question

    # QA must emit a WARNING (OCR_REQUIRED or LOW_CONFIDENCE — either is correct)
    exceptions = run_qa([record], expected_document_count=1)
    warning_codes = {e.code for e in exceptions if e.severity == "WARNING"}
    assert warning_codes & {"PDF_OCR_REQUIRED", "PDF_LOW_CONFIDENCE"}


# ---------------------------------------------------------------------------
# Test 6: source verification fields stay in internal reports, not runsheet
# ---------------------------------------------------------------------------

def test_source_verification_fields_not_in_runsheet(tmp_path):
    pdf_path = _write_pdf(
        tmp_path / "patent_001.pdf",
        (
            "Patent. United States of America to John Smith. "
            "Section 1 Township 4S Range 64W, 160 acres."
        ),
    )
    zip_path = _zip_with_pdf(tmp_path / "docs.zip", pdf_path)
    template_path = tmp_path / "template.xlsx"
    _create_template(template_path)

    dataframe = pd.DataFrame(
        [
            {
                "Document Type": "Patent",
                "Instrument Number": "1001",
                "Instrument Date": "1901-01-01",
                "Recording Date": "1901-01-02",
                "Grantor": "United States of America",
                "Grantee": "John Smith",
                "Original File Name": "patent_001.pdf",
            }
        ]
    )

    result = run_full_pipeline_from_dataframe(
        dataframe=dataframe,
        template_path=template_path,
        metadata=ProjectMetadata(project_name="SV Test"),
        output_root=tmp_path / "output",
        document_package_path=zip_path,
    )

    from openpyxl import load_workbook
    runsheet_wb = load_workbook(result.artifacts["preliminary_runsheet"])
    sheet = runsheet_wb["Runsheet"]
    visible_headers = [sheet.cell(row=1, column=c).value for c in range(1, 12)]
    assert visible_headers == OFFICIAL_HEADERS
    assert sheet.max_column == 11

    internal_fields = {
        "source_document_id", "matched_source_file", "source_match_confidence",
        "source_match_method", "pdf_extraction_status", "pdf_extraction_confidence",
        "pdf_page_count", "pdf_text_preview", "possible_document_type_keywords",
        "possible_legal_description_snippets",
    }
    assert internal_fields.isdisjoint(set(visible_headers))


# ---------------------------------------------------------------------------
# Test 7: Export All produces 08_source_match_report.xlsx and
#         09_pdf_extraction_report.xlsx
# ---------------------------------------------------------------------------

def test_export_all_produces_source_match_and_pdf_extraction_reports(tmp_path):
    template_path = tmp_path / "template.xlsx"
    _create_template(template_path)

    dataframe = pd.DataFrame(
        [
            {
                "Document Type": "Patent",
                "Instrument Number": "1001",
                "Instrument Date": "1901-01-01",
                "Recording Date": "1901-01-02",
                "Grantor": "United States of America",
                "Grantee": "Owner One",
            }
        ]
    )

    result = run_full_pipeline_from_dataframe(
        dataframe=dataframe,
        template_path=template_path,
        metadata=ProjectMetadata(project_name="Report Test"),
        output_root=tmp_path / "output",
    )

    assert "source_match_report" in result.artifacts
    assert result.artifacts["source_match_report"].exists()
    assert result.artifacts["source_match_report"].name == "08_source_match_report.xlsx"

    assert "pdf_extraction_report" in result.artifacts
    assert result.artifacts["pdf_extraction_report"].exists()
    assert result.artifacts["pdf_extraction_report"].name == "09_pdf_extraction_report.xlsx"

    # Verify required columns are present in both reports
    required_cols = {
        "source_document_id", "document_type", "instrument_number",
        "status", "matched_source_file", "pdf_extraction_status",
        "review_question", "qa_exception_codes",
    }
    for artifact_key in ("source_match_report", "pdf_extraction_report"):
        df = pd.read_excel(result.artifacts[artifact_key])
        assert required_cols.issubset(set(df.columns)), (
            f"{artifact_key} missing columns: {required_cols - set(df.columns)}"
        )
