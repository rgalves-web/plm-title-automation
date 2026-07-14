import json
from datetime import date

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from scripts.generate_sample_data import generate_sample_dataset
from src.classifier import classify_documents
from src.dedupe import deduplicate_documents
from src.document_register import create_document_register
from src.excel_writer import (
    generate_runsheet,
    generate_runsheet_with_report,
    template_integrity_report_path,
    validate_template,
)
from src.intake import read_document_list
from src.pipeline import run_full_pipeline_from_dataframe
from src.schemas import ClassificationStatus, LeaseStatus, ProjectMetadata, SourceDocument


OFFICIAL_HEADERS = [
    "NO.",
    "INSTRUMENT",
    "INSTRUMENT NUMBER",
    "EFFECTIVE DATE",
    "INSTRUMENT DATE",
    "RECORDED DATE",
    "GRANTOR",
    "GRANTEE",
    "ACREAGE",
    "CONVEYANCE",
    "COMMENTS",
]
INTERNAL_HEADERS = {
    "Source Document ID",
    "Status",
    "Redline Reason",
    "Review Question",
    "Lease Status",
    "Duplicate Of",
}


def test_validate_template_and_write_official_plm_runsheet(tmp_path):
    paths = generate_sample_dataset(tmp_path)
    dataframe = read_document_list(paths["document_list"])
    records = create_document_register(dataframe)
    deduplicate_documents(records)
    classify_documents(records)

    check = validate_template(paths["template"])
    assert check.passed
    assert check.missing_columns == []

    output_path = generate_runsheet(
        paths["template"],
        records,
        tmp_path,
        project_name="Test Project",
    )
    workbook = load_workbook(output_path)
    sheet = workbook["Runsheet"]
    headers = [sheet.cell(row=1, column=column).value for column in range(1, 12)]

    assert headers == OFFICIAL_HEADERS
    assert INTERNAL_HEADERS.isdisjoint(set(headers))
    assert sheet.max_column == 11
    assert sheet["A2"].value == 1
    assert sheet["B2"].value == "Patent"
    assert sheet["C2"].value == "1/10\n1001"
    assert sheet["K3"].value and "EXPIRED" in sheet["K3"].value
    assert sheet["B4"].value == "Easement"
    # REDLINE row: identified by fill color, not by text in Comments
    assert "REDLINE" not in (sheet["K4"].value or "")
    assert sheet.cell(row=4, column=1).fill.fill_type == "solid"
    assert template_integrity_report_path(output_path).exists()


def test_template_preservation_and_integrity_report(tmp_path):
    template_path = tmp_path / "preservation_template.xlsx"
    _create_official_template(template_path, sheet_name="Arapahoe Custom", header_row=5)

    record = SourceDocument(
        source_document_id="DOC-000001",
        source_row_number=2,
        document_type="Patent",
        effective_date=date(1901, 1, 1),
        instrument_date=date(1901, 1, 2),
        recording_date=date(1901, 2, 1),
        grantor="United States of America",
        grantee="John Smith, a married man",
        book="1",
        page="10",
        instrument_number="1001",
        acreage="160",
        conveyance="Grant of Section 1 Township 4S Range 64W",
        comments="Preserve exact names.",
        status=ClassificationStatus.INCLUDE_FULL,
    )

    check = validate_template(template_path)
    assert check.passed
    assert check.header_row == 5
    assert check.sheet_name == "Arapahoe Custom"

    result = generate_runsheet_with_report(
        template_path,
        [record],
        tmp_path,
        project_name="Preservation",
    )

    workbook = load_workbook(result.output_path, data_only=False)
    sheet = workbook["Arapahoe Custom"]
    notes = workbook["Notes"]

    assert workbook.sheetnames == ["Arapahoe Custom", "Notes"]
    assert [sheet.cell(row=5, column=column).value for column in range(1, 12)] == OFFICIAL_HEADERS
    assert sheet.max_column == 11
    assert sheet.column_dimensions["A"].width == 28
    assert sheet.row_dimensions[6].height == 42
    assert sheet.row_dimensions[8].height == 31
    assert sheet["A6"].alignment.wrap_text is True
    assert sheet["A6"].border.left.style == "thin"
    assert sheet["E6"].number_format == "mm/dd/yyyy"
    assert notes["B2"].value == "=1+1"
    assert notes.column_dimensions["B"].width == 22

    with result.integrity_report_path.open("r", encoding="utf-8") as handle:
        report = json.load(handle)

    assert report["passed"] is True
    assert report["checks"]["template_file_copied_before_edit"] is True
    assert report["checks"]["sheet_names_preserved"] is True
    assert report["checks"]["column_widths_preserved"] is True
    assert report["checks"]["row_heights_outside_inserted_range_preserved"] is True
    assert report["checks"]["wrap_text_outside_inserted_range_preserved"] is True
    assert report["checks"]["borders_outside_inserted_range_preserved"] is True
    assert report["checks"]["number_formats_outside_inserted_range_preserved"] is True
    assert report["checks"]["formulas_outside_inserted_range_preserved"] is True


def test_official_plm_template_does_not_add_internal_audit_columns(tmp_path):
    template_path = tmp_path / "arapahoe_style_template.xlsx"
    _create_official_template(template_path, sheet_name="Runsheet - NE", header_row=2)

    records = [
        SourceDocument(
            source_document_id="DOC-000001",
            source_row_number=2,
            document_type="Patent",
            instrument_number="P123",
            grantor="United States of America",
            grantee="Owner One",
            instrument_date=date(1901, 1, 1),
            recording_date=date(1901, 1, 2),
            status=ClassificationStatus.INCLUDE_FULL,
        ),
        SourceDocument(
            source_document_id="DOC-000002",
            source_row_number=3,
            document_type="Correction Deed",
            instrument_number="C123",
            grantor="Grantor Name",
            grantee="Grantee Name",
            instrument_date=date(1910, 1, 1),
            recording_date=date(1910, 1, 2),
            status=ClassificationStatus.REVIEW,
            review_question="Confirm whether correction affects minerals.",
        ),
        SourceDocument(
            source_document_id="DOC-000003",
            source_row_number=4,
            document_type="Easement",
            instrument_number="E123",
            grantor="Grantor Name",
            grantee="Pipeline Co.",
            instrument_date=date(1920, 1, 1),
            recording_date=date(1920, 1, 2),
            comments="Non-title instrument.",
            status=ClassificationStatus.REDLINE,
            redline_reason="Non-title easement.",
        ),
        SourceDocument(
            source_document_id="DOC-000004",
            source_row_number=5,
            document_type="Patent",
            instrument_number="P123",
            grantor="United States of America",
            grantee="Owner One",
            status=ClassificationStatus.DUPLICATE,
            duplicate_of="DOC-000001",
        ),
    ]

    result = generate_runsheet_with_report(
        template_path,
        records,
        tmp_path,
        project_name="Arapahoe Style",
    )
    output_workbook = load_workbook(result.output_path)
    output_sheet = output_workbook["Runsheet - NE"]
    headers = [output_sheet.cell(row=2, column=column).value for column in range(1, 12)]

    assert headers == OFFICIAL_HEADERS
    assert INTERNAL_HEADERS.isdisjoint(set(headers))
    assert output_sheet.max_column == 11
    assert output_sheet["A3"].value == 1
    assert output_sheet["B3"].value == "Patent"
    assert output_sheet["B4"].value == "Correction Deed"
    # REVIEW row: status text must NOT appear in Comments column (belongs in QA log)
    k4 = output_sheet["K4"].value or ""
    assert "REVIEW REQUIRED" not in k4
    assert "Confirm whether correction affects minerals." not in k4
    # REVIEW row identified by yellow fill
    assert output_sheet.cell(row=4, column=1).fill.fill_type == "solid"
    assert output_sheet["B5"].value == "Easement"
    # REDLINE row: status text must NOT appear in Comments column
    k5 = output_sheet["K5"].value or ""
    assert "REDLINE" not in k5
    assert "Non-title easement." not in k5
    # REDLINE row identified by red fill
    assert output_sheet.cell(row=5, column=1).fill.fill_type == "solid"
    assert output_sheet["B6"].value is None


def test_internal_audit_fields_stay_in_internal_exports_only(tmp_path):
    template_path = tmp_path / "official_template.xlsx"
    _create_official_template(template_path)
    dataframe = pd.DataFrame(
        [
            {
                "Document Type": "Patent",
                "Instrument Number": "1001",
                "Instrument Date": "1901-01-01",
                "Recording Date": "1901-01-02",
                "Grantor": "United States of America",
                "Grantee": "Owner One",
            },
            {
                "Document Type": "Correction Deed",
                "Instrument Number": "2002",
                "Instrument Date": "1910-01-01",
                "Recording Date": "1910-01-02",
                "Grantor": "Owner One",
                "Grantee": "Owner Two",
            },
            {
                "Document Type": "Easement",
                "Instrument Number": "3003",
                "Instrument Date": "1920-01-01",
                "Recording Date": "1920-01-02",
                "Grantor": "Owner Two",
                "Grantee": "Pipeline Co.",
            },
            {
                "Document Type": "Patent",
                "Instrument Number": "1001",
                "Instrument Date": "1901-01-01",
                "Recording Date": "1901-01-02",
                "Grantor": "United States of America",
                "Grantee": "Owner One",
            },
        ]
    )

    result = run_full_pipeline_from_dataframe(
        dataframe=dataframe,
        template_path=template_path,
        metadata=ProjectMetadata(project_name="Audit Separation"),
        output_root=tmp_path,
    )

    runsheet = load_workbook(result.artifacts["preliminary_runsheet"])["Runsheet"]
    visible_headers = [runsheet.cell(row=1, column=column).value for column in range(1, 12)]
    assert visible_headers == OFFICIAL_HEADERS
    assert INTERNAL_HEADERS.isdisjoint(set(visible_headers))
    assert runsheet.max_column == 11
    assert [runsheet.cell(row=row, column=2).value for row in range(2, 5)] == [
        "Patent",
        "Correction Deed",
        "Easement",
    ]
    # Comments column is title-analysis only — no status strings
    assert "REVIEW REQUIRED" not in (runsheet["K3"].value or "")
    assert "REDLINE" not in (runsheet["K4"].value or "")
    # Visual distinction via fill color, not text
    assert runsheet.cell(row=3, column=1).fill.fill_type == "solid"  # REVIEW → yellow
    assert runsheet.cell(row=4, column=1).fill.fill_type == "solid"  # REDLINE → red
    assert runsheet["B6"].value is None

    internal_artifacts = [
        "document_register",
        "classification_results",
        "review_required",
        "redline_rows",
    ]
    for artifact_name in internal_artifacts:
        columns = list(pd.read_excel(result.artifacts[artifact_name]).columns)
        assert "source_document_id" in columns
        assert "status" in columns

    classification_results = pd.read_excel(result.artifacts["classification_results"])
    assert "DUPLICATE" in set(classification_results["status"])
    duplicate_rows = classification_results[classification_results["status"] == "DUPLICATE"]
    assert duplicate_rows["duplicate_of"].iloc[0] == "DOC-000001"

    review_required = pd.read_excel(result.artifacts["review_required"])
    assert "review_question" in review_required.columns
    assert not review_required.empty

    redline_rows = pd.read_excel(result.artifacts["redline_rows"])
    assert "redline_reason" in redline_rows.columns
    assert not redline_rows.empty

    qa_log_columns = list(pd.read_excel(result.artifacts["qa_exception_log"]).columns)
    assert {"severity", "code", "message", "source_document_id", "field"}.issubset(
        qa_log_columns
    )


def _create_official_template(
    path,
    sheet_name="Runsheet",
    header_row=1,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    notes = workbook.create_sheet("Notes")

    for column_number, header in enumerate(OFFICIAL_HEADERS, start=1):
        cell = sheet.cell(row=header_row, column=column_number)
        cell.value = header
        cell.font = Font(bold=True)

    template_row = header_row + 1
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = PatternFill("solid", fgColor="E2F0D9")
    for row_number in (header_row, template_row, template_row + 1):
        for column_number in range(1, len(OFFICIAL_HEADERS) + 1):
            cell = sheet.cell(row=row_number, column=column_number)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = border
            if row_number == template_row:
                cell.fill = copy_fill(fill)
    sheet.cell(row=template_row, column=5).number_format = "mm/dd/yyyy"
    sheet.column_dimensions["A"].width = 28
    sheet.row_dimensions[header_row].height = 24
    sheet.row_dimensions[template_row].height = 42
    sheet.row_dimensions[template_row + 1].height = 31

    notes["A1"] = "Preserve this sheet"
    notes["B2"] = "=1+1"
    notes.column_dimensions["B"].width = 22
    notes.row_dimensions[2].height = 33
    workbook.save(path)


def copy_fill(fill):
    return PatternFill(
        fill_type=fill.fill_type,
        fgColor=fill.fgColor.rgb,
        bgColor=fill.bgColor.rgb,
    )
