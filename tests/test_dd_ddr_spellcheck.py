"""Spell-check pass (PLM checklist item 11): reports, never corrects."""
from __future__ import annotations

import openpyxl
import pytest

from src.dd.ddr_spellcheck import spell_check_workbook, spellcheck_report

pytest.importorskip("spellchecker")


def _wb(cells: dict[str, object], sheet="Summary"):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(sheet)
    for coord, value in cells.items():
        ws[coord] = value
    return wb


def test_a_real_misspelling_is_reported():
    found = spell_check_workbook(_wb({"B2": "the lease is recieved and filed"}))
    assert any(f.word == "recieved" for f in found)
    assert any(f.suggestion == "received" for f in found)


def test_proper_nouns_are_never_judged():
    # There is no dictionary of Montague County landowners.
    found = spell_check_workbook(_wb({"B2": "Hrabovsky, Brumbaugh, Croxton and Sanderfer"}))
    assert found == []


def test_domain_and_legal_terms_are_not_misspellings():
    found = spell_check_workbook(_wb({
        "B2": "lessor pooled leasehold orri nri wellbore pugh casinghead per annum incl. cert"}))
    assert found == []


def test_identifiers_and_numbers_are_skipped():
    found = spell_check_workbook(_wb({"B2": "0220344-013 K1400813 A-122 42-337-34216 140.9963"}))
    assert found == []


def test_a_word_used_by_the_source_document_is_correct_by_definition():
    text = "the lands are described as the Calhoun CSL abstracted therein"
    found = spell_check_workbook(_wb({"B2": "abstracted therein"}), source_text=text)
    assert found == []


def test_the_audit_trail_is_not_spell_checked():
    wb = _wb({"B2": "clean"})
    aud = wb.create_sheet("DD_Audit_Trail")
    aud["A2"] = "recieved a garbled ocr strng here"
    assert spell_check_workbook(wb) == []


def test_findings_carry_cell_word_suggestion_and_context():
    found = spell_check_workbook(_wb({"B2": "the royalty was recieved late"}))
    f = next(x for x in found if x.word == "recieved")
    assert f.sheet == "Summary" and f.cell == "B2"
    assert "royalty" in f.context and "recieved" in f.describe()


def test_a_word_is_reported_once_per_sheet_not_once_per_cell():
    found = spell_check_workbook(_wb({"B2": "recieved", "B3": "recieved", "B4": "recieved"}))
    assert len([f for f in found if f.word == "recieved"]) == 1


def test_report_says_it_did_not_correct_anything():
    text = spellcheck_report(spell_check_workbook(_wb({"B2": "recieved"})))
    assert "NOT auto-corrected" in text


def test_clean_workbook_reports_clean():
    assert "no unrecognised words" in spellcheck_report([])
