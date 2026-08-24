"""Gera a unidade sintética — material de teste do repo, NÃO dado de cliente.

Todo documento carrega o rótulo "DOCUMENTO SINTÉTICO — FIXTURE DE TESTE" na
primeira linha. A unidade: JOSH SINTÉTICO UNIT, 100 acres fictícios em
"Palo Duro County" (condado que não existe), 2 tratos, 5 leases declaradas no
Exhibit "A" da DPU — das quais só 3 têm o instrumento em FILES/. Uma das duas
retidas espera em `retidos/` para o teste do `continuar`.

Iscas deliberadas, para provar comportamento:
  - a lease -003 tem registros DIVERGENTES entre o Exhibit "A" (Vol. 100,
    Pg. 3) e o próprio instrumento (Vol. 100, Pg. 33) — deve virar PERGUNTA;
  - as leases -004 e -005 são declaradas e ausentes — devem virar PEDIDO;
  - o parecer traz um requirement (probate de Ana Souza) — deve chegar à
    memória e ao DDR.

Idempotente: rodar de novo reescreve tudo igual.
"""
from __future__ import annotations

import csv
import io
from pathlib import Path

AQUI = Path(__file__).parent
ROTULO = "DOCUMENTO SINTETICO - FIXTURE DE TESTE - NAO E DADO DE CLIENTE"

LEASES = {
    "0999001-001": ("Ana Souza", "January 5, 2011", "Vol. 100, Pg. 1"),
    "0999001-002": ("Bruno Lima", "January 12, 2011", "Vol. 100, Pg. 2"),
    "0999001-003": ("Carla Nunes", "February 2, 2011", "Vol. 100, Pg. 3"),
    "0999001-004": ("Daniel Prado", "February 20, 2011", "Vol. 100, Pg. 4"),
    "0999001-005": ("Elisa Ramos", "March 1, 2011", "Vol. 100, Pg. 5"),
}
# O instrumento da -003 imprime OUTRO registro que o Exhibit "A" — a isca da
# pergunta confirmar-no-cartório.
RECORDING_NO_INSTRUMENTO = {"0999001-003": "Vol. 100, Pg. 33"}
PRESENTES = ("0999001-001", "0999001-002", "0999001-003")
RETIDA = "0999001-004"          # vai para retidos/, entra no teste do continuar


def _pdf(paginas: list[str]) -> bytes:
    import fitz
    doc = fitz.open()
    for texto in paginas:
        page = doc.new_page()
        y = 60
        for linha in texto.splitlines():
            page.insert_text((50, y), linha[:110], fontsize=9)
            y += 13
    data = doc.tobytes()
    doc.close()
    return data


def _lease_pdf(numero: str) -> bytes:
    lessor, data, _ = LEASES[numero]
    recording = RECORDING_NO_INSTRUMENTO.get(numero, LEASES[numero][2])
    return _pdf([f"""{ROTULO}

OIL AND GAS LEASE

EOG SINTETICA LEASE NO. {numero}

THIS AGREEMENT made this {data}, between {lessor}, as Lessor,
and EOG SINTETICA RESOURCES INC., as Lessee.

Lessor leases to Lessee the lands in the JOSH SINTETICO UNIT area,
Palo Duro County, Texas, for the purpose of producing oil and gas.

ROYALTY: Lessor shall receive one-eighth (1/8) of all oil and gas produced.

TERM: This lease shall remain in force for three (3) years from this date.

RECORDED: {recording}, Official Records, Palo Duro County, Texas.
"""])


def _dpu_pdf() -> bytes:
    exhibit_a = "\n".join(
        f"  {n} | {lessor} | EOG SINTETICA RESOURCES INC. | {data} | {rec}"
        for n, (lessor, data, rec) in LEASES.items())
    return _pdf([
        f"""{ROTULO}

DECLARATION OF POOLED UNIT

JOSH SINTETICO UNIT
Palo Duro County, Texas

EOG SINTETICA RESOURCES INC., as Operator, declares the pooled unit named
JOSH SINTETICO UNIT, comprising 100.00 acres, more or less, effective
March 15, 2011, recorded at Vol. 200, Pg. 50, Official Records,
Palo Duro County, Texas.

The unit is formed for the production of oil and gas from the Sintetica
Shale formation and no other.
""",
        f"""{ROTULO}

EXHIBIT "A" TO DECLARATION OF POOLED UNIT - JOSH SINTETICO UNIT

Schedule of Oil and Gas Leases pooled into the unit:

  LEASE NO. | LESSOR | LESSEE | DATE | RECORDING

{exhibit_a}

Five (5) leases in total are pooled into the JOSH SINTETICO UNIT.
""",
        f"""{ROTULO}

EXHIBIT "C" TO DECLARATION OF POOLED UNIT - JOSH SINTETICO UNIT

Schedule of Tracts:

  TRACT 1 - 40.00 acres out of the Sintetica Survey, Abstract 1 (A-1),
            covered by Leases 0999001-001, 0999001-002 and 0999001-003.
  TRACT 2 - 60.00 acres out of the Sintetica Survey, Abstract 2 (A-2),
            covered by Leases 0999001-004 and 0999001-005.

Total: 100.00 acres, two (2) tracts.
"""])


def _parecer_pdf() -> bytes:
    return _pdf([f"""{ROTULO}

ORIGINAL TITLE OPINION

RE: JOSH SINTETICO UNIT - 100.00 acres, Sintetica Survey,
    Palo Duro County, Texas

Escritorio Sintetico de Pareceres LLP examined title to the captioned lands
as of June 1, 2011 (the certification date), and finds title vested as
reflected in the ownership schedules of this opinion.

REQUIREMENT NO. 1: Obtain and record certified copies of the probate
proceedings of the Estate of Ana Souza, whose lease 0999001-001 depends on
the devise of the minerals under Tract 1. Until probate is of record, the
lease is subject to the claims of the heirs.

No further requirements are made.
"""])


def _poco_csv() -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow([ROTULO])
    writer.writerow(["API", "WELL NAME", "OPERATOR", "STATUS",
                     "TOTAL DEPTH", "COMPLETION DATE"])
    writer.writerow(["42-999-00001", "JOSH SINTETICO 1H",
                     "EOG SINTETICA RESOURCES INC.", "PRODUCING",
                     "7500", "2012-01-20"])
    return buf.getvalue().encode("utf-8")


def _template_xlsx(path: Path) -> None:
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    resumo = wb.create_sheet("Summary")
    resumo["A1"] = "Due Diligence Report (TEMPLATE SINTETICO)"
    resumo["A2"], resumo["B2"] = "Unidade:", "***"
    resumo["A3"], resumo["B3"] = "Condado:", "***"
    resumo["A4"], resumo["B4"] = "Acres:", "***"
    resumo["A5"], resumo["B5"] = "Operador:", "***"

    lr = wb.create_sheet("Lease Review")
    for col, titulo in zip("ABCDEFG", ["LEASE NO.", "LESSOR", "LESSEE",
                                       "DATA", "REGISTRO", "ROYALTY",
                                       "TERMO"]):
        lr[f"{col}1"] = titulo
    # uma linha de amostra de OUTRA unidade — a sobra que o ship deve limpar
    lr["A2"], lr["B2"] = "0888888-001", "Amostra De Outra Unidade"

    ow = wb.create_sheet("Ownership - Worksheet")
    for col, titulo in zip("ABCD", ["TRACT", "ACRES", "LEASES", "SURVEY"]):
        ow[f"{col}1"] = titulo
    # A "célula azul" do template real: um total por fórmula. A caneta recusa
    # escrita aqui (S02) — a fixture existe para exercitar essa recusa.
    ow["E1"] = "TOTAL ACRES"
    ow["E2"] = "=SUM(B2:B99)"

    rs = wb.create_sheet("Runsheet")
    for col, titulo in zip("ABCDE", ["NO.", "INSTRUMENTO", "REGISTRO",
                                     "DATA", "PARTES"]):
        rs[f"{col}1"] = titulo

    po = wb.create_sheet("Wells")
    for col, titulo in zip("ABCDEF", ["API", "NOME", "OPERADOR", "STATUS",
                                      "PROFUNDIDADE", "COMPLETACAO"]):
        po[f"{col}1"] = titulo

    cur = wb.create_sheet("Curative")
    for col, titulo in zip("ABC", ["NO.", "REQUIREMENT", "FONTE"]):
        cur[f"{col}1"] = titulo

    wb.save(path)
    wb.close()


def gerar() -> Path:
    files = AQUI / "FILES"
    retidos = AQUI / "retidos"
    files.mkdir(exist_ok=True)
    retidos.mkdir(exist_ok=True)

    (files / "dpu - josh sintetico unit - rec.pdf").write_bytes(_dpu_pdf())
    for numero in PRESENTES:
        lessor = LEASES[numero][0].lower().replace(" ", " ")
        (files / f"{numero} ({lessor}) lease.pdf").write_bytes(
            _lease_pdf(numero))
    (files / "titulo - original title opinion.pdf").write_bytes(_parecer_pdf())
    (files / "pocos - rrc sintetico.csv").write_bytes(_poco_csv())

    (retidos / f"{RETIDA} (daniel prado) lease.pdf").write_bytes(
        _lease_pdf(RETIDA))

    _template_xlsx(AQUI / "template.xlsx")
    return AQUI


if __name__ == "__main__":
    print(gerar())
