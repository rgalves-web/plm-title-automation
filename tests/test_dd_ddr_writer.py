"""Tests for src/dd/ddr_writer.py against a synthetic-but-structurally-faithful
DDR template (mirrors the real "DDR - Template (PLM).xlsx" facts confirmed by
direct inspection: sheet names, card-block heights/gaps, formula shapes) —
never the real client template, matching the synthetic-fixture convention used
in test_dd_manifest_builder.py.
"""
from __future__ import annotations

from fractions import Fraction

import pytest
from openpyxl import Workbook, load_workbook

from src.dd.ddr_writer import (
    DELETE_TAB_NAME,
    OWNERSHIP_1H_SHEET,
    fill_ddr_template,
)
from src.dd.manifest import ReferenceContractRef, RepresentedInterest, UnitManifest, WellRef
from src.dd.unit_aggregation import LeaseTractRow, aggregate_unit, reconcile

# The actual triage signals MOR emits (confirmed in test_excel_writer.py) —
# bare "REVIEW"/"INCLUDE" are too generic and collide with legitimate prose
# (sheet titles like "PLM Title Review", scope text like "Review of the
# chain of title...").
TRIAGE_STRINGS = ("REVIEW REQUIRED", "REDLINE", "DUPLICATE OF")


def _build_synthetic_ddr_template(path) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    # --- Summary ---
    summary = wb.create_sheet("Summary")
    labels = [
        "Client:", "Acquisition:", "Operator(s):", "State:", "Counties:",
        "DSU STR:", "PLM Landman:", "Date Submitted:", "Potential Defect(s):",
        "Outstanding Curative Requirements:", "Liens and Encumbrances:", "Section Notes:",
    ]
    for i, label in enumerate(labels, start=2):
        summary.cell(i, 1).value = label
        summary.cell(i, 2).value = "EXAMPLE VALUE"

    # --- Ownership Report ---
    ownership_report = wb.create_sheet("Ownership Report")
    ownership_report.merge_cells("A1:X1")
    ownership_report.cell(1, 1).value = "Verified Lease Ownership Interests EXAMPLE"
    ownership_report.merge_cells("A3:E3")
    ownership_report.cell(3, 1).value = "Unit Verified Totals (EXAMPLE UNIT)"
    headers_24 = [
        "Unit Tract No.", "Lease No.", "Interest Type", "WI Owner", "Lessor Name",
        "Lessee Name", "Royalty", "Primary Term(Months)", "Start Date", "End Date",
        "Record Book", "Record Page", "Lease Type", "County", "Legal Description",
        "Tract Depth:From", "Tract Depth:To", "Tract Gross Acres", "MI Value", "WI Value",
        "Company Net Acres", "NRI Value", "ORRI Burdens", "Company Net Revenue Acres",
    ]
    for col, header in enumerate(headers_24, start=1):
        ownership_report.cell(7, col).value = header
    ownership_report.cell(5, 2).value = "=SUM(U8:U24)"
    ownership_report.cell(5, 3).value = "=SUM(X8:X24)"
    ownership_report.cell(5, 4).value = "=B5/A5"
    ownership_report.cell(5, 5).value = "=C5/A5"
    for r in range(8, 25):
        ownership_report.cell(r, 18).value = 629.53
        ownership_report.cell(r, 19).value = "=7/1280"
        ownership_report.cell(r, 20).value = 1
        ownership_report.cell(r, 21).value = f"=R{r}*S{r}*T{r}"
        ownership_report.cell(r, 22).value = 0.75
        ownership_report.cell(r, 24).value = f"=U{r}*V{r}"

    # --- Ownership - Wells ---
    ownership_wells = wb.create_sheet("Ownership - Wells")
    well_headers = [
        "API#", "Well Name", "Operator", "Well Type", "Well Status", "Drill Type",
        "Total Depth Drilled", "Spud Date", "Completion Date", "Target Zone",
        "Gaps in Production", "Unit Size", "Seller WI(PSA)", "Seller NRI(PSA)",
        "Seller Wellbore WI(Verified)", "WI Delta", "Seller Wellbore NRI(Verified)", "NRI Delta",
    ]
    for col, header in enumerate(well_headers, start=1):
        ownership_wells.cell(2, col).value = header
    ownership_wells.cell(3, 16).value = "=M3-O3"
    ownership_wells.cell(3, 18).value = "=N3-Q3"
    # Sample well rows with masked API numbers, as in the real template.
    for r in (4, 5):
        ownership_wells.cell(r, 1).value = "42-***-*****"
        ownership_wells.cell(r, 2).value = "*** Unit 3-11 2H"
        ownership_wells.cell(r, 3).value = "*** E&P LLC"

    # --- Lease Review ---
    lease_review = wb.create_sheet("Lease Review")
    lease_headers = [
        "TITLE OPINION LEASE NO.", "TYPE", "LESSOR", "LESSEE", "EFFECTIVE DATE",
        "RECORDING INFORMATION", "LEGAL DESCRIPTION", "TOTAL LEASE ACREAGE",
        "TERM OF LEASE", "ROYALTY", "SHUT-IN ROYALTY CLAUSE", "DEPTH LIMITATIONS",
        "HORIZONTAL PUGH CLAUSE", "VERTICAL PUGH CLAUSE",
    ]
    for col, header in enumerate(lease_headers, start=1):
        lease_review.cell(1, col).value = header
    for r in range(2, 20):
        lease_review.cell(r, 3).value = "EXAMPLE LESSOR"

    # --- Agreements (card layout, 2 example blocks of 19 rows + 1 gap) ---
    agreements = wb.create_sheet("Agreements")
    agreement_labels = [
        "JOA", "EXAMPLE-001", "Operator:", "Non-Operating Parties:", "Effective Date:",
        "Form:", "Contract Area:", "Depth Limitation:", "Non-Consent Penalties:",
        "Carried Non-Consent Interests:", "Preferential Right to Purchase (Y/N):",
        "Maintenance of Uniform Interest (Y/N):", "Tax Partnership:", "Unusual Provisions: ",
        "Area of Mutual Interest:", "Stricken Language:", "Confidentiality Clause:",
        "Amended:", "Additional Notes:",
    ]
    for block_offset in (0, 20):
        for i, label in enumerate(agreement_labels):
            agreements.cell(1 + block_offset + i, 1).value = label
        agreements.cell(1 + block_offset, 2).value = "*** Resources Corporation"
        agreements.cell(3 + block_offset, 2).value = "*** Resources Corporation"
        agreements.merge_cells(start_row=1 + block_offset, start_column=1, end_row=1 + block_offset, end_column=2)
        agreements.merge_cells(start_row=2 + block_offset, start_column=1, end_row=2 + block_offset, end_column=2)

    # --- Unit and Pooling Agreements (card layout, 1 example block of 8 rows) ---
    unit_pooling = wb.create_sheet("Unit and Pooling Agreements")
    unit_pooling_labels = [
        "Unit/Pooling Agreement", "EXAMPLE-UNIT-001", "Type: ", "Recording Information:",
        "Effective Date:", "Unit Description:", "Formations:", "Notes:",
    ]
    for i, label in enumerate(unit_pooling_labels):
        unit_pooling.cell(1 + i, 1).value = label
    unit_pooling.merge_cells("A1:B1")
    unit_pooling.merge_cells("A2:B2")

    # --- Runsheet (title + sample rows mirror the real template's baked-in
    # "636.13 acres ... Reeves County" example and masked instrument rows) ---
    runsheet = wb.create_sheet("Runsheet")
    runsheet.merge_cells("A1:K1")
    runsheet.cell(1, 1).value = (
        "Supplemental Title Runsheet\n\n636.13 acres, more or less, out of "
        "Section **, Block *, Reeves County, Texas"
    )
    runsheet_headers = [
        "NO.", "INSTRUMENT", "INSTRUMENT NUMBER", "EFFECTIVE DATE", "INSTRUMENT DATE",
        "RECORDED DATE", "GRANTOR", "GRANTEE",
    ]
    for col, header in enumerate(runsheet_headers, start=1):
        runsheet.cell(2, col).value = header
    runsheet.merge_cells("A3:K3")
    runsheet.cell(3, 1).value = "EXAMPLE scope note"
    runsheet.cell(4, 1).value = 1
    runsheet.cell(4, 2).value = "Correction Special Warranty Deed"
    runsheet.cell(4, 3).value = "202100****"
    runsheet.cell(5, 1).value = 2
    runsheet.cell(5, 2).value = "Wellbore Only Assignment"
    runsheet.cell(5, 3).value = "202200****"
    runsheet.cell(6, 1).value = 3
    runsheet.cell(6, 2).value = "Merger Certificate"
    runsheet.cell(6, 7).value = "Forest Oil Corporation"

    # --- PLM Title Review (top block rows1-7, 1 example TO card at row9) ---
    plm_title_review = wb.create_sheet("PLM Title Review")
    plm_title_review.cell(1, 1).value = "PLM Title Review"
    plm_title_review.cell(2, 1).value = "Title Certification Dates:"
    plm_title_review.cell(2, 2).value = "9/12/2018 - 11/17/2024"
    plm_title_review.cell(3, 1).value = "Legal Description:"
    plm_title_review.cell(4, 1).value = "Scope of Title Work:"
    plm_title_review.cell(5, 1).value = "Exceptions from Title Review:"
    plm_title_review.cell(6, 1).value = "Landman:"
    plm_title_review.cell(7, 1).value = "Date:"
    plm_title_review.merge_cells("A9:B9")
    plm_title_review.cell(9, 1).value = "Title Opinion Summary"
    to_card_labels = [
        None, "Type of Title Opinion:", "Certification Date(s):",
        "Legal Description:", "Material Outstanding Title Requirements:",
    ]
    for i, label in enumerate(to_card_labels):
        if label is not None:
            plm_title_review.cell(9 + i, 1).value = label
    plm_title_review.cell(10, 2).value = "EXAMPLE Original Title Opinion"

    # --- Disclaimer ---
    disclaimer = wb.create_sheet("Disclaimer")
    for r in range(1, 7):
        disclaimer.cell(r, 1).value = f"Disclaimer line {r}"

    # --- Ownership - 1H (No FO Option) + delete tab ---
    wb.create_sheet(OWNERSHIP_1H_SHEET).cell(1, 1).value = "A.E. Carroll example (1H)"
    wb.create_sheet(DELETE_TAB_NAME).cell(1, 1).value = "A.E. Carroll example (delete)"

    wb.save(path)


def _manifest(**overrides) -> UnitManifest:
    defaults = dict(
        unit_name="BOWEN A UNIT",
        operator="EOG RESOURCES INC",
        county="MONTAGUE",
        seller_name="EOG RESOURCES INC",
        unit_gross_acres="629.53",
        represented_interest=RepresentedInterest(current_wi="0.9975", current_nri="0.78813617"),
        wells=[WellRef(well_name="BOWEN A UNIT #1H", current_wi="0.9975", current_nri="0.78813617")],
    )
    defaults.update(overrides)
    return UnitManifest(**defaults)


def _aggregation_and_reconciliation(manifest: UnitManifest):
    rows = [LeaseTractRow("L1", Fraction(manifest.unit_gross_acres), Fraction(7, 1280), Fraction(1), Fraction("0.25"))]
    aggregation = aggregate_unit(Fraction(manifest.unit_gross_acres), rows)
    reconciliation = reconcile(
        aggregation,
        manifest.represented_interest.current_wi,
        manifest.represented_interest.current_nri,
    )
    return aggregation, reconciliation


def _fill(tmp_path, manifest, reference_contracts=None, to_summaries=None, lease_review_rows=None,
          runsheet_clean_note=None):
    template_path = tmp_path / "template.xlsx"
    _build_synthetic_ddr_template(template_path)
    if reference_contracts is not None:
        manifest = manifest.model_copy(update={"reference_contracts": reference_contracts})
    aggregation, reconciliation = _aggregation_and_reconciliation(manifest)
    return fill_ddr_template(
        template_path=template_path,
        manifest=manifest,
        output_dir=tmp_path / "out",
        aggregation=aggregation,
        reconciliation=reconciliation,
        lease_review_rows=lease_review_rows or [],
        runsheet_records=[],
        runsheet_note="Supplemental scope note",
        to_summaries=to_summaries or [],
        runsheet_clean_note=runsheet_clean_note,
    )


# ---------------------------------------------------------------------------
# Sheet removal
# ---------------------------------------------------------------------------

def test_delete_tab_removed_unconditionally(tmp_path):
    result = _fill(tmp_path, _manifest())
    wb = load_workbook(result.output_path)
    assert DELETE_TAB_NAME not in wb.sheetnames


def test_ownership_1h_retained_unconditionally(tmp_path):
    # Peer DDRs keep this tab regardless of farmout scenario; its example
    # text is handled by the placeholder scrub, not by tab removal.
    for manifest in (_manifest(), _manifest(has_farmout_option_scenario=True)):
        result = _fill(tmp_path, manifest)
        wb = load_workbook(result.output_path)
        assert OWNERSHIP_1H_SHEET in wb.sheetnames


def test_disclaimer_copied_through_unchanged(tmp_path):
    result = _fill(tmp_path, _manifest())
    wb = load_workbook(result.output_path)
    disclaimer = wb["Disclaimer"]
    assert [disclaimer.cell(r, 1).value for r in range(1, 7)] == [f"Disclaimer line {r}" for r in range(1, 7)]


# ---------------------------------------------------------------------------
# Card-block insertion — regression test for the off-by-one bug where the
# inserted block's row_template was stamped one row too early (onto the gap
# row), leaving stale template text (e.g. "JOA") behind.
# ---------------------------------------------------------------------------

def test_agreements_card_block_insertion_leaves_no_leftover_template_text(tmp_path):
    contracts = [
        ReferenceContractRef(file_number="K1400170", record_type="FARMOUT AGREEMENT", lessor="A", lessee="EOG"),
        ReferenceContractRef(file_number="K1400170-AAA", record_type="ASSIGNMENT", lessor="B", lessee="EOG"),
        ReferenceContractRef(file_number="K1400430", record_type="ASSIGNMENT", lessor="C", lessee="EOG"),
    ]
    result = _fill(tmp_path, _manifest(), reference_contracts=contracts)
    wb = load_workbook(result.output_path)
    ws = wb["Agreements"]

    # Gap row directly before the 3rd (inserted) block must be blank.
    assert ws.cell(40, 1).value is None
    assert ws.cell(40, 2).value is None

    # The 3rd block starts at row 41 with the correct title/file number.
    assert ws.cell(41, 1).value == "ASSIGNMENT"
    assert ws.cell(42, 1).value == "K1400430"
    assert ws.cell(43, 2).value == "EOG"  # Operator: row
    assert ws.cell(44, 2).value == "C"  # Counterparty: row

    # No row in the new block carries leftover example values from the template.
    for r in range(41, 60):
        for col in (1, 2):
            value = ws.cell(r, col).value
            assert value != "*** Resources Corporation"
            assert value != "JOA"


def test_plm_title_review_card_block_insertion_leaves_no_leftover_template_text(tmp_path):
    to_summaries = [
        {"file_number": "T0141486", "certification_date": "2008-05-02", "type_description": "Original Title Opinion", "legal_description": "Tract A", "requirements": "-", "confidence": 0.9},
        {"file_number": "T0141487", "certification_date": "2008-05-02", "type_description": "Original Title Opinion", "legal_description": "Tract B", "requirements": "-", "confidence": 0.9},
    ]
    result = _fill(tmp_path, _manifest(), to_summaries=to_summaries)
    wb = load_workbook(result.output_path)
    ws = wb["PLM Title Review"]

    # Gap row directly before the 2nd (inserted) card must be blank.
    assert ws.cell(14, 1).value is None
    assert ws.cell(14, 2).value is None

    # The 2nd card starts at row 15 with the correct content.
    assert ws.cell(15, 1).value == "Title Opinion Summary"
    assert ws.cell(16, 2).value == "Original Title Opinion"
    assert ws.cell(17, 2).value == "2008-05-02"
    assert ws.cell(18, 2).value == "Tract B"

    assert ws.cell(15, 2).value != "EXAMPLE Original Title Opinion"


def test_card_block_deletion_when_fewer_entries_than_template(tmp_path):
    # No UNIT AGREEMENT contracts -> Unit and Pooling Agreements gets 0 entries,
    # template's 1 example block must be cleared, not left with example values.
    result = _fill(tmp_path, _manifest(), reference_contracts=[])
    wb = load_workbook(result.output_path)
    ws = wb["Unit and Pooling Agreements"]
    assert ws.cell(1, 1).value is None  # example title ("Unit/Pooling Agreement") cleared
    assert ws.cell(2, 1).value is None  # example identifier cleared
    for col in range(2, ws.max_column + 1):
        assert ws.cell(1, col).value is None


# ---------------------------------------------------------------------------
# No triage strings in client-facing cells
# ---------------------------------------------------------------------------

def test_no_triage_strings_in_any_client_facing_cell(tmp_path):
    contracts = [
        ReferenceContractRef(file_number="K1400170", record_type="FARMOUT AGREEMENT", lessor="A", lessee="EOG"),
        ReferenceContractRef(file_number="K1400725", record_type="UNIT AGREEMENT", lessor="D", lessee="EOG"),
    ]
    to_summaries = [
        {"file_number": "T0141486", "certification_date": "2008-05-02", "type_description": "Original Title Opinion", "legal_description": "Tract A", "requirements": "-", "confidence": 0.5},
    ]
    lease_review_rows = [{"lease_no": "L1", "type": "OGL", "lessor": "Bowen", "lessee": "EOG", "confidence": 0.5}]
    result = _fill(tmp_path, _manifest(), reference_contracts=contracts, to_summaries=to_summaries, lease_review_rows=lease_review_rows)
    wb = load_workbook(result.output_path)

    for sheet_name in wb.sheetnames:
        if sheet_name == "DD_Audit_Trail":
            continue
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str):
                    for triage_word in TRIAGE_STRINGS:
                        assert triage_word not in value.upper(), (sheet_name, cell.coordinate, value)


# ---------------------------------------------------------------------------
# Audit trail
# ---------------------------------------------------------------------------

def test_audit_trail_sheet_present_and_hidden(tmp_path):
    result = _fill(tmp_path, _manifest())
    wb = load_workbook(result.output_path)
    assert "DD_Audit_Trail" in wb.sheetnames
    assert wb["DD_Audit_Trail"].sheet_state == "hidden"
    assert len(result.audit_rows) > 0


# ---------------------------------------------------------------------------
# Placeholder scrub + Maxwell-style smoke assertions
# ---------------------------------------------------------------------------

from src.dd.ddr_writer import PLACEHOLDER_MARKERS  # noqa: E402


def test_no_placeholder_marker_survives_in_any_tab(tmp_path):
    # The synthetic template bakes in every marker family the real template
    # carries: "***" (Agreements/Wells/Runsheet), "Reeves"/"636.13" (Runsheet
    # title), "3-11" (Wells sample names), "A.E. Carroll" (1H tab),
    # "Forest Oil" (Runsheet sample grantor). None may survive the fill.
    result = _fill(tmp_path, _manifest())
    wb = load_workbook(result.output_path)
    for sheet_name in wb.sheetnames:
        if sheet_name == "DD_Audit_Trail":
            continue
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str):
                    for marker in PLACEHOLDER_MARKERS:
                        assert marker not in value, (sheet_name, cell.coordinate, value)


def test_runsheet_clean_note_row_written_when_no_records(tmp_path):
    note = (
        "No instruments of record affecting EOG Resources' interest were recorded "
        "between the controlling certification dates and present."
    )
    result = _fill(tmp_path, _manifest(), runsheet_clean_note=note)
    wb = load_workbook(result.output_path)
    ws = wb["Runsheet"]
    assert ws.cell(4, 1).value == "-"
    assert ws.cell(4, 2).value == note
    # The template's masked sample rows must be gone.
    assert ws.cell(5, 2).value is None
    assert ws.cell(6, 2).value is None


def test_summary_review_narrative_defaults_to_none_identified(tmp_path):
    result = _fill(tmp_path, _manifest())
    wb = load_workbook(result.output_path)
    ws = wb["Summary"]
    # Rows 10-12: Potential Defect(s) / Outstanding Curative / Liens.
    for row in (10, 11, 12):
        assert ws.cell(row, 2).value == "None identified", row


def test_unit_totals_computed_not_hardcoded(tmp_path):
    lease_review_rows = [{"lease_no": "L1", "lessor": "Cox", "lessee": "EOG",
                          "mi_value": Fraction(7, 1280), "wi_value": Fraction(1)}]
    result = _fill(tmp_path, _manifest(), lease_review_rows=lease_review_rows)
    wb = load_workbook(result.output_path)

    # Ownership Report totals are live SUM/ratio formulas over the detail rows.
    ownership = wb["Ownership Report"]
    assert str(ownership.cell(5, 2).value).startswith("=SUM(U")
    assert str(ownership.cell(5, 3).value).startswith("=SUM(X")
    assert ownership.cell(5, 4).value == "=B5/A5"

    # Ownership - Wells verified WI equals the Fraction-exact aggregation
    # result (7/1280 for the fixture rows), written as its float image.
    wells = wb["Ownership - Wells"]
    assert wells.cell(3, 15).value == pytest.approx(float(Fraction(7, 1280)))
    assert wells.cell(3, 16).value == "=M3-O3"


def test_unit_pooling_card_carries_enriched_fields(tmp_path):
    dpu = ReferenceContractRef(
        file_number="K1400813",
        record_type="DESIGNATION OF POOLED UNIT",
        lessee="EOG RESOURCES INC",
        recording_info="Volume 491, Page 160",
        effective_date="8/12/2009",
        unit_description="160 acres in the Calhoun County School Land Survey, League 12, A-122",
    )
    result = _fill(tmp_path, _manifest(), reference_contracts=[dpu])
    wb = load_workbook(result.output_path)
    ws = wb["Unit and Pooling Agreements"]
    assert ws.cell(2, 1).value == "K1400813"
    assert ws.cell(3, 2).value == "DESIGNATION OF POOLED UNIT"
    assert ws.cell(4, 2).value == "Volume 491, Page 160"
    assert ws.cell(5, 2).value == "8/12/2009"
    assert "Calhoun County School Land Survey" in ws.cell(6, 2).value
