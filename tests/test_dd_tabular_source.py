"""Client workbook columns resolve by HEADER and fail loudly (systemic fix).

Each test replays a real failure from this project's export drift, where a
hardcoded column position produced silent, confident nonsense instead of an
error.
"""
from __future__ import annotations

import openpyxl
import pytest

from src.dd.tabular_source import (
    ColumnSpec,
    TabularSourceError,
    read_rows,
    resolve_header_columns,
    resolve_sheet,
)

# The XREF's real header: note RECORD TYPE appears twice — once for the lease,
# once for the contract.
_XREF_HEADER = [
    "INCLUDE FLAG", "LEASE NO", "LESSOR", "LESSEE", "ACTIVE FLAG", "RECORD TYPE",
    "RIGHTS HELD", "PROPERTY STATUS", "LEASE DATE", "EFF DATE", "EXP DATE",
    "EXT EXP DATE", "STATE", "COUNTY", "WELL OR CONTRACT NB",
    "WELL OR CONTRACT NAME", "ACTIVE FLAG", "RECORD TYPE", "RIGHTS HELD",
]
_XREF_SPECS = [
    ColumnSpec("lease_no", ["LEASE NO"]),
    ColumnSpec("contract_no", ["WELL OR CONTRACT NB"]),
    ColumnSpec("contract_name", ["WELL OR CONTRACT NAME"]),
    ColumnSpec("record_type", ["RECORD TYPE"], after="contract_no"),
]


def test_resolves_columns_by_name():
    got = resolve_header_columns(_XREF_HEADER, _XREF_SPECS, source="XREF")
    assert got["lease_no"] == 1 and got["contract_no"] == 14


def test_duplicate_header_resolves_to_the_one_after_its_anchor():
    # RECORD TYPE at 5 is the LEASE's type; at 17 it is the CONTRACT's. Taking
    # the first match classified every contract by its lease type.
    got = resolve_header_columns(_XREF_HEADER, _XREF_SPECS, source="XREF")
    assert got["record_type"] == 17


def test_missing_column_raises_and_names_it():
    # Real case: "UNIT LEVEL NAME" (column AA) vanished from the export.
    specs = _XREF_SPECS + [ColumnSpec("unit_name", ["UNIT LEVEL NAME"])]
    with pytest.raises(TabularSourceError) as exc:
        resolve_header_columns(_XREF_HEADER, specs, source="XREF")
    message = str(exc.value)
    assert "unit_name" in message and "UNIT LEVEL NAME" in message
    assert "LEASE NO" in message           # shows the header actually seen
    assert "reading by position" in message  # tells the reader what to do


def test_moved_column_is_read_correctly_not_by_stale_position():
    # Real case: Well List "UNIT NAME" moved from index 28 to 23. A positional
    # read matched no wells at all -> no represented interest -> A5 = 0.
    old = ["OPERATOR", "COUNTY", "WELL NAME"] + [f"C{i}" for i in range(3, 28)] + ["UNIT NAME"]
    new = ["OPERATOR", "COUNTY", "WELL NAME"] + [f"C{i}" for i in range(3, 23)] + ["UNIT NAME"]
    specs = [ColumnSpec("unit_name", ["UNIT NAME"])]
    assert resolve_header_columns(old, specs)["unit_name"] == 28
    assert resolve_header_columns(new, specs)["unit_name"] == 23


def test_optional_column_may_be_absent():
    specs = _XREF_SPECS + [ColumnSpec("unit_name", ["UNIT LEVEL NAME"], required=False)]
    got = resolve_header_columns(_XREF_HEADER, specs, source="XREF")
    assert "unit_name" not in got


def test_alternate_spellings_are_accepted_best_first():
    header = ["LEASE NO", "WELL OR CONTRACT NB", "WELL OR CONTRACT NAME", "RECORD TYPE3"]
    specs = [ColumnSpec("record_type", ["RECORD TYPE3", "RECORD TYPE"])]
    assert resolve_header_columns(header, specs)["record_type"] == 3


def _wb(sheets: dict[str, list[list]]):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        ws = workbook.create_sheet(name)
        for row in rows:
            ws.append(row)
    return workbook


_WELL_SPECS = [ColumnSpec("unit_name", ["UNIT NAME"]), ColumnSpec("wi", ["CURRENT WI"])]


def test_sheet_is_found_by_header_when_the_tab_was_renamed():
    # Real case: the Well List's "Wells" sheet became "Sheet1".
    workbook = _wb({"Sheet1": [["UNIT NAME", "CURRENT WI"], ["MAXWELL A UNIT", 0.99]]})
    ws = resolve_sheet(workbook, _WELL_SPECS, preferred="Wells", source="Well List")
    assert ws.title == "Sheet1"


def test_preferred_sheet_wins_when_present():
    workbook = _wb({"Wells": [["UNIT NAME", "CURRENT WI"], ["A", 1]],
                    "Sheet1": [["UNIT NAME", "CURRENT WI"], ["B", 2]]})
    assert resolve_sheet(workbook, _WELL_SPECS, preferred="Wells").title == "Wells"


def test_excluded_sheet_is_never_chosen():
    # "Wells Removed" holds out-of-scope wells; picking it would ship removals.
    workbook = _wb({"Wells Removed": [["UNIT NAME", "CURRENT WI"], ["X", 1]],
                    "Sheet1": [["UNIT NAME", "CURRENT WI"], ["Y", 2]]})
    ws = resolve_sheet(workbook, _WELL_SPECS, preferred="Wells",
                       exclude=("removed",), source="Well List")
    assert ws.title == "Sheet1"


def test_no_matching_sheet_raises_listing_what_was_tried():
    workbook = _wb({"Notes": [["SOMETHING", "ELSE"], ["a", "b"]]})
    with pytest.raises(TabularSourceError) as exc:
        resolve_sheet(workbook, _WELL_SPECS, preferred="Wells", source="Well List")
    assert "Well List" in str(exc.value) and "Notes" in str(exc.value)


def test_read_rows_uses_the_resolved_map():
    workbook = _wb({"Sheet1": [["UNIT NAME", "CURRENT WI"], ["MAXWELL A UNIT", 0.99]]})
    ws = workbook["Sheet1"]
    columns = resolve_header_columns(
        next(ws.iter_rows(min_row=1, max_row=1, values_only=True)), _WELL_SPECS)
    rows = read_rows(ws, columns)
    assert rows == [{"unit_name": "MAXWELL A UNIT", "wi": 0.99}]


def test_short_row_does_not_raise_index_error():
    # A ragged row must yield None, not blow up mid-read.
    workbook = _wb({"Sheet1": [["UNIT NAME", "CURRENT WI"], ["MAXWELL A UNIT"]]})
    ws = workbook["Sheet1"]
    columns = resolve_header_columns(
        next(ws.iter_rows(min_row=1, max_row=1, values_only=True)), _WELL_SPECS)
    assert read_rows(ws, columns)[0]["wi"] is None
