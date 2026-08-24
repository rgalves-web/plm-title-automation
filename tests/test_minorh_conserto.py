"""Conserto definitivo (ordem de 23/08) — testes do encanamento novo.

Etapas B e C: a caneta protegida (fórmulas, hash, write-roots), o teto por
unit, as horas do operador, o kit carbon-copy com md5, o publicar congelado, a
quarentena por identidade e a trilha com supersessão. Nenhuma asserção conclui
title — só bytes, contagens e recusas.
"""
from __future__ import annotations

import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from src.examinador.cache import PageCache
from src.examinador.corpus import Corpus
from src.examinador.formulas import (FormulaAlterada, erro_celula_de_formula,
                                     formula_cells, formulas_hash)
from src.examinador.ledger import Ledger
from src.examinador.run import write_workbook
from src.minorh.estado import Estado


def _template_com_formula(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ownership - Worksheet"
    ws["A1"] = "rotulo"
    ws["B2"] = "=SUM(C2:D2)"          # a célula azul
    ws["C2"] = 1
    ws["D2"] = 2
    wb.save(path)
    wb.close()
    return path


# --- fórmulas: indexação e hash ---------------------------------------------

def test_formula_cells_indexa_so_as_formulas(tmp_path: Path):
    tpl = _template_com_formula(tmp_path / "tpl.xlsx")
    got = formula_cells(tpl)
    assert got == {("Ownership - Worksheet", "B2"): "=SUM(C2:D2)"}


def test_o_hash_muda_quando_uma_formula_muda(tmp_path: Path):
    tpl = _template_com_formula(tmp_path / "tpl.xlsx")
    igual = _template_com_formula(tmp_path / "copia.xlsx")
    assert formulas_hash(tpl) == formulas_hash(igual)
    wb = openpyxl.load_workbook(igual)
    wb["Ownership - Worksheet"]["B2"] = "=SUM(C2:E2)"
    wb.save(igual)
    wb.close()
    assert formulas_hash(tpl) != formulas_hash(igual)


# --- B1/B2: a caneta recusa e o hash prova ----------------------------------

def test_write_workbook_preserva_formula_e_prova_por_hash(tmp_path: Path):
    tpl = _template_com_formula(tmp_path / "tpl.xlsx")
    ledger = Ledger()
    ledger.write_cell("Ownership - Worksheet", "B2", "0.80", "d.pdf", 1,
                      "trecho", "alta")
    ledger.write_cell("Ownership - Worksheet", "E2", "0.80", "d.pdf", 1,
                      "trecho", "alta")
    out = write_workbook(ledger, tmp_path / "ddr.xlsx", template=tpl)
    assert formulas_hash(out) == formulas_hash(tpl)
    wb = openpyxl.load_workbook(out)
    assert wb["Ownership - Worksheet"]["B2"].value == "=SUM(C2:D2)"
    assert wb["Ownership - Worksheet"]["E2"].value == "0.80"
    wb.close()
    recusadas = json.loads((tmp_path / "formulas_preservadas.json")
                           .read_text(encoding="utf-8"))
    assert any("B2" in r for r in recusadas)


def test_faixa_de_branco_nao_apaga_formula(tmp_path: Path):
    tpl = _template_com_formula(tmp_path / "tpl.xlsx")
    ledger = Ledger()
    ledger.record_blank("Ownership - Worksheet", "A1:D2", "limpeza", "op")
    out = write_workbook(ledger, tmp_path / "ddr.xlsx", template=tpl)
    wb = openpyxl.load_workbook(out)
    ws = wb["Ownership - Worksheet"]
    assert ws["B2"].value == "=SUM(C2:D2)"      # a fórmula fica
    assert ws["A1"].value is None               # o resto da faixa esvazia
    assert ws["C2"].value is None
    wb.close()


def test_o_erro_nomeado_volta_ao_modelo_na_sessao(tmp_path: Path):
    from src.minorh.ship import _dispatch_escrita
    tpl = _template_com_formula(tmp_path / "tpl.xlsx")
    ledger = Ledger()
    d = _dispatch_escrita(ledger, tpl)
    with pytest.raises(ValueError) as err:
        d["write_cell"]("Ownership - Worksheet", "B2", "0.80", "d.pdf", "1",
                        "t", "alta")
    assert "contém fórmula do template" in str(err.value)
    assert "célula branca" in str(err.value)
    assert ledger.cells == []                   # nada gravado
    with pytest.raises(ValueError):
        d["record_blank"]("Ownership - Worksheet", "b2", "m", "o")
    # célula branca passa; faixa passa (write_workbook preserva por dentro)
    d["write_cell"]("Ownership - Worksheet", "E2", "0.80", "d.pdf", "1", "t",
                    "alta")
    d["record_blank"]("Ownership - Worksheet", "A8:X8", "sobras", "op")
    assert len(ledger.cells) == 1 and len(ledger.blanks) == 1


def test_mensagem_de_erro_nomeada():
    got = erro_celula_de_formula("Aba", "B2")
    assert got.startswith("Aba!B2 contém fórmula do template")


# --- C4: trilha com supersessão ---------------------------------------------

def test_trilha_marca_a_entrada_supersedida(tmp_path: Path):
    tpl = _template_com_formula(tmp_path / "tpl.xlsx")
    ledger = Ledger()
    ledger.write_cell("Ownership - Worksheet", "E2", "primeiro", "d.pdf", 1,
                      "t", "alta")
    ledger.write_cell("Ownership - Worksheet", "E2", "segundo", "d.pdf", 2,
                      "t", "alta")
    ledger.write_cell("Ownership - Worksheet", "F2", "unico", "d.pdf", 1,
                      "t", "alta")
    out = write_workbook(ledger, tmp_path / "ddr.xlsx", template=tpl)
    wb = openpyxl.load_workbook(out)
    linhas = list(wb["DD_Audit_Trail"].iter_rows(min_row=2, values_only=True))
    wb.close()
    por_valor = {str(r[2]): str(r[9] or "") for r in linhas}
    assert por_valor["primeiro"] == "supersedida_por: 2"
    assert por_valor["segundo"] == ""
    assert por_valor["unico"] == ""


# --- B4: teto por unit -------------------------------------------------------

def _ctx_estado(tmp_path: Path, gasto: float) -> dict[str, Any]:
    from src.minorh.memoria import Memoria
    estado = Estado(tmp_path)
    if gasto:
        estado.registrar_execucao("unidade", {"entrada": 1, "saida": 1,
                                              "cache_escrita": 0,
                                              "cache_leitura": 0,
                                              "custo_usd": gasto})
    memoria = Memoria(tmp_path / "_minorh" / "memoria.json")
    return {"estado": estado, "memoria": memoria}


def test_orcamento_e_o_minimo_entre_comando_e_o_que_resta(tmp_path, monkeypatch):
    from src.minorh.cli import _orcamento
    monkeypatch.setenv("MINORH_TETO_UNIT_USD", "25")
    ctx = _ctx_estado(tmp_path, 20.0)
    assert _orcamento(ctx, 40.0) == 5.0          # resta menos que o comando
    assert _orcamento(ctx, 3.0) == 3.0           # o comando pede menos


def test_teto_batido_recusa_limpo_e_diz_como_autorizar(tmp_path, monkeypatch):
    from src.minorh.cli import _orcamento
    monkeypatch.setenv("MINORH_TETO_UNIT_USD", "1")
    ctx = _ctx_estado(tmp_path, 1.5)
    ctx["memoria"].pendente("lease x", "lessor", "PED-001")
    with pytest.raises(SystemExit) as err:
        _orcamento(ctx, 40.0)
    texto = str(err.value)
    assert "TETO DA UNIDADE ATINGIDO" in texto
    assert "$1.5" in texto
    assert "pendente" in texto
    assert "MINORH_TETO_UNIT_USD=" in texto      # como autorizar exceção


# --- B5: horas do operador ---------------------------------------------------

def test_horas_acumulam_e_sobrevivem(tmp_path: Path):
    estado = Estado(tmp_path)
    estado.declarar_censo(1, "conta")            # cria o estado
    estado.registrar_horas(0.5, "aprovações")
    estado.registrar_horas(0.25)
    relido = Estado(tmp_path)
    assert relido.horas_acumuladas() == 0.75
    assert relido.horas[0]["nota"] == "aprovações"


# --- C1: kit — alocação do modelo, cópia da caneta --------------------------

def _corpus(tmp_path: Path, nomes: dict[str, bytes]) -> Corpus:
    files = tmp_path / "FILES"
    files.mkdir(exist_ok=True)
    for nome, dados in nomes.items():
        (files / nome).write_bytes(dados)
    return Corpus(files, cache=PageCache(tmp_path / "cache"))


@dataclass
class _Uso:
    input_tokens: int = 10
    output_tokens: int = 5
    cache_creation_input_tokens: int = 0
    cache_read_input_tokens: int = 0


@dataclass
class _Msg:
    content: list[Any]
    usage: _Uso = field(default_factory=_Uso)


@dataclass
class _Use:
    name: str
    input: dict[str, Any]
    id: str = "u1"
    type: str = "tool_use"


@dataclass
class _Texto:
    text: str
    type: str = "text"


class _ClientRoteirizado:
    def __init__(self, roteiro: list[list[Any]]) -> None:
        self.roteiro = roteiro
        self.messages = self

    def create(self, **kw: Any) -> _Msg:
        return _Msg(content=self.roteiro.pop(0))


def test_alocar_valida_forma_e_registra_decisoes(tmp_path: Path):
    from src.minorh.kit import alocar
    from src.minorh.memoria import Memoria
    corpus = _corpus(tmp_path, {"lease.pdf": b"a", "print.png": b"b"})
    memoria = Memoria(tmp_path / "_minorh" / "memoria.json")
    client = _ClientRoteirizado([
        [_Use("alocar_arquivo", {"documento": "nao-existe.pdf",
                                 "destino": "Leases", "por_que": "x"})],
        [_Use("alocar_arquivo", {"documento": "lease.pdf",
                                 "destino": "Pasta Inventada",
                                 "por_que": "x"})],
        [_Use("alocar_arquivo", {"documento": "lease.pdf",
                                 "destino": "Leases",
                                 "novo_nome": "L1 - 100.5 - Ana OGL.pdf",
                                 "por_que": "lease da unidade"})],
        [_Use("nao_alocar", {"documento": "print.png",
                             "motivo": "screenshot, nao e o documento"})],
        [_Texto("feito")],
    ])
    alocacao, log = alocar(pasta=tmp_path, corpus=corpus, memoria=memoria,
                           modelo="claude-opus-5", dispatch_leitura={},
                           client=client)
    assert alocacao["alocados"] == [{
        "documento": "lease.pdf", "destino": "Leases",
        "novo_nome": "L1 - 100.5 - Ana OGL.pdf",
        "por_que": "lease da unidade"}]
    assert alocacao["fora_do_kit"][0]["documento"] == "print.png"
    assert alocacao["sem_decisao"] == []
    # os dois primeiros turnos viraram erro devolvido ao modelo
    erros = [c for t in log.turns for c in t.calls if c["erro"]]
    assert len(erros) == 2


def test_montar_copia_com_md5_e_declara_o_resto(tmp_path: Path):
    from src.minorh.kit import PASTAS, RAIZ_KIT, montar
    corpus = _corpus(tmp_path, {"lease.pdf": b"conteudo do lease",
                                "solto.pdf": b"x"})
    ddr = tmp_path / "DDR.xlsx"
    ddr.write_bytes(b"bytes do ddr")
    alocacao = {"alocados": [{"documento": "lease.pdf", "destino": "Leases",
                              "novo_nome": "L1 - 100.5 - Ana OGL.pdf",
                              "por_que": "ok"}],
                "fora_do_kit": [],
                "sem_decisao": ["solto.pdf"]}
    got = montar(pasta=tmp_path, corpus=corpus, alocacao=alocacao, ddr=ddr)
    kit = tmp_path / "_minorh" / "ship" / "kit" / RAIZ_KIT
    for pasta in PASTAS:
        assert (kit / pasta).is_dir()            # carbon copy: as 7 pastas
    assert (kit / "Leases" / "L1 - 100.5 - Ana OGL.pdf").read_bytes() == \
        b"conteudo do lease"
    ddr_kit = kit / "PLM Report" / f"DDR - {tmp_path.name.title()}.xlsx"
    assert ddr_kit.read_bytes() == b"bytes do ddr"
    assert got["md5_roundtrip"] == "2/2"
    assert got["sem_decisao"] == ["solto.pdf"]   # declarado, não escondido
    assert "ausente" in got["checklist"]
    rel = json.loads((tmp_path / "_minorh" / "ship" / "RELATORIO_KIT.json")
                     .read_text(encoding="utf-8"))
    assert rel["md5_roundtrip"] == "2/2"


def test_publicar_congela_o_pacote_entregue(tmp_path: Path):
    from src.minorh.kit import KitCongelado, montar, publicar
    corpus = _corpus(tmp_path, {"lease.pdf": b"conteudo"})
    ddr = tmp_path / "DDR.xlsx"
    ddr.write_bytes(b"ddr")
    montar(pasta=tmp_path, corpus=corpus, ddr=ddr,
           alocacao={"alocados": [{"documento": "lease.pdf",
                                   "destino": "Leases",
                                   "novo_nome": "lease.pdf", "por_que": "x"}],
                     "fora_do_kit": [], "sem_decisao": []})
    got = publicar(tmp_path)
    assert (tmp_path / "SUBMISSION PLM" / "1. Landman Submitted" /
            "Leases" / "lease.pdf").read_bytes() == b"conteudo"
    assert got["md5_roundtrip"] == "2/2"
    with pytest.raises(KitCongelado):
        publicar(tmp_path)                       # pacote entregue é congelado


# --- C2: quarentena por identidade ------------------------------------------

def test_quarentena_deixa_o_kit_do_proprio_run_e_bloqueia_o_gabarito(
        tmp_path: Path):
    from src.knowledge.quarantine import QuarantineBreach, quarantined
    proprio = tmp_path / "_minorh" / "ship" / "kit" / "1. Landman Submitted"
    proprio.mkdir(parents=True)
    (proprio / "a.txt").write_text("meu", encoding="utf-8")
    gabarito = tmp_path / "UNIT" / "1. Landman Submitted"
    gabarito.mkdir(parents=True)
    (gabarito / "b.txt").write_text("entregue a mao", encoding="utf-8")

    with quarantined(allow_roots=(tmp_path / "_minorh",)) as tentativas:
        assert (proprio / "a.txt").read_text(encoding="utf-8") == "meu"
        with pytest.raises(QuarantineBreach):
            (gabarito / "b.txt").read_text(encoding="utf-8")
    assert tentativas == [str(gabarito / "b.txt")]


def test_quarentena_sem_allow_roots_segue_bloqueando_tudo(tmp_path: Path):
    from src.knowledge.quarantine import QuarantineBreach, quarantined
    alvo = tmp_path / "gabarito" / "x.txt"
    alvo.parent.mkdir(parents=True)
    alvo.write_text("g", encoding="utf-8")
    with quarantined():
        with pytest.raises(QuarantineBreach):
            alvo.read_text(encoding="utf-8")


# --- B3: o ship escreve só nos write-roots -----------------------------------

def test_escrita_do_ship_fora_do_root_e_excecao(tmp_path: Path):
    from src.knowledge.write_roots import WriteOutsideRoots, write_scope
    ledger = Ledger()
    tpl = _template_com_formula(tmp_path / "tpl.xlsx")
    dentro = tmp_path / "unidade" / "_minorh"
    with write_scope(dentro, label="teste"):
        write_workbook(ledger, dentro / "ship" / "DDR.xlsx", template=tpl)
        with pytest.raises(WriteOutsideRoots):
            write_workbook(ledger, tmp_path / "fora.xlsx", template=tpl)
