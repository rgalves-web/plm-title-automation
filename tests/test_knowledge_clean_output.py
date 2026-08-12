"""PR-8 — the four gates that should have existed before anything was written.

Every case here is something that actually shipped in a generated DDR on
2026-08-11: 247 cells of another operator's unit, a `delete - …` worksheet, the
wrong Ownership tab name, no audit trail, and a checklist asserting the
hand-built kit's numbers.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from src.knowledge.clean_output import (AUDIT_TRAIL_SHEET, NotCleanError,
                                        PLACEHOLDER_MARKERS, assert_clean,
                                        check_checklist, check_clean_output)
from src.knowledge.quarantine import QuarantineBreach, is_quarantined, quarantined

EXPECTED = {"Summary", "Lease Review", AUDIT_TRAIL_SHEET}


def _book(tmp_path: Path, *, sheets=("Summary", "Lease Review", AUDIT_TRAIL_SHEET),
          contaminate: str | None = None) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name in sheets:
        ws = wb.create_sheet(name)
        ws["A1"] = "clean value"
    if contaminate:
        wb[sheets[0]]["B2"] = contaminate
    path = tmp_path / "DDR.xlsx"
    wb.save(path)
    return path


# --- 1. clean output, over bytes re-read from disk -------------------------

@pytest.mark.parametrize("marker", PLACEHOLDER_MARKERS)
def test_every_template_marker_fails_the_gate(tmp_path, marker):
    report = check_clean_output(_book(tmp_path, contaminate=f"x {marker} y"),
                                expected_sheets=EXPECTED)
    assert report.failed
    assert any(f.kind == "placeholder" for f in report.findings)


def test_a_delete_tab_fails(tmp_path):
    path = _book(tmp_path, sheets=("Summary", "Lease Review", AUDIT_TRAIL_SHEET,
                                   "delete - A.E. Carroll 4H"))
    report = check_clean_output(path, expected_sheets=EXPECTED)
    assert any(f.kind == "delete_tab" for f in report.findings)


def test_a_sheet_set_that_differs_from_the_template_fails_both_ways(tmp_path):
    # The real failure was BOTH: "Ownership Report" shipped where PLM's
    # instruction sheet names "Ownership - Worksheet".
    path = _book(tmp_path, sheets=("Summary", "Lease Review", AUDIT_TRAIL_SHEET,
                                   "Ownership Report"))
    report = check_clean_output(path, expected_sheets=EXPECTED | {
        "Ownership - Worksheet"})
    kinds = [(f.kind, f.where) for f in report.findings]
    assert ("sheet_set", "Ownership Report") in kinds        # unexpected present
    assert ("sheet_set", "Ownership - Worksheet") in kinds   # expected absent


def test_a_deliverable_without_the_audit_trail_fails(tmp_path):
    path = _book(tmp_path, sheets=("Summary", "Lease Review"))
    report = check_clean_output(path, expected_sheets={"Summary", "Lease Review"})
    assert any(f.kind == "no_audit_trail" for f in report.findings)


def test_a_clean_deliverable_passes(tmp_path):
    report = assert_clean(_book(tmp_path), expected_sheets=EXPECTED)
    assert not report.failed and report.cells_checked


def test_the_gate_reads_the_file_and_not_the_workbook_that_made_it(tmp_path):
    # A gate that inspects the in-memory object is testing its own intentions.
    path = _book(tmp_path)
    wb = openpyxl.load_workbook(path)
    wb["Summary"]["C3"] = "*** Energy"       # only in memory
    assert not check_clean_output(path, expected_sheets=EXPECTED).failed
    wb.save(path)
    assert check_clean_output(path, expected_sheets=EXPECTED).failed


def test_assert_clean_raises_with_the_findings_named(tmp_path):
    with pytest.raises(NotCleanError, match="Reeves"):
        assert_clean(_book(tmp_path, contaminate="Reeves County"),
                     expected_sheets=EXPECTED)


# --- 2. the checklist asserts ---------------------------------------------

def _checklist(tmp_path: Path, rows: dict[int, tuple[str, str]]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["D1"], ws["E1"] = "Folder", "Completed (Yes/No)", "Notes"
    for row, (answer, note) in rows.items():
        ws.cell(row=row, column=4).value = answer
        ws.cell(row=row, column=5).value = note
    path = tmp_path / "chk.xlsx"
    wb.save(path)
    return path


def test_yes_beside_an_empty_folder_fails(tmp_path):
    path = _checklist(tmp_path, {2: ("Yes", "")})
    report = check_checklist(path, placed={"Leases": []},
                             row_to_folder={2: "Leases"})
    assert any(f.kind == "unearned_yes" for f in report.findings)


def test_a_number_nobody_computed_fails(tmp_path):
    # "32 documents uploaded" beside a folder holding one is a false statement
    # to the client, made by a machine, in a file with our name on it.
    path = _checklist(tmp_path, {12: ("Yes", "32 documents uploaded")})
    report = check_checklist(path, placed={"Title Documents": ["a"]},
                             row_to_folder={12: "Title Documents"},
                             derived_numbers={12: {"1"}})
    assert any(f.kind == "undeclared_number" and "32" in f.detail
               for f in report.findings)


def test_a_derived_number_passes(tmp_path):
    path = _checklist(tmp_path, {12: ("Yes", "1 placed; 18 declared pending")})
    report = check_checklist(path, placed={"Title Documents": ["a"]},
                             row_to_folder={12: "Title Documents"},
                             derived_numbers={12: {"1", "18"}})
    assert not report.failed


def test_no_without_a_pendencia_fails(tmp_path):
    path = _checklist(tmp_path, {15: ("No", "")})
    report = check_checklist(path, placed={"Well Files": []},
                             row_to_folder={15: "Well Files"})
    assert any(f.kind == "undeclared_no" for f in report.findings)


# --- 3. the answer sheet is quarantined ------------------------------------

@pytest.mark.parametrize("path", [
    "docs/JOSH_B_final_2026-08-10.md",
    r"docs\JOSH_B_kickoff_download_list.md",
    "/somewhere/gabarito_josh_b.zip",
])
def test_the_answer_sheets_are_recognised(path):
    assert is_quarantined(path)


def test_a_corpus_document_is_not_quarantined():
    assert not is_quarantined("data/units/_josh_b_replay/corpus/LEASE.TIF")
    assert not is_quarantined("docs/FETCH_BRIEF_SPEC.md")


def test_reading_the_answer_sheet_inside_a_run_raises(tmp_path):
    with pytest.raises(QuarantineBreach, match="not a replay"):
        with quarantined():
            open("docs/JOSH_B_final_2026-08-10.md", encoding="utf-8")


def test_every_read_path_is_guarded_not_just_builtins_open():
    for call in (lambda: Path("docs/JOSH_B_status_2026-08-05.md").read_text(),
                 lambda: Path("docs/JOSH_B_status_2026-08-05.md").read_bytes(),
                 lambda: Path("docs/JOSH_B_status_2026-08-05.md").open()):
        with pytest.raises(QuarantineBreach):
            with quarantined():
                call()


def test_the_guard_is_lifted_when_the_block_ends():
    with quarantined():
        pass
    # The comparison harness reads the gabarito — AFTER the run, outside it.
    assert Path("docs/JOSH_B_final_2026-08-10.md").read_text(encoding="utf-8")


def test_ordinary_reads_inside_the_block_still_work(tmp_path):
    sample = tmp_path / "corpus.txt"
    sample.write_text("lease", encoding="utf-8")
    with quarantined():
        assert sample.read_text(encoding="utf-8") == "lease"


# --- 6. the standing regression -------------------------------------------
#
# From this PR on, the cold-built DDR passes the clean-output gate on every run.
# Bowen and Maxwell stay in READ mode; this is the artefact the system MAKES,
# and it is the one that shipped contaminated.

COLD_DDR = Path("data/units/_josh_b_replay/KIT/DDR - Josh B Unit.xlsx")


@pytest.mark.skipif(not COLD_DDR.exists(),
                    reason="cold kit not built in this checkout (client data)")
def test_the_cold_built_ddr_passes_the_clean_output_gate():
    from src.knowledge.kit_builder import EXPECTED_SHEETS
    report = assert_clean(COLD_DDR, expected_sheets=EXPECTED_SHEETS)
    assert not report.failed, report.summary()


@pytest.mark.skipif(not COLD_DDR.exists(),
                    reason="cold kit not built in this checkout (client data)")
def test_the_cold_built_ddr_carries_its_evidence():
    wb = openpyxl.load_workbook(COLD_DDR, data_only=True)
    try:
        ws = wb[AUDIT_TRAIL_SHEET]
        assert ws.max_row > 1, "the audit trail is empty"
        header = [str(c.value or "").upper() for c in next(ws.iter_rows(max_row=1))]
        assert "CITED TO" in header and "REASON IF BLANK" in header
    finally:
        wb.close()


@pytest.mark.skipif(not COLD_DDR.exists(), reason="cold kit not built")
def test_the_writer_calls_the_gate_so_a_new_write_path_cannot_skip_it():
    import inspect

    from src.knowledge import kit_builder
    assert "assert_clean" in inspect.getsource(kit_builder.write_ddr)
